#!/usr/bin/env python3
from __future__ import annotations

if __package__ in {None, ""}:
    import sys
    from pathlib import Path

    APP_BOOT_DIR = Path(__file__).resolve().parents[2]
    PACKAGE_PARENT = APP_BOOT_DIR.parent
    if str(PACKAGE_PARENT) not in sys.path:
        sys.path.insert(0, str(PACKAGE_PARENT))

import csv
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from cimmyt_app.ce_pipeline.phase01.phase1 import (
    build_column_order,
    read_definition,
    read_sheet_headers_and_rows,
    write_xlsx,
)
from cimmyt_app.ce_pipeline.phase03.script03 import create_phase03_workbook
from cimmyt_app.ce_pipeline.phase04.phase04 import create_phase04_workbook
from cimmyt_app.ce_pipeline.phase05.phase05 import run_ce_phase05
from cimmyt_app.pipeline.common import dataframe_to_geojson
from cimmyt_app.pipeline.model_registry import get_registered_model, write_model_metadata
from cimmyt_app.pipeline.run_model_pipeline import ModelPipelineOutputs, run_model_pipeline
from cimmyt_app.preprocess.runner import (
    PreprocessOutputs,
    resolve_point_uploaded_workbook_mode,
    run_preprocess_pipeline,
)
from cimmyt_app.workbook_preview import (
    CLEAR_SELECTION_ID_SENTINEL,
    PREDICTION_GROUP_KEY_HEADER,
    PREDICTION_GROUP_LABEL_HEADER,
    PREDICTION_GROUP_REP_HEADER,
    PREDICTION_INTERNAL_ROW_ID_HEADER,
    _build_prediction_internal_row_id,
)


@dataclass
class SavedModelPredictionOutputs:
    preprocess: PreprocessOutputs
    model: ModelPipelineOutputs


@dataclass
class CeSavedModelPredictionOutputs:
    summary: dict[str, object]
    geojson: dict[str, object]
    summary_file: Path
    geojson_file: Path


PREDICTION_PROFILE_KEY_HEADER = "Prediction Profile Key"
PREDICTION_PROFILE_LABEL_HEADER = "Prediction Profile"
PREDICTION_PROFILE_COLUMNS_HEADER = "Prediction Profile Columns"
PREDICTION_PROFILE_TARGET_MEAN_HEADER = "Prediction Target Mean Threshold"
PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER = "Prediction Profile Target Mean"
PREDICTION_PROFILE_ABOVE_MEAN_HEADER = "Prediction Profile Above Target Mean"
PREDICTION_PREDICTED_COLUMN = "Grain Yield predicted"
MANUAL_GRID_GEOJSON_BASE_HEADERS = (
    "_GPS coordinates_latitude",
    "_GPS coordinates_longitude",
    "Forecast grid cell id",
    "Forecast grid cell index",
    "Name",
    "Farm",
    PREDICTION_PREDICTED_COLUMN,
    PREDICTION_PROFILE_KEY_HEADER,
    PREDICTION_PROFILE_LABEL_HEADER,
    PREDICTION_PROFILE_ABOVE_MEAN_HEADER,
)


def _normalize_manual_grid_display_columns(
    prediction_df: pd.DataFrame,
    *,
    manual_bbox_grid: dict[str, object] | None = None,
) -> pd.DataFrame:
    normalized_df = prediction_df.copy()
    if (
        "Forecast grid cell id" not in normalized_df.columns
        and "Forecast locality" in normalized_df.columns
    ):
        normalized_df["Forecast grid cell id"] = normalized_df["Forecast locality"]

    if "Forecast grid cell index" not in normalized_df.columns:
        cell_index_by_id: dict[str, object] = {}
        cells = manual_bbox_grid.get("cells", []) if isinstance(manual_bbox_grid, dict) else []
        for cell in cells if isinstance(cells, list) else []:
            if not isinstance(cell, dict):
                continue
            cell_id = str(cell.get("grid_cell_id") or "").strip()
            if not cell_id:
                continue
            cell_index_by_id[cell_id] = cell.get("grid_cell_index")
        if cell_index_by_id and "Forecast grid cell id" in normalized_df.columns:
            normalized_df["Forecast grid cell index"] = normalized_df["Forecast grid cell id"].map(
                lambda value: cell_index_by_id.get(str(value or "").strip())
            )

    return normalized_df


def _unique_headers(values: list[object]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        header = str(value or "").strip()
        if not header or header in seen:
            continue
        seen.add(header)
        ordered.append(header)
    return ordered


def build_required_prediction_headers(selection_summary: dict[str, object]) -> list[str]:
    initial_settings = selection_summary.get("initial_settings", {}) if isinstance(selection_summary, dict) else {}
    divisions = selection_summary.get("divisions", {}) if isinstance(selection_summary, dict) else {}
    return _unique_headers([
        initial_settings.get("longitude_column", ""),
        initial_settings.get("latitude_column", ""),
        initial_settings.get("planting_date_column", ""),
        initial_settings.get("harvesting_date_column", ""),
        initial_settings.get("soil_texture_column", ""),
        initial_settings.get("soil_depth_column", ""),
        *(
            divisions.get("germplams_identifiers", [])
            if isinstance(divisions.get("germplams_identifiers"), list)
            else []
        ),
        *(
            divisions.get("DG", [])
            if isinstance(divisions.get("DG"), list)
            else []
        ),
        *(
            divisions.get("categorical_data", [])
            if isinstance(divisions.get("categorical_data"), list)
            else []
        ),
        *(
            divisions.get("cuantitative_data", [])
            if isinstance(divisions.get("cuantitative_data"), list)
            else []
        ),
    ])


def build_required_division_headers(selection_summary: dict[str, object]) -> list[str]:
    divisions = selection_summary.get("divisions", {}) if isinstance(selection_summary, dict) else {}
    return _unique_headers([
        *(
            divisions.get("germplams_identifiers", [])
            if isinstance(divisions.get("germplams_identifiers"), list)
            else []
        ),
        *(
            divisions.get("DG", [])
            if isinstance(divisions.get("DG"), list)
            else []
        ),
        *(
            divisions.get("categorical_data", [])
            if isinstance(divisions.get("categorical_data"), list)
            else []
        ),
        *(
            divisions.get("cuantitative_data", [])
            if isinstance(divisions.get("cuantitative_data"), list)
            else []
        ),
    ])


def _read_workbook_headers(workbook_path: Path) -> list[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        return [str(value).strip() for value in header_row if value is not None]
    finally:
        workbook.close()


def _normalize_header_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ")).strip().casefold()


def _resolve_existing_header(headers: list[str], expected_header: str) -> str:
    normalized_expected = _normalize_header_name(expected_header)
    if not normalized_expected:
        return ""
    for header in headers:
        if _normalize_header_name(header) == normalized_expected:
            return str(header or "").strip()
    return ""


def _find_missing_headers(required_headers: list[str], actual_headers: list[str]) -> list[str]:
    actual_lookup = {_normalize_header_name(header) for header in actual_headers}
    return [header for header in required_headers if _normalize_header_name(header) not in actual_lookup]


def _read_saved_selection_summary(metadata: dict[str, object], model_dir: Path) -> dict[str, object]:
    inline_summary = metadata.get("selection_summary")
    if isinstance(inline_summary, dict):
        return inline_summary

    candidate_paths = [
        str(metadata.get("selection_summary_file") or "").strip(),
        str(model_dir / "selection_summary.json"),
    ]
    for candidate in candidate_paths:
        if not candidate:
            continue
        path = Path(candidate)
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    return {}


def _read_saved_normalization_stats(metadata: dict[str, object], model_dir: Path) -> dict[str, object]:
    inline_stats = metadata.get("phase04_normalization_stats")
    if isinstance(inline_stats, dict):
        return inline_stats

    candidate_paths = [
        str(metadata.get("phase04_normalization_stats_file") or "").strip(),
        str(model_dir / "phase04_normalization_stats.json"),
    ]
    for candidate in candidate_paths:
        if not candidate:
            continue
        path = Path(candidate)
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    return {}


def _resolve_saved_model(selection_model_id: str) -> tuple[object, dict[str, object], Path]:
    registered_model = get_registered_model(selection_model_id)
    if registered_model is None:
        raise FileNotFoundError(f"The selected Grain Yield model was not found: {selection_model_id}")
    selection_summary = _read_saved_selection_summary(registered_model.metadata, registered_model.model_dir)
    if not selection_summary:
        raise FileNotFoundError(
            f"The selected Grain Yield model does not include selection_summary.json metadata: {selection_model_id}"
        )
    selection_summary_file = Path(
        str(registered_model.metadata.get("selection_summary_file") or registered_model.model_dir / "selection_summary.json")
    )
    return registered_model, selection_summary, selection_summary_file


def _bbox_centroid(
    latitude_min: float,
    latitude_max: float,
    longitude_min: float,
    longitude_max: float,
) -> tuple[float, float]:
    return ((latitude_min + latitude_max) / 2.0, (longitude_min + longitude_max) / 2.0)


def _filter_records_by_selected_names(
    records: list[dict[str, str]],
    selected_names: list[str] | None,
    *,
    selection_mode: str = "all_matches",
) -> list[dict[str, str]]:
    if not selected_names:
        return records
    ordered_names = [str(name or "").strip() for name in selected_names if str(name or "").strip()]
    normalized = set(ordered_names)
    if not normalized:
        return records
    matching_records = [record for record in records if str(record.get("Name", "")).strip() in normalized]
    if selection_mode != "first_match_only":
        return matching_records
    first_matches: list[dict[str, str]] = []
    seen_names: set[str] = set()
    ordered_lookup = ordered_names + sorted(normalized)
    for selected_name in ordered_lookup:
        if selected_name in seen_names:
            continue
        for record in records:
            if str(record.get("Name", "")).strip() == selected_name:
                first_matches.append(record)
                seen_names.add(selected_name)
                break
    return first_matches


def _annotate_prediction_internal_row_ids(
    records: list[dict[str, str]],
) -> list[dict[str, str]]:
    annotated: list[dict[str, str]] = []
    for row_number, record in enumerate(records, start=2):
        updated = dict(record)
        updated[PREDICTION_INTERNAL_ROW_ID_HEADER] = _build_prediction_internal_row_id(row_number)
        annotated.append(updated)
    return annotated


def _filter_records_by_selected_row_ids(
    records: list[dict[str, str]],
    selected_row_ids: list[str] | None,
) -> list[dict[str, str]]:
    normalized_ids = {
        str(value or "").strip()
        for value in (selected_row_ids or [])
        if str(value or "").strip()
    }
    if not normalized_ids:
        return records
    return [
        record
        for record in records
        if str(record.get(PREDICTION_INTERNAL_ROW_ID_HEADER, "") or "").strip() in normalized_ids
    ]


def _is_missing_selected_id_value(value: object) -> bool:
    normalized = str(value or "").strip()
    return normalized.casefold() in {"", "null", "none", "nan"}


def _apply_selected_id_fallbacks(
    records: list[dict[str, str]],
    selected_id_header: str,
) -> list[dict[str, str]]:
    normalized_header = str(selected_id_header or "").strip()
    if not normalized_header:
        return records

    grouped_records: dict[str, list[dict[str, str]]] = {}
    for record in records:
        name = str(record.get("Name", "")).strip()
        grouped_records.setdefault(name, []).append(dict(record))

    output_records: list[dict[str, str]] = []
    for _, name_records in grouped_records.items():
        if len(name_records) <= 1:
            output_records.extend(name_records)
            continue

        used_values: set[str] = set()
        for record in name_records:
            raw_value = str(record.get(normalized_header, "") or "").strip()
            if not _is_missing_selected_id_value(raw_value):
                used_values.add(raw_value)

        next_fallback = 1
        for record in name_records:
            raw_value = str(record.get(normalized_header, "") or "").strip()
            if _is_missing_selected_id_value(raw_value):
                while str(next_fallback) in used_values:
                    next_fallback += 1
                raw_value = str(next_fallback)
                next_fallback += 1
                record[normalized_header] = raw_value
                used_values.add(raw_value)
        output_records.extend(name_records)
    return output_records


def _parse_float(value: object) -> float | None:
    text = str(value or "").strip().replace(",", ".")
    if not text:
        return None
    try:
        numeric = float(text)
    except ValueError:
        return None
    return numeric if numeric == numeric else None


def _build_profile_candidate_headers(
    selection_summary: dict[str, object],
    headers: list[str],
    *,
    selected_id_header: str = "",
) -> list[str]:
    initial_settings = selection_summary.get("initial_settings", {}) if isinstance(selection_summary, dict) else {}
    divisions = selection_summary.get("divisions", {}) if isinstance(selection_summary, dict) else {}
    excluded_headers = {
        str(initial_settings.get("target_column", "")).strip(),
        str(initial_settings.get("longitude_column", "")).strip(),
        str(initial_settings.get("latitude_column", "")).strip(),
        str(initial_settings.get("planting_date_column", "")).strip(),
        str(initial_settings.get("harvesting_date_column", "")).strip(),
        str(initial_settings.get("soil_texture_column", "")).strip(),
        str(initial_settings.get("soil_depth_column", "")).strip(),
        "Name",
        str(selected_id_header or "").strip(),
        PREDICTION_INTERNAL_ROW_ID_HEADER,
        PREDICTION_GROUP_KEY_HEADER,
        PREDICTION_GROUP_LABEL_HEADER,
        PREDICTION_GROUP_REP_HEADER,
    }
    candidate_headers = _unique_headers([
        *(divisions.get("germplams_identifiers", []) if isinstance(divisions.get("germplams_identifiers"), list) else []),
        *(divisions.get("categorical_data", []) if isinstance(divisions.get("categorical_data"), list) else []),
        *(divisions.get("cuantitative_data", []) if isinstance(divisions.get("cuantitative_data"), list) else []),
    ])
    filtered = [
        header
        for header in candidate_headers
        if header in headers and header and header not in excluded_headers and not header.startswith("DG")
    ]
    if filtered:
        return filtered
    return [
        header
        for header in headers
        if header and header not in excluded_headers and not header.startswith("DG")
    ]


def _build_germplasm_profile_metadata(
    records: list[dict[str, str]],
    headers: list[str],
    selection_summary: dict[str, object],
    *,
    selected_id_header: str = "",
) -> tuple[list[dict[str, str]], dict[str, object]]:
    if not records:
        return [], {
            "multi_profile_mode": False,
            "profile_count": 0,
            "varying_columns": [],
            "target_mean_threshold": None,
            "selected_name": "",
            "profile_summaries": [],
        }

    initial_settings = selection_summary.get("initial_settings", {}) if isinstance(selection_summary, dict) else {}
    target_column = str(initial_settings.get("target_column", "")).strip()
    normalized_selected_id_header = str(selected_id_header or "").strip()
    candidate_headers = _build_profile_candidate_headers(
        selection_summary,
        headers,
        selected_id_header=normalized_selected_id_header,
    )

    varying_columns: list[dict[str, object]] = []
    varying_header_names: list[str] = []
    for header in candidate_headers:
        distinct_values = []
        seen: set[str] = set()
        for record in records:
            value = str(record.get(header, "") or "").strip()
            if not value:
                continue
            key = value.casefold()
            if key in seen:
                continue
            seen.add(key)
            distinct_values.append(value)
        if len(distinct_values) <= 1:
            continue
        varying_header_names.append(header)
        varying_columns.append({
            "header": header,
            "values": distinct_values[:8],
            "unique_count": len(distinct_values),
        })

    indexed_records = list(enumerate(records))
    profile_groups: dict[str, dict[str, object]] = {}
    if not normalized_selected_id_header:
        fallback_name = str(records[0].get("Name", "") or "").strip() if records else ""
        profile_label = fallback_name or "selected-germplasm"
        profile_groups[profile_label] = {
            "key": profile_label,
            "label": profile_label,
            "rows": indexed_records,
        }
    else:
        for index, record in indexed_records:
            selected_name = str(record.get("Name", "") or "").strip()
            selected_id_value = (
                str(record.get(normalized_selected_id_header, "") or "").strip()
                if normalized_selected_id_header
                else ""
            )
            profile_label = (
                f"{selected_name} - {selected_id_value}"
                if selected_name and selected_id_value
                else selected_name or selected_id_value or f"profile-{index + 1}"
            )
            profile_key = profile_label
            if profile_key not in profile_groups:
                profile_groups[profile_key] = {
                    "key": profile_key,
                    "label": profile_label,
                    "rows": [],
                }
            profile_groups[profile_key]["rows"].append((index, record))

    target_values = [_parse_float(record.get(target_column, "")) for record in records] if target_column else []
    target_values = [value for value in target_values if value is not None]
    target_mean_threshold = round(sum(target_values) / len(target_values), 6) if target_values else None

    profile_summaries: list[dict[str, object]] = []
    annotated_records_by_index: dict[int, dict[str, str]] = {}
    for palette_index, profile in enumerate(profile_groups.values()):
        profile_rows = profile.get("rows", [])
        profile_target_values = [
            _parse_float(record.get(target_column, ""))
            for _, record in profile_rows
            if target_column
        ]
        profile_target_values = [value for value in profile_target_values if value is not None]
        profile_target_mean = (
            round(sum(profile_target_values) / len(profile_target_values), 6)
            if profile_target_values
            else None
        )
        profile_summary = {
            "key": str(profile.get("key") or ""),
            "label": str(profile.get("label") or ""),
            "row_count": len(profile_rows),
            "rep_count": len(profile_rows),
            "palette_index": palette_index,
            "target_mean": profile_target_mean,
        }
        profile_summaries.append(profile_summary)
        for row_index, record in profile_rows:
            updated = dict(record)
            updated[PREDICTION_PROFILE_KEY_HEADER] = profile_summary["key"]
            updated[PREDICTION_PROFILE_LABEL_HEADER] = profile_summary["label"]
            updated[PREDICTION_PROFILE_COLUMNS_HEADER] = ", ".join(varying_header_names)
            updated[PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER] = "" if profile_target_mean is None else str(profile_target_mean)
            updated[PREDICTION_PROFILE_TARGET_MEAN_HEADER] = "" if target_mean_threshold is None else str(target_mean_threshold)
            updated[PREDICTION_GROUP_KEY_HEADER] = profile_summary["key"]
            updated[PREDICTION_GROUP_LABEL_HEADER] = profile_summary["label"]
            updated[PREDICTION_GROUP_REP_HEADER] = str(len(profile_rows))
            annotated_records_by_index[row_index] = updated

    ordered_records = [annotated_records_by_index[index] for index, _ in indexed_records]

    return ordered_records, {
        "multi_profile_mode": len(profile_groups) > 1 if normalized_selected_id_header else False,
        "profile_count": len(profile_groups) if normalized_selected_id_header else 1,
        "varying_columns": varying_columns,
        "target_mean_threshold": target_mean_threshold,
        "selected_name": "",
        "selected_id_header": normalized_selected_id_header,
        "profile_summaries": profile_summaries,
        "source_row_count": len(records),
        "uses_source_row_dates": True,
    }


def build_prediction_phase02_workbook(
    uploaded_workbook: Path,
    *,
    selection_summary: dict[str, object],
    output_file: Path,
    forecast_planting_date: str,
    forecast_harvesting_date: str,
    regional_bounds_label: str,
    regional_bounds_latitude_min: float,
    regional_bounds_latitude_max: float,
    regional_bounds_longitude_min: float,
    regional_bounds_longitude_max: float,
    selected_germplasm_names: list[str] | None = None,
    selected_germplasm_row_ids: list[str] | None = None,
    selected_germplasm_selection_mode: str = "all_matches",
    selected_id_header: str = "",
) -> dict[str, object]:
    definition = selection_summary
    ordered_headers = build_column_order(definition)
    headers, records = read_sheet_headers_and_rows(uploaded_workbook)
    if not headers:
        raise ValueError("The uploaded prediction workbook is empty.")
    resolved_farm_header = _resolve_existing_header(headers, "Farm")
    if not resolved_farm_header:
        raise ValueError("The uploaded workbook must contain the required Farm column.")
    requested_clear_selection = str(selected_id_header or "").strip() == CLEAR_SELECTION_ID_SENTINEL
    normalized_selected_id_header = (
        ""
        if requested_clear_selection
        else _resolve_existing_header(headers, str(selected_id_header or "").strip())
    )
    if normalized_selected_id_header and normalized_selected_id_header in headers and normalized_selected_id_header not in ordered_headers:
        ordered_headers.append(normalized_selected_id_header)
    annotated_records = _annotate_prediction_internal_row_ids(records)
    if PREDICTION_INTERNAL_ROW_ID_HEADER not in ordered_headers:
        ordered_headers.append(PREDICTION_INTERNAL_ROW_ID_HEADER)
    for hidden_header in (PREDICTION_GROUP_KEY_HEADER, PREDICTION_GROUP_LABEL_HEADER, PREDICTION_GROUP_REP_HEADER):
        if hidden_header not in ordered_headers:
            ordered_headers.append(hidden_header)

    required_division_headers = build_required_division_headers(selection_summary)
    missing_headers = _find_missing_headers(required_division_headers, headers)
    if missing_headers:
        raise ValueError(
            "The uploaded workbook does not match the saved-model stepper definition. Missing required columns:\n- "
            + "\n- ".join(missing_headers)
        )

    filtered_records = _filter_records_by_selected_names(
        annotated_records,
        selected_germplasm_names,
        selection_mode=selected_germplasm_selection_mode,
    )
    filtered_records = _filter_records_by_selected_row_ids(
        filtered_records,
        selected_germplasm_row_ids,
    )
    if not filtered_records:
        raise ValueError("The selected germplasm filter produced no records for prediction.")
    filtered_records = _apply_selected_id_fallbacks(filtered_records, normalized_selected_id_header)

    profiled_records, profile_metadata = _build_germplasm_profile_metadata(
        filtered_records,
        headers,
        selection_summary,
        selected_id_header=normalized_selected_id_header,
    )

    initial_settings = selection_summary.get("initial_settings", {}) if isinstance(selection_summary, dict) else {}
    target_column = str(initial_settings.get("target_column", "")).strip()
    longitude_column = str(initial_settings.get("longitude_column", "")).strip()
    latitude_column = str(initial_settings.get("latitude_column", "")).strip()
    planting_date_column = str(initial_settings.get("planting_date_column", "")).strip()
    harvesting_date_column = str(initial_settings.get("harvesting_date_column", "")).strip()
    soil_texture_column = str(initial_settings.get("soil_texture_column", "")).strip()
    soil_depth_column = str(initial_settings.get("soil_depth_column", "")).strip()
    centroid_latitude, centroid_longitude = _bbox_centroid(
        regional_bounds_latitude_min,
        regional_bounds_latitude_max,
        regional_bounds_longitude_min,
        regional_bounds_longitude_max,
    )

    output_records: list[dict[str, str]] = []
    for record in profiled_records:
        output_record = {header: str(record.get(header, "") or "").strip() for header in ordered_headers}
        if PREDICTION_PROFILE_KEY_HEADER not in output_record:
            output_record[PREDICTION_PROFILE_KEY_HEADER] = str(record.get(PREDICTION_PROFILE_KEY_HEADER, "") or "").strip()
        if PREDICTION_PROFILE_LABEL_HEADER not in output_record:
            output_record[PREDICTION_PROFILE_LABEL_HEADER] = str(record.get(PREDICTION_PROFILE_LABEL_HEADER, "") or "").strip()
        if PREDICTION_PROFILE_COLUMNS_HEADER not in output_record:
            output_record[PREDICTION_PROFILE_COLUMNS_HEADER] = str(record.get(PREDICTION_PROFILE_COLUMNS_HEADER, "") or "").strip()
        if PREDICTION_PROFILE_TARGET_MEAN_HEADER not in output_record:
            output_record[PREDICTION_PROFILE_TARGET_MEAN_HEADER] = str(record.get(PREDICTION_PROFILE_TARGET_MEAN_HEADER, "") or "").strip()
        if PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER not in output_record:
            output_record[PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER] = str(record.get(PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER, "") or "").strip()
        for hidden_header in (PREDICTION_INTERNAL_ROW_ID_HEADER, PREDICTION_GROUP_KEY_HEADER, PREDICTION_GROUP_LABEL_HEADER, PREDICTION_GROUP_REP_HEADER):
            if hidden_header not in output_record:
                output_record[hidden_header] = str(record.get(hidden_header, "") or "").strip()
        if target_column:
            output_record[target_column] = ""
        if longitude_column:
            output_record[longitude_column] = str(centroid_longitude)
        if latitude_column:
            output_record[latitude_column] = str(centroid_latitude)
        if planting_date_column:
            output_record[planting_date_column] = str(record.get(planting_date_column, "") or "").strip()
        if harvesting_date_column:
            output_record[harvesting_date_column] = str(record.get(harvesting_date_column, "") or "").strip()
        if soil_texture_column and soil_texture_column not in output_record:
            output_record[soil_texture_column] = ""
        if soil_depth_column and soil_depth_column not in output_record:
            output_record[soil_depth_column] = ""
        output_records.append(output_record)

    output_headers = _unique_headers([
        *ordered_headers,
        PREDICTION_PROFILE_KEY_HEADER,
        PREDICTION_PROFILE_LABEL_HEADER,
        PREDICTION_PROFILE_COLUMNS_HEADER,
        PREDICTION_PROFILE_TARGET_MEAN_HEADER,
        PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER,
        PREDICTION_INTERNAL_ROW_ID_HEADER,
        PREDICTION_GROUP_KEY_HEADER,
        PREDICTION_GROUP_LABEL_HEADER,
        PREDICTION_GROUP_REP_HEADER,
    ])
    write_xlsx(output_file, output_headers, output_records)
    return {
        "phase": "phase02prediction",
        "source_workbook": str(uploaded_workbook),
        "output_workbook": str(output_file),
        "required_division_headers": required_division_headers,
        "selected_germplasm_count": len(selected_germplasm_names or []),
        "selected_germplasm_row_id_count": len(selected_germplasm_row_ids or []),
        "selected_germplasm_selection_mode": selected_germplasm_selection_mode,
        "output_row_count": len(output_records),
        "regional_bounds_label": regional_bounds_label,
        "regional_bounds_latitude_min": regional_bounds_latitude_min,
        "regional_bounds_latitude_max": regional_bounds_latitude_max,
        "regional_bounds_longitude_min": regional_bounds_longitude_min,
        "regional_bounds_longitude_max": regional_bounds_longitude_max,
        "centroid_latitude": centroid_latitude,
        "centroid_longitude": centroid_longitude,
        "forecast_planting_date": forecast_planting_date,
        "forecast_harvesting_date": forecast_harvesting_date,
        "prediction_profile_analysis": profile_metadata,
        "uses_source_row_dates": True,
    }


def _filter_prediction_output_files_by_profile_keys(
    prediction_xlsx: Path,
    prediction_csv: Path,
    allowed_profile_keys: set[str],
) -> None:
    if allowed_profile_keys and prediction_xlsx.exists():
        prediction_df = pd.read_excel(prediction_xlsx)
        if PREDICTION_PROFILE_KEY_HEADER in prediction_df.columns:
            filtered_df = prediction_df[
                prediction_df[PREDICTION_PROFILE_KEY_HEADER]
                .astype(str)
                .str.strip()
                .isin(allowed_profile_keys)
            ].copy()
            filtered_df.to_excel(prediction_xlsx, index=False)
    if allowed_profile_keys and prediction_csv.exists():
        with prediction_csv.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            fieldnames = list(reader.fieldnames or [])
            rows = list(reader)
        if PREDICTION_PROFILE_KEY_HEADER in fieldnames:
            filtered_rows = [
                row
                for row in rows
                if str(row.get(PREDICTION_PROFILE_KEY_HEADER, "") or "").strip() in allowed_profile_keys
            ]
            with prediction_csv.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(filtered_rows)


def _postprocess_multi_profile_prediction_results(
    *,
    geojson: dict[str, object],
    geojson_file: Path,
    prediction_xlsx: Path,
    prediction_csv: Path,
    profile_metadata: dict[str, object],
) -> tuple[dict[str, object], dict[str, object]]:
    if not profile_metadata.get("multi_profile_mode"):
        return geojson, {
            "multi_profile_mode": False,
        }

    target_mean_threshold = _parse_float(profile_metadata.get("target_mean_threshold"))
    features = geojson.get("features", []) if isinstance(geojson, dict) else []
    profile_stats: dict[str, dict[str, object]] = {}
    for feature in features:
        if not isinstance(feature, dict):
            continue
        properties = feature.get("properties", {}) if isinstance(feature.get("properties"), dict) else {}
        profile_key = str(properties.get(PREDICTION_PROFILE_KEY_HEADER, "") or "").strip()
        if not profile_key:
            continue
        predicted_value = _parse_float(properties.get(PREDICTION_PREDICTED_COLUMN))
        bucket = profile_stats.setdefault(profile_key, {"values": [], "label": str(properties.get(PREDICTION_PROFILE_LABEL_HEADER, "") or "").strip()})
        if predicted_value is not None:
            bucket["values"].append(predicted_value)

    profile_summary_updates: list[dict[str, object]] = []
    allowed_profile_keys: set[str] = set()
    for profile_summary in profile_metadata.get("profile_summaries", []) if isinstance(profile_metadata.get("profile_summaries"), list) else []:
        profile_key = str((profile_summary or {}).get("key") or "").strip()
        stats = profile_stats.get(profile_key, {})
        predicted_values = stats.get("values", []) if isinstance(stats, dict) else []
        predicted_mean = round(sum(predicted_values) / len(predicted_values), 6) if predicted_values else None
        is_above_mean = bool(
            target_mean_threshold is not None
            and predicted_mean is not None
            and predicted_mean > target_mean_threshold
        )
        profile_summary_updates.append({
            **profile_summary,
            "predicted_mean": predicted_mean,
            "above_target_mean": is_above_mean,
        })
        if is_above_mean:
            allowed_profile_keys.add(profile_key)

    if not allowed_profile_keys and profile_summary_updates:
        ranked_summaries = [item for item in profile_summary_updates if item.get("predicted_mean") is not None]
        if ranked_summaries:
            fallback = max(ranked_summaries, key=lambda item: float(item.get("predicted_mean") or 0))
            allowed_profile_keys.add(str(fallback.get("key") or ""))

    annotated_features = []
    for feature in features:
        if not isinstance(feature, dict):
            continue
        properties = feature.get("properties", {}) if isinstance(feature.get("properties"), dict) else {}
        profile_key = str(properties.get(PREDICTION_PROFILE_KEY_HEADER, "") or "").strip()
        updated_feature = dict(feature)
        updated_properties = dict(properties)
        updated_properties[PREDICTION_PROFILE_ABOVE_MEAN_HEADER] = (
            "Yes" if profile_key in allowed_profile_keys else "No"
        )
        updated_feature["properties"] = updated_properties
        annotated_features.append(updated_feature)

    filtered_geojson = dict(geojson)
    filtered_metadata = dict(filtered_geojson.get("metadata", {}) if isinstance(filtered_geojson.get("metadata"), dict) else {})
    filtered_metadata["total_features"] = len(annotated_features)
    filtered_geojson["metadata"] = filtered_metadata
    filtered_geojson["features"] = annotated_features
    geojson_file.write_text(json.dumps(filtered_geojson, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    return filtered_geojson, {
        "multi_profile_mode": True,
        "target_mean_threshold": target_mean_threshold,
        "profile_summaries": profile_summary_updates,
        "displayed_profile_keys": sorted(key for key in allowed_profile_keys if key),
        "displayed_profile_count": len([key for key in allowed_profile_keys if key]),
        "profile_color_mode": "categorical",
        "uses_source_row_dates": True,
        "varying_columns": profile_metadata.get("varying_columns", []),
        "selected_name": profile_metadata.get("selected_name", ""),
    }


def _first_non_blank_value(series: pd.Series) -> object:
    for value in series:
        if str(value or "").strip():
            return value
    return series.iloc[0] if len(series.index) else ""


def _should_preserve_manual_grid_prediction_rows(
    *,
    climate_scope: str,
    selected_id_header: str,
    selected_germplasm_names: list[str] | None,
) -> bool:
    if str(climate_scope or "").strip() != "regional_manual":
        return False
    if str(selected_id_header or "").strip():
        return False
    distinct_names = {
        str(name or "").strip()
        for name in (selected_germplasm_names or [])
        if str(name or "").strip()
    }
    return len(distinct_names) == 1


def _aggregate_prediction_dataframe_by_selected_id(
    prediction_df: pd.DataFrame,
    *,
    selected_id_header: str,
) -> pd.DataFrame:
    normalized_selected_id_header = str(selected_id_header or "").strip()
    required_columns = {"Name", PREDICTION_PREDICTED_COLUMN}
    if normalized_selected_id_header:
        required_columns.add(normalized_selected_id_header)
    if not required_columns.issubset(set(prediction_df.columns)):
        return prediction_df

    grouping_columns = [
        column
        for column in [
            "Country",
            "_GPS coordinates_latitude",
            "_GPS coordinates_longitude",
            "Forecast grid cell id",
            "Forecast grid cell index",
            "Name",
            PREDICTION_PROFILE_KEY_HEADER,
            PREDICTION_PROFILE_LABEL_HEADER,
            PREDICTION_PROFILE_COLUMNS_HEADER,
            PREDICTION_PROFILE_TARGET_MEAN_HEADER,
            PREDICTION_PROFILE_GROUP_TARGET_MEAN_HEADER,
            PREDICTION_PROFILE_ABOVE_MEAN_HEADER,
            PREDICTION_GROUP_KEY_HEADER,
            PREDICTION_GROUP_LABEL_HEADER,
        ]
        if column in prediction_df.columns
    ]
    if normalized_selected_id_header and normalized_selected_id_header in prediction_df.columns:
        grouping_columns.append(normalized_selected_id_header)
    if "Name" not in grouping_columns:
        return prediction_df

    aggregated_rows: list[dict[str, object]] = []
    for _, group_df in prediction_df.groupby(grouping_columns, dropna=False, sort=False):
        aggregated_row: dict[str, object] = {}
        for column in prediction_df.columns:
            if column == PREDICTION_INTERNAL_ROW_ID_HEADER:
                continue
            if column == PREDICTION_GROUP_REP_HEADER:
                aggregated_row[column] = len(group_df.index)
                continue
            if column == PREDICTION_PREDICTED_COLUMN:
                numeric_series = pd.to_numeric(group_df[column], errors="coerce")
                aggregated_row[column] = round(float(numeric_series.mean()), 6) if numeric_series.notna().any() else _first_non_blank_value(group_df[column])
                continue
            if column in grouping_columns:
                aggregated_row[column] = _first_non_blank_value(group_df[column])
                continue
            aggregated_row[column] = _first_non_blank_value(group_df[column])
        aggregated_rows.append(aggregated_row)
    return pd.DataFrame(aggregated_rows, columns=[column for column in prediction_df.columns if column != PREDICTION_INTERNAL_ROW_ID_HEADER])


def _prepare_prediction_output_dataframe_for_display(
    prediction_df: pd.DataFrame,
    *,
    latitude_column: str,
    longitude_column: str,
    selected_id_header: str,
) -> pd.DataFrame:
    prepared_df = prediction_df.copy()
    report_only_drop_columns = [
        PREDICTION_PROFILE_KEY_HEADER,
        PREDICTION_PROFILE_LABEL_HEADER,
        PREDICTION_PROFILE_COLUMNS_HEADER,
        PREDICTION_PROFILE_TARGET_MEAN_HEADER,
        PREDICTION_GROUP_KEY_HEADER,
        PREDICTION_GROUP_LABEL_HEADER,
    ]
    existing_drop_columns = [column for column in report_only_drop_columns if column in prepared_df.columns]
    if existing_drop_columns:
        prepared_df = prepared_df.drop(columns=existing_drop_columns)
    if "Country" in prepared_df.columns:
        prepared_df = prepared_df.drop(columns=["Country"])
    if selected_id_header and selected_id_header in prepared_df.columns:
        reordered_columns = [selected_id_header] + [column for column in prepared_df.columns if column != selected_id_header]
        prepared_df = prepared_df.loc[:, reordered_columns]

    sort_columns = [
        column
        for column in (longitude_column, latitude_column, PREDICTION_PREDICTED_COLUMN)
        if str(column or "").strip()
    ]
    existing_sort_columns = [column for column in sort_columns if column in prepared_df.columns]
    if existing_sort_columns:
        existing_ascending = [True, True, False][:len(existing_sort_columns)]
        prepared_df = prepared_df.sort_values(
            by=existing_sort_columns,
            ascending=existing_ascending,
            kind="mergesort",
        )
    return prepared_df


def _sort_prediction_output_files(
    *,
    prediction_xlsx: Path,
    prediction_csv: Path,
    latitude_column: str,
    longitude_column: str,
    selected_id_header: str,
) -> None:
    if prediction_xlsx.exists():
        prediction_df = pd.read_excel(prediction_xlsx)
        prediction_df = _prepare_prediction_output_dataframe_for_display(
            prediction_df,
            latitude_column=latitude_column,
            longitude_column=longitude_column,
            selected_id_header=selected_id_header,
        )
        prediction_df.to_excel(prediction_xlsx, index=False)

    if prediction_csv.exists():
        prediction_df = pd.read_csv(prediction_csv)
        prediction_df = _prepare_prediction_output_dataframe_for_display(
            prediction_df,
            latitude_column=latitude_column,
            longitude_column=longitude_column,
            selected_id_header=selected_id_header,
        )
        prediction_df.to_csv(prediction_csv, index=False)


def _should_use_manual_grid_dataframe_finalization(
    *,
    climate_scope: str,
    selected_id_header: str,
) -> bool:
    return (
        str(climate_scope or "").strip() == "regional_manual"
        and bool(str(selected_id_header or "").strip())
    )


def _build_lightweight_manual_grid_geojson_dataframe(
    prediction_df: pd.DataFrame,
    *,
    selected_id_header: str,
) -> pd.DataFrame:
    selected_columns: list[str] = []
    seen_columns: set[str] = set()
    requested_columns = [
        *MANUAL_GRID_GEOJSON_BASE_HEADERS,
        str(selected_id_header or "").strip(),
    ]
    for column in requested_columns:
        normalized_column = str(column or "").strip()
        if not normalized_column or normalized_column in seen_columns or normalized_column not in prediction_df.columns:
            continue
        seen_columns.add(normalized_column)
        selected_columns.append(normalized_column)
    if not selected_columns:
        return prediction_df
    return prediction_df.loc[:, selected_columns].copy()


def _aggregate_manual_grid_prediction_dataframe_for_display(
    prediction_df: pd.DataFrame,
) -> pd.DataFrame:
    required_columns = {
        "Forecast grid cell id",
        "_GPS coordinates_latitude",
        "_GPS coordinates_longitude",
        PREDICTION_PREDICTED_COLUMN,
    }
    if not required_columns.issubset(set(prediction_df.columns)):
        return prediction_df

    aggregated_rows: list[dict[str, object]] = []
    for grid_cell_id, group_df in prediction_df.groupby("Forecast grid cell id", dropna=False, sort=False):
        if pd.isna(grid_cell_id) or not str(grid_cell_id).strip():
            continue
        numeric_series = pd.to_numeric(group_df[PREDICTION_PREDICTED_COLUMN], errors="coerce")
        best_index = numeric_series.idxmax() if numeric_series.notna().any() else group_df.index[0]
        representative = group_df.loc[best_index].copy()
        representative["Prediction cell row count"] = int(len(group_df.index))
        representative["Prediction cell replication count"] = int(len(group_df.index))
        aggregated_rows.append(dict(representative))

    if not aggregated_rows:
        return prediction_df
    return pd.DataFrame(aggregated_rows, columns=list(prediction_df.columns) + [
        column for column in ("Prediction cell row count", "Prediction cell replication count")
        if column not in prediction_df.columns
    ])


def _restore_selected_id_column_in_phase04_workbook(
    *,
    phase03_workbook: Path,
    phase04_workbook: Path,
    selected_id_header: str,
) -> None:
    normalized_header = str(selected_id_header or "").strip()
    if not normalized_header or not phase03_workbook.exists() or not phase04_workbook.exists():
        return

    phase03_df = pd.read_excel(phase03_workbook, dtype=object)
    phase04_df = pd.read_excel(phase04_workbook, dtype=object)
    phase03_df.columns = [str(column).strip() for column in phase03_df.columns]
    phase04_df.columns = [str(column).strip() for column in phase04_df.columns]
    if normalized_header not in phase03_df.columns or normalized_header in phase04_df.columns:
        return

    if "CE Row ID" in phase03_df.columns and "CE Row ID" in phase04_df.columns:
        source_values = (
            phase03_df.loc[:, ["CE Row ID", normalized_header]]
            .copy()
        )
        source_values["CE Row ID"] = source_values["CE Row ID"].astype(str).str.strip()
        source_values = source_values.drop_duplicates(subset=["CE Row ID"], keep="first")
        merged_df = phase04_df.merge(source_values, on="CE Row ID", how="left")
        phase04_df = merged_df
    elif len(phase03_df.index) == len(phase04_df.index):
        phase04_df[normalized_header] = phase03_df[normalized_header].tolist()
    else:
        return

    phase04_df.to_excel(phase04_workbook, index=False)


def update_saved_model_prediction_metadata(
    *,
    model_dir: Path,
    metadata: dict[str, object],
    selection_summary_file: Path | None,
    phase04_selected_fields_csv: Path | None,
    phase04_workbook: Path | None,
    phase04_training_input_csv: Path | None,
    phase04_normalization_stats_file: Path | None = None,
) -> dict[str, object]:
    updated_metadata = dict(metadata)
    selection_summary_payload: dict[str, object] = {}

    if selection_summary_file is not None and selection_summary_file.exists():
        model_selection_summary = model_dir / "selection_summary.json"
        shutil.copy2(selection_summary_file, model_selection_summary)
        selection_summary_payload = json.loads(selection_summary_file.read_text(encoding="utf-8"))
        updated_metadata["selection_summary_file"] = str(model_selection_summary)
        updated_metadata["selection_summary"] = selection_summary_payload
        updated_metadata["required_prediction_headers"] = build_required_prediction_headers(selection_summary_payload)
        updated_metadata["required_division_headers"] = build_required_division_headers(selection_summary_payload)

    if phase04_selected_fields_csv is not None and phase04_selected_fields_csv.exists():
        saved_selected_fields_csv = model_dir / "phase04_selected_fields.csv"
        shutil.copy2(phase04_selected_fields_csv, saved_selected_fields_csv)
        updated_metadata["phase04_selected_fields_csv"] = str(saved_selected_fields_csv)

    if phase04_workbook is not None and phase04_workbook.exists():
        saved_phase04_workbook = model_dir / "phase04.xlsx"
        shutil.copy2(phase04_workbook, saved_phase04_workbook)
        updated_metadata["phase04_workbook"] = str(saved_phase04_workbook)

    if phase04_training_input_csv is not None and phase04_training_input_csv.exists():
        saved_phase04_training_csv = model_dir / "phase04_training_input.csv"
        shutil.copy2(phase04_training_input_csv, saved_phase04_training_csv)
        updated_metadata["phase04_training_input_csv"] = str(saved_phase04_training_csv)

    if phase04_normalization_stats_file is not None and phase04_normalization_stats_file.exists():
        saved_normalization_stats = model_dir / "phase04_normalization_stats.json"
        shutil.copy2(phase04_normalization_stats_file, saved_normalization_stats)
        updated_metadata["phase04_normalization_stats_file"] = str(saved_normalization_stats)
        updated_metadata["phase04_normalization_stats"] = json.loads(
            phase04_normalization_stats_file.read_text(encoding="utf-8")
        )

    write_model_metadata(model_dir, updated_metadata)
    return updated_metadata


def validate_saved_model_prediction_workbook(
    uploaded_workbook: Path,
    *,
    selected_model_id: str,
    climate_scope: str = "point",
    selected_germplasm_projection_mode: str = "",
    projection_template_workbook: Path | None = None,
) -> None:
    registered_model, selection_summary, _ = _resolve_saved_model(selected_model_id)

    if selected_germplasm_projection_mode == "all_markers" and projection_template_workbook is not None:
        return

    if climate_scope == "point":
        workbook_mode, workbook_errors = resolve_point_uploaded_workbook_mode(uploaded_workbook)
        if workbook_mode == "coordinate" and not workbook_errors:
            return

    required_headers = build_required_division_headers(selection_summary)
    if not required_headers:
        return

    actual_headers = _read_workbook_headers(uploaded_workbook)
    missing_headers = _find_missing_headers(required_headers, actual_headers)
    if missing_headers:
        raise ValueError(
            "The uploaded workbook does not match the saved-model stepper definition. Missing required columns:\n- "
            + "\n- ".join(missing_headers)
        )

    if _normalize_header_name("Name") not in {_normalize_header_name(header) for header in actual_headers}:
        raise ValueError("The uploaded workbook does not contain the required Name column.")


def validate_saved_model_prediction_workbook_for_names(
    uploaded_workbook: Path,
    *,
    selected_model_id: str,
    climate_scope: str = "regional_manual",
) -> None:
    validate_saved_model_prediction_workbook(
        uploaded_workbook,
        selected_model_id=selected_model_id,
        climate_scope=climate_scope,
    )


def run_saved_model_prediction(
    uploaded_workbook: Path,
    run_dir: Path,
    *,
    progress_callback=None,
    selected_model_id: str,
    forecast_planting_date: str | None = None,
    forecast_harvesting_date: str | None = None,
    climate_scope: str = "point",
    regional_country: str | None = None,
    selected_germplasm_names: list[str] | None = None,
    selected_germplasm_projection_mode: str | None = None,
    projection_template_workbook: Path | None = None,
    regional_bounds_label: str | None = None,
    regional_bounds_latitude_min: float | None = None,
    regional_bounds_latitude_max: float | None = None,
    regional_bounds_longitude_min: float | None = None,
    regional_bounds_longitude_max: float | None = None,
    nasa_grid_resolution_km: int = 10,
) -> SavedModelPredictionOutputs:
    validate_saved_model_prediction_workbook(
        uploaded_workbook,
        selected_model_id=selected_model_id,
        climate_scope=climate_scope,
        selected_germplasm_projection_mode=selected_germplasm_projection_mode or "",
        projection_template_workbook=projection_template_workbook,
    )
    preprocess = run_preprocess_pipeline(
        uploaded_workbook,
        run_dir / "preprocess",
        progress_callback=progress_callback,
        forecast_planting_date=forecast_planting_date,
        forecast_harvesting_date=forecast_harvesting_date,
        climate_scope=climate_scope,
        regional_country=regional_country,
        selected_model_id=selected_model_id,
        selected_germplasm_names=selected_germplasm_names,
        selected_germplasm_projection_mode=selected_germplasm_projection_mode,
        projection_template_workbook=projection_template_workbook,
        regional_bounds_label=regional_bounds_label,
        regional_bounds_latitude_min=regional_bounds_latitude_min,
        regional_bounds_latitude_max=regional_bounds_latitude_max,
        regional_bounds_longitude_min=regional_bounds_longitude_min,
        regional_bounds_longitude_max=regional_bounds_longitude_max,
        nasa_grid_resolution_km=nasa_grid_resolution_km,
    )
    model = run_model_pipeline(
        preprocess.phase06_xlsx,
        run_dir,
        progress_callback=progress_callback,
        selected_model_id=selected_model_id,
    )
    return SavedModelPredictionOutputs(preprocess=preprocess, model=model)


def run_ce_saved_model_prediction(
    uploaded_workbook: Path,
    run_dir: Path,
    *,
    selected_model_id: str,
    forecast_planting_date: str,
    forecast_harvesting_date: str,
    selected_germplasm_names: list[str],
    selected_germplasm_row_ids: list[str],
    selected_germplasm_selection_mode: str = "all_matches",
    regional_bounds_label: str,
    regional_bounds_latitude_min: float,
    regional_bounds_latitude_max: float,
    regional_bounds_longitude_min: float,
    regional_bounds_longitude_max: float,
    nasa_grid_resolution_km: int = 10,
    use_source_row_dates: bool = True,
    selected_id_header: str = "",
    progress_callback=None,
) -> CeSavedModelPredictionOutputs:
    selected_id_header = str(selected_id_header or "").strip()
    registered_model, selection_summary, selection_summary_file = _resolve_saved_model(selected_model_id)
    saved_normalization_stats = _read_saved_normalization_stats(registered_model.metadata, registered_model.model_dir)
    validate_saved_model_prediction_workbook(
        uploaded_workbook,
        selected_model_id=selected_model_id,
        climate_scope="regional_manual",
    )

    ce_dir = run_dir / "ce_pipeline_prediction"
    phase02_dir = ce_dir / "phase02"
    phase03_dir = ce_dir / "phase03"
    phase04_dir = ce_dir / "phase04"
    phase05_dir = ce_dir / "phase05"
    for directory in (phase02_dir, phase03_dir, phase04_dir, phase05_dir):
        directory.mkdir(parents=True, exist_ok=True)

    if progress_callback:
        progress_callback(
            12,
            "Preparing ce prediction input",
            "Validating the uploaded workbook against the saved model definition and building the prediction phase02 workbook.",
            None,
        )

    phase02_prediction_path = phase02_dir / "phase02prediction.xlsx"
    phase02_log = build_prediction_phase02_workbook(
        uploaded_workbook,
        selection_summary=selection_summary,
        output_file=phase02_prediction_path,
        forecast_planting_date=forecast_planting_date,
        forecast_harvesting_date=forecast_harvesting_date,
        regional_bounds_label=regional_bounds_label,
        regional_bounds_latitude_min=regional_bounds_latitude_min,
        regional_bounds_latitude_max=regional_bounds_latitude_max,
        regional_bounds_longitude_min=regional_bounds_longitude_min,
        regional_bounds_longitude_max=regional_bounds_longitude_max,
        selected_germplasm_names=selected_germplasm_names,
        selected_germplasm_row_ids=selected_germplasm_row_ids,
        selected_germplasm_selection_mode=selected_germplasm_selection_mode,
        selected_id_header=selected_id_header,
    )
    (phase02_dir / "phase02prediction_log.json").write_text(
        json.dumps(phase02_log, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    if progress_callback:
        progress_callback(
            22,
            "Preparing phase03 climate enrichment",
            "Expanding the selected germplasm rows across the manual bounding-box grid and preparing CHIRPS/CHIRTS-daily climate downloads.",
            {"phase": "phase03", "climate_days_processed": 0, "climate_days_total": 0},
        )

    phase03_prediction_path = phase03_dir / "phase03prediction.xlsx"
    phase03_log_path = phase03_dir / "phase03prediction_log.json"
    phase03_log = create_phase03_workbook(
        phase02_prediction_path,
        selection_summary_file,
        phase03_prediction_path,
        phase03_log_path,
        progress_file=None,
        skip_soil_enrichment=False,
        progress_callback=progress_callback,
        climate_override={
            "climate_scope": "regional_manual",
            "forecast_years_back": 5,
            "selected_model_id": selected_model_id,
            "regional_bounds_label": regional_bounds_label,
            "regional_bounds_latitude_min": regional_bounds_latitude_min,
            "regional_bounds_latitude_max": regional_bounds_latitude_max,
            "regional_bounds_longitude_min": regional_bounds_longitude_min,
            "regional_bounds_longitude_max": regional_bounds_longitude_max,
            "forecast_planting_date": forecast_planting_date,
            "forecast_harvesting_date": forecast_harvesting_date,
            "nasa_grid_resolution_km": nasa_grid_resolution_km,
            "use_source_row_dates": use_source_row_dates,
        },
    )

    if progress_callback:
        progress_callback(
            72,
            "Building phase04prediction",
            "Preparing the model-ready ce prediction dataset while preserving traceability columns.",
            None,
        )

    phase04_prediction_path = phase04_dir / "phase04prediction.xlsx"
    phase04_log_path = phase04_dir / "phase04prediction_log.json"
    phase04_log = create_phase04_workbook(
        phase03_prediction_path,
        selection_summary_file,
        phase04_prediction_path,
        phase04_log_path,
        allow_missing_target=True,
        stream_traceability_workbook=True,
        progress_callback=progress_callback,
        normalization_overrides=saved_normalization_stats,
    )
    _restore_selected_id_column_in_phase04_workbook(
        phase03_workbook=phase03_prediction_path,
        phase04_workbook=phase04_prediction_path,
        selected_id_header=selected_id_header,
    )
    phase04_training_csv = phase04_dir / "phase04prediction_training_input.csv"

    if progress_callback:
        progress_callback(
            90,
            "Running saved model prediction",
            "Executing the selected Grain Yield model using phase04prediction.xlsx as the prediction input.",
            None,
        )

    target_column = str((selection_summary.get("initial_settings", {}) if isinstance(selection_summary, dict) else {}).get("target_column", "")).strip()
    phase05_outputs = run_ce_phase05(
        phase04_workbook=phase04_prediction_path,
        training_input_csv=phase04_training_csv,
        run_dir=phase05_dir,
        source_name=uploaded_workbook.name,
        target_column=target_column,
        selected_id_header=selected_id_header,
        selected_model_id=selected_model_id,
        progress_file=None,
        progress_callback=progress_callback,
    )

    final_prediction_xlsx = ce_dir / "prediction.xlsx"
    final_prediction_csv = ce_dir / "prediction.csv"
    display_geojson_file = ce_dir / "prediction_display.geojson"
    initial_settings = selection_summary.get("initial_settings", {}) if isinstance(selection_summary, dict) else {}
    if _should_use_manual_grid_dataframe_finalization(
        climate_scope="regional_manual",
        selected_id_header=selected_id_header,
    ):
        phase05_prediction_csv = Path(phase05_outputs.summary["prediction_csv"])
        if phase05_prediction_csv.exists():
            final_prediction_df = pd.read_csv(phase05_prediction_csv)
        else:
            final_prediction_df = pd.read_excel(Path(phase05_outputs.summary["prediction_xlsx"]))
    else:
        shutil.copy2(Path(phase05_outputs.summary["prediction_xlsx"]), final_prediction_xlsx)
        shutil.copy2(Path(phase05_outputs.summary["prediction_csv"]), final_prediction_csv)
        final_prediction_df = pd.read_excel(final_prediction_xlsx)
    if not _should_preserve_manual_grid_prediction_rows(
        climate_scope="regional_manual",
        selected_id_header=selected_id_header,
        selected_germplasm_names=selected_germplasm_names,
    ):
        final_prediction_df = _aggregate_prediction_dataframe_by_selected_id(
            final_prediction_df,
            selected_id_header=selected_id_header,
        )
    final_prediction_df = _prepare_prediction_output_dataframe_for_display(
        final_prediction_df,
        latitude_column=str(initial_settings.get("latitude_column", "")).strip(),
        longitude_column=str(initial_settings.get("longitude_column", "")).strip(),
        selected_id_header=selected_id_header,
    )
    if _should_use_manual_grid_dataframe_finalization(
        climate_scope="regional_manual",
        selected_id_header=selected_id_header,
    ):
        final_prediction_df = _normalize_manual_grid_display_columns(
            final_prediction_df,
            manual_bbox_grid=(phase03_log.get("nasa", {}) if isinstance(phase03_log, dict) else {}).get("manual_bbox_grid"),
        )
    final_prediction_df.to_excel(final_prediction_xlsx, index=False)
    final_prediction_df.to_csv(final_prediction_csv, index=False)
    geojson_source_df = (
        _build_lightweight_manual_grid_geojson_dataframe(
            final_prediction_df,
            selected_id_header=selected_id_header,
        )
        if _should_use_manual_grid_dataframe_finalization(
            climate_scope="regional_manual",
            selected_id_header=selected_id_header,
        )
        else final_prediction_df
    )
    aggregated_geojson = dataframe_to_geojson(geojson_source_df, phase05_outputs.geojson_file)
    display_prediction_df = (
        _aggregate_manual_grid_prediction_dataframe_for_display(final_prediction_df)
        if _should_use_manual_grid_dataframe_finalization(
            climate_scope="regional_manual",
            selected_id_header=selected_id_header,
        )
        else final_prediction_df
    )
    display_geojson_source_df = (
        _build_lightweight_manual_grid_geojson_dataframe(
            display_prediction_df,
            selected_id_header=selected_id_header,
        )
        if _should_use_manual_grid_dataframe_finalization(
            climate_scope="regional_manual",
            selected_id_header=selected_id_header,
        )
        else display_prediction_df
    )
    display_geojson = dataframe_to_geojson(display_geojson_source_df, display_geojson_file)

    profile_metadata = phase02_log.get("prediction_profile_analysis", {}) if isinstance(phase02_log, dict) else {}
    filtered_geojson, profile_summary_updates = _postprocess_multi_profile_prediction_results(
        geojson=display_geojson,
        geojson_file=display_geojson_file,
        prediction_xlsx=final_prediction_xlsx,
        prediction_csv=final_prediction_csv,
        profile_metadata=profile_metadata if isinstance(profile_metadata, dict) else {},
    )
    if not _should_use_manual_grid_dataframe_finalization(
        climate_scope="regional_manual",
        selected_id_header=selected_id_header,
    ):
        _sort_prediction_output_files(
            prediction_xlsx=final_prediction_xlsx,
            prediction_csv=final_prediction_csv,
            latitude_column=str(initial_settings.get("latitude_column", "")).strip(),
            longitude_column=str(initial_settings.get("longitude_column", "")).strip(),
            selected_id_header=selected_id_header,
        )

    summary = {
        "workflow": "ce_saved_model_prediction",
        "run_dir": str(run_dir),
        "input_workbook": str(uploaded_workbook),
        "selected_model_id": selected_model_id,
        "selected_germplasm_names": selected_germplasm_names,
        "selected_germplasm_row_ids": selected_germplasm_row_ids,
        "selected_germplasm_selection_mode": selected_germplasm_selection_mode,
        "selected_id_header": selected_id_header,
        "climate_scope": "regional_manual",
        "phase06_xlsx": str(phase03_prediction_path),
        "phase02prediction_xlsx": str(phase02_prediction_path),
        "phase03prediction_xlsx": str(phase03_prediction_path),
        "phase04prediction_xlsx": str(phase04_prediction_path),
        "phase04prediction_training_csv": str(phase04_training_csv),
        "prediction_csv": str(final_prediction_csv),
        "prediction_xlsx": str(final_prediction_xlsx),
        "geojson_file": str(phase05_outputs.geojson_file),
        "display_geojson_file": str(display_geojson_file),
        "prediction_model_name": phase05_outputs.summary.get("prediction_model_name", ""),
        "prediction_feature_count": len(filtered_geojson.get("features", [])) if isinstance(filtered_geojson, dict) else len(display_prediction_df.index),
        "prediction_output_row_count": len(final_prediction_df.index),
        "prediction_model_id": phase05_outputs.summary.get("prediction_model_id", ""),
        "prediction_model_created_at": phase05_outputs.summary.get("prediction_model_created_at", ""),
        "prediction_model_display_name": phase05_outputs.summary.get("prediction_model_display_name", ""),
        "prediction_model_description": phase05_outputs.summary.get("prediction_model_description", ""),
        "prediction_model_metric_name": phase05_outputs.summary.get("prediction_model_metric_name", ""),
        "prediction_model_metric_value": phase05_outputs.summary.get("prediction_model_metric_value"),
        "manual_bbox_grid": (phase03_log.get("nasa", {}) if isinstance(phase03_log, dict) else {}).get("manual_bbox_grid"),
        "source_row_dates_used": use_source_row_dates,
        **profile_summary_updates,
        "preprocess_log": {
            "overview": {
                "input_rows": phase02_log.get("output_row_count", 0),
                "climate_exportable_rows": (phase03_log.get("nasa", {}) if isinstance(phase03_log, dict) else {}).get("climate_input_row_count", 0),
                "nasa_unique_queries": (phase03_log.get("nasa", {}) if isinstance(phase03_log, dict) else {}).get("nasa_unique_queries", 0),
                "soil_enriched_cells": (phase03_log.get("soil", {}) if isinstance(phase03_log, dict) else {}).get("soil_enriched_cells", 0),
            },
            "climate": (phase03_log.get("nasa", {}) if isinstance(phase03_log, dict) else {}),
            "soil": (phase03_log.get("soil", {}) if isinstance(phase03_log, dict) else {}),
            "phase04": phase04_log,
        },
        "task_report": [
            {"phase": "phase02prediction", "title": "Building ce prediction phase02 workbook", "output_file": str(phase02_prediction_path)},
            {"phase": "phase03prediction", "title": "Building ce prediction phase03 workbook", "output_file": str(phase03_prediction_path)},
            {"phase": "phase04prediction", "title": "Building ce prediction phase04 workbook", "output_file": str(phase04_prediction_path)},
            *phase05_outputs.summary.get("task_report", []),
        ],
    }
    summary_file = run_dir / "summary.json"
    summary_file.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return CeSavedModelPredictionOutputs(
        summary=summary,
        geojson=filtered_geojson,
        summary_file=summary_file,
        geojson_file=display_geojson_file,
    )
