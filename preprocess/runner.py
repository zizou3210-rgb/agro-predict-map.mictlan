from __future__ import annotations

import csv
import json
import re
import shutil
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook

from config_env import (
    get_prediction_bridge_dev_row_limit,
    get_preprocess_validation_enabled,
)
from preprocess import ea_pipeline
from preprocess.soil_enrichment import enrich_manual_bbox_prediction_soils


PHASE_DIR_NAMES = {
    "phase00": "phase00_selected_input",
    "phase01": "phase01_quality_prepared",
    "phase02": "phase02_climate_enriched",
    "phase03": "phase03_management_derived",
    "phase04": "phase04_fertilization_consolidated",
    "phase05": "phase05_final_with_idpk",
    "phase06": "phase06_phase1_compatible",
}

REQUIRED_PHASES = (
    "phase00",
    "phase01",
    "phase02",
    "phase03",
    "phase04",
    "phase05",
    "phase06",
)


@dataclass
class PreprocessOutputs:
    run_dir: Path
    input_workbook: Path
    phase06_xlsx: Path
    summary_file: Path
    summary: dict[str, object]


ProgressCallback = Callable[[int, str, str, dict[str, object] | None], None]
LEGACY_OPTIONAL_HEADERS = {"Plot No."}
DG_HEADER_PATTERN = re.compile(r"^DG\d+$")
TEMP_SOURCE_ROW_ID_HEADER = "__app_source_row_id"
APP_SHARED_NASA_CACHE = Path(__file__).resolve().parents[1] / ".cache" / "nasa_power_cache.json"
APP_SHARED_SOIL_CACHE = Path(__file__).resolve().parents[1] / ".cache" / "soilgrids_cache.json"
POINT_UPLOAD_REQUIRED_HEADERS = (
    "_GPS coordinates_latitude",
    "_GPS coordinates_longitude",
)
ALL_MARKERS_BASE_LOCATION_COLUMNS = [
    "Country",
    "GPS coordinates",
    "_GPS coordinates_latitude",
    "_GPS coordinates_longitude",
    "_GPS coordinates_altitude",
    "_GPS coordinates_precision",
    "Trial series name",
    "Farm",
    "Site Number",
    "Plot",
    "Rep",
]
COORDINATE_UPLOAD_OVERRIDE_COLUMNS = [
    "GPS coordinates",
    "_GPS coordinates_latitude",
    "_GPS coordinates_longitude",
    "_GPS coordinates_altitude",
    "_GPS coordinates_precision",
]
POINT_SAVED_MODEL_BASE_OVERRIDE_COLUMNS = [
    *ALL_MARKERS_BASE_LOCATION_COLUMNS,
    "Date of planting",
    "Date of thinning",
    "Date_of_harvesting",
    "Grain Yield (T/Ha)",
]
MANUAL_BBOX_CLEAR_COLUMNS = [
    "Country",
    "GPS coordinates",
    "_GPS coordinates_latitude",
    "_GPS coordinates_longitude",
    "_GPS coordinates_altitude",
    "_GPS coordinates_precision",
    "Date of planting",
    "Date of thinning",
    "Date_of_harvesting",
    "Grain Yield (T/Ha)",
    "Soil type/texture",
    "Soil Depth (cm)-How deep do you expect maize roots to grow in this soil?",
]
PREDICTION_BRIDGE_REQUIRED_DATE_COLUMNS = ("Date of planting", "Date_of_harvesting")
PREDICTION_BRIDGE_GRAIN_YIELD_HEADER = "Grain Yield (T/Ha)"


def normalize_selected_germplasm_name(value: object) -> str:
    normalized = ea_pipeline.normalize_upper(value)
    return ea_pipeline.NAME_REPLACEMENTS.get(normalized, normalized)


def read_canonical_headers() -> list[str]:
    template_csv = (
        Path(__file__).resolve().parents[1] / "template" / "canonical_template_superset_88.csv"
    )
    with template_csv.open("r", encoding="utf-8", newline="") as handle:
        return next(csv.reader(handle))


def load_first_sheet_headers(workbook_path: Path) -> list[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        return [str(value).strip() for value in header_row if value is not None]
    finally:
        workbook.close()


def list_distinct_germplasm_names(uploaded_test_xlsx: Path) -> list[str]:
    validation_errors = validate_uploaded_test_workbook(uploaded_test_xlsx)
    if validation_errors:
        raise ValueError(
            "The uploaded workbook is not a valid app test.xlsx contract:\n- "
            + "\n- ".join(validation_errors)
        )

    workbook = load_workbook(uploaded_test_xlsx, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows)
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        if "Name" not in headers:
            raise ValueError("The uploaded workbook does not contain the required Name column.")
        name_index = headers.index("Name")
        names: set[str] = set()
        for row in rows:
            if name_index >= len(row):
                continue
            raw_value = row[name_index]
            name = str(raw_value).strip() if raw_value is not None else ""
            if name:
                names.add(name)
        return sorted(names)
    finally:
        workbook.close()


def validate_uploaded_test_workbook(uploaded_test_xlsx: Path) -> list[str]:
    expected_headers = read_canonical_headers()
    actual_headers = load_first_sheet_headers(uploaded_test_xlsx)
    errors: list[str] = []
    if actual_headers != expected_headers:
        expected_only = [header for header in expected_headers if header not in actual_headers]
        actual_only = [
            header
            for header in actual_headers
            if header not in expected_headers
            and header not in LEGACY_OPTIONAL_HEADERS
            and not DG_HEADER_PATTERN.match(header)
        ]
        legacy_optional = [header for header in actual_headers if header in LEGACY_OPTIONAL_HEADERS]
        if expected_only:
            errors.append("Missing required headers: " + ", ".join(expected_only))
        if actual_only:
            errors.append("Unexpected headers: " + ", ".join(actual_only))
        if not errors and not legacy_optional:
            errors.append("The workbook headers do not match the canonical test.xlsx order.")
    return errors


def normalize_uploaded_test_workbook(uploaded_test_xlsx: Path, destination_path: Path) -> None:
    expected_headers = read_canonical_headers()
    workbook = load_workbook(uploaded_test_xlsx, read_only=True, data_only=True)
    normalized = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = normalized.create_sheet(source_sheet.title)
        header_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        actual_headers = [str(value).strip() if value is not None else "" for value in header_row]
        header_to_index = {header: idx for idx, header in enumerate(actual_headers)}
        dg_headers = [header for header in actual_headers if DG_HEADER_PATTERN.match(header)]
        output_headers = [*expected_headers, *dg_headers]

        output_sheet.append(output_headers)
        for row in source_sheet.iter_rows(min_row=2, values_only=True):
            normalized_row = [
                (
                    row[header_to_index[header]]
                    if header in header_to_index and header_to_index[header] < len(row)
                    else None
                )
                for header in output_headers
            ]
            output_sheet.append(normalized_row)
        normalized.save(destination_path)
    finally:
        workbook.close()


def split_dg_columns_from_records(
    headers: list[str],
    records: list[dict[str, object]],
) -> tuple[list[str], list[dict[str, object]], list[str], dict[str, dict[str, object]]]:
    dg_headers = [header for header in headers if DG_HEADER_PATTERN.match(header)]
    if not dg_headers:
        return headers, records, [], {}

    stripped_headers = [header for header in headers if header not in dg_headers]
    if TEMP_SOURCE_ROW_ID_HEADER not in stripped_headers:
        stripped_headers.append(TEMP_SOURCE_ROW_ID_HEADER)

    dg_by_row_id: dict[str, dict[str, object]] = {}
    stripped_records: list[dict[str, object]] = []
    for index, record in enumerate(records, start=1):
        row_id = str(index)
        dg_by_row_id[row_id] = {header: record.get(header) for header in dg_headers}
        stripped_record = {
            header: value
            for header, value in record.items()
            if header not in dg_headers
        }
        stripped_record[TEMP_SOURCE_ROW_ID_HEADER] = row_id
        stripped_records.append(stripped_record)
    return stripped_headers, stripped_records, dg_headers, dg_by_row_id


def restore_dg_columns_into_phase05(
    headers: list[str],
    records: list[dict[str, object]],
    dg_headers: list[str],
    dg_by_row_id: dict[str, dict[str, object]],
) -> tuple[list[str], list[dict[str, object]]]:
    if not dg_headers or not dg_by_row_id:
        return headers, records

    restored_headers = list(headers)
    if TEMP_SOURCE_ROW_ID_HEADER not in restored_headers:
        raise ValueError("The preprocess row-id header was not preserved through phase05.")
    insertion_index = restored_headers.index(TEMP_SOURCE_ROW_ID_HEADER)
    restored_headers[insertion_index:insertion_index] = dg_headers

    restored_records: list[dict[str, object]] = []
    for record in records:
        restored = dict(record)
        row_id = str(restored.get(TEMP_SOURCE_ROW_ID_HEADER, "")).strip()
        dg_values = dg_by_row_id.get(row_id, {})
        for header in dg_headers:
            restored[header] = dg_values.get(header)
        restored_records.append(restored)
    return restored_headers, restored_records


def load_nasa_cache_payload(cache_path: Path) -> dict[str, object]:
    if not cache_path.exists():
        return {}
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return payload if isinstance(payload, dict) else {}


def merge_nasa_cache_files(source_cache: Path, destination_cache: Path) -> None:
    source_payload = load_nasa_cache_payload(source_cache)
    if not source_payload:
        return
    destination_payload = load_nasa_cache_payload(destination_cache)
    destination_payload.update(source_payload)
    destination_cache.parent.mkdir(parents=True, exist_ok=True)
    destination_cache.write_text(
        json.dumps(destination_payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def is_saved_model_manual_bbox_prediction(
    *,
    climate_scope: str | None,
    selected_model_id: str | None,
) -> bool:
    return bool(selected_model_id) and (climate_scope or "point") == "regional_manual"


def validate_uploaded_coordinate_workbook(uploaded_xlsx: Path) -> list[str]:
    actual_headers = load_first_sheet_headers(uploaded_xlsx)
    errors: list[str] = []
    missing_headers = [header for header in POINT_UPLOAD_REQUIRED_HEADERS if header not in actual_headers]
    unexpected_headers = [
        header
        for header in actual_headers
        if header not in POINT_UPLOAD_REQUIRED_HEADERS and header != "id"
    ]
    if missing_headers:
        errors.append("Missing required headers: " + ", ".join(missing_headers))
    if unexpected_headers:
        errors.append("Unexpected headers: " + ", ".join(unexpected_headers))
    return errors


def normalize_uploaded_coordinate_workbook(uploaded_xlsx: Path, destination_path: Path) -> None:
    expected_headers = read_canonical_headers()
    workbook = load_workbook(uploaded_xlsx, read_only=True, data_only=True)
    normalized = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = normalized.create_sheet(source_sheet.title)
        header_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        actual_headers = [str(value).strip() if value is not None else "" for value in header_row]
        header_to_index = {header: idx for idx, header in enumerate(actual_headers)}
        latitude_index = header_to_index["_GPS coordinates_latitude"]
        longitude_index = header_to_index["_GPS coordinates_longitude"]

        output_sheet.append(expected_headers)
        for row in source_sheet.iter_rows(min_row=2, values_only=True):
            latitude = row[latitude_index] if latitude_index < len(row) else None
            longitude = row[longitude_index] if longitude_index < len(row) else None
            normalized_row = [None] * len(expected_headers)
            if "_GPS coordinates_latitude" in expected_headers:
                normalized_row[expected_headers.index("_GPS coordinates_latitude")] = latitude
            if "_GPS coordinates_longitude" in expected_headers:
                normalized_row[expected_headers.index("_GPS coordinates_longitude")] = longitude
            if "GPS coordinates" in expected_headers and latitude is not None and longitude is not None:
                normalized_row[expected_headers.index("GPS coordinates")] = f"{latitude},{longitude}"
            output_sheet.append(normalized_row)
        normalized.save(destination_path)
    finally:
        workbook.close()


def resolve_point_uploaded_workbook_mode(
    uploaded_xlsx: Path,
) -> tuple[str | None, list[str]]:
    coordinate_errors = validate_uploaded_coordinate_workbook(uploaded_xlsx)
    if not coordinate_errors:
        return "coordinate", []
    test_workbook_errors = validate_uploaded_test_workbook(uploaded_xlsx)
    if not test_workbook_errors:
        return "test", []
    return None, coordinate_errors


def filter_normalized_test_workbook_by_germplasm_names(
    normalized_workbook: Path,
    destination_path: Path,
    selected_germplasm_names: list[str],
) -> None:
    selected_names = {
        normalize_selected_germplasm_name(name)
        for name in selected_germplasm_names
        if str(name).strip()
    }
    if not selected_names:
        shutil.copy2(normalized_workbook, destination_path)
        return

    workbook = load_workbook(normalized_workbook, read_only=True, data_only=True)
    filtered = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = filtered.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        if "Name" not in headers:
            raise ValueError("The normalized workbook does not contain the required Name column.")
        name_index = headers.index("Name")
        output_sheet.append(headers)

        for row in rows:
            if name_index >= len(row):
                continue
            raw_value = row[name_index]
            name = normalize_selected_germplasm_name(raw_value)
            if name in selected_names:
                output_sheet.append(list(row))

        filtered.save(destination_path)
    finally:
        workbook.close()


def parse_workbook_numeric_value(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            numeric = float(value)
        except (TypeError, ValueError):
            return None
        if numeric != numeric:
            return None
        return numeric
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def filter_normalized_test_workbook_for_prediction_bridge(
    normalized_workbook: Path,
    destination_path: Path,
    *,
    row_limit: int | None = None,
) -> None:
    workbook = load_workbook(normalized_workbook, read_only=True, data_only=True)
    filtered = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = filtered.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        header_to_index = {header: idx for idx, header in enumerate(headers)}
        output_sheet.append(headers)

        planting_index = header_to_index.get(PREDICTION_BRIDGE_REQUIRED_DATE_COLUMNS[0])
        harvesting_index = header_to_index.get(PREDICTION_BRIDGE_REQUIRED_DATE_COLUMNS[1])
        grain_yield_index = header_to_index.get(PREDICTION_BRIDGE_GRAIN_YIELD_HEADER)

        kept_rows = 0
        for row in rows:
            if planting_index is None or harvesting_index is None or grain_yield_index is None:
                if row_limit and kept_rows >= row_limit:
                    break
                output_sheet.append(list(row))
                kept_rows += 1
                continue
            planting_value = row[planting_index] if planting_index < len(row) else None
            harvesting_value = row[harvesting_index] if harvesting_index < len(row) else None
            grain_yield_value = row[grain_yield_index] if grain_yield_index < len(row) else None

            if planting_value is None or str(planting_value).strip() == "":
                continue
            if harvesting_value is None or str(harvesting_value).strip() == "":
                continue
            parsed_grain_yield = parse_workbook_numeric_value(grain_yield_value)
            if parsed_grain_yield is not None and parsed_grain_yield <= 0:
                continue

            if row_limit and kept_rows >= row_limit:
                break
            output_sheet.append(list(row))
            kept_rows += 1

        filtered.save(destination_path)
    finally:
        workbook.close()


def limit_normalized_workbook_rows(
    normalized_workbook: Path,
    destination_path: Path,
    *,
    row_limit: int | None = None,
) -> None:
    if not row_limit or row_limit <= 0:
        shutil.copy2(normalized_workbook, destination_path)
        return

    workbook = load_workbook(normalized_workbook, read_only=True, data_only=True)
    limited = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = limited.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        output_sheet.append(headers)

        kept_rows = 0
        for row in rows:
            if kept_rows >= row_limit:
                break
            output_sheet.append(list(row))
            kept_rows += 1

        limited.save(destination_path)
    finally:
        workbook.close()


def clear_normalized_workbook_columns(
    normalized_workbook: Path,
    destination_path: Path,
    *,
    columns_to_clear: list[str],
) -> None:
    workbook = load_workbook(normalized_workbook, read_only=True, data_only=True)
    cleared = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = cleared.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        header_to_index = {header: idx for idx, header in enumerate(headers)}
        clear_indices = [header_to_index[column] for column in columns_to_clear if column in header_to_index]
        output_sheet.append(headers)

        for row in rows:
            row_values = list(row)
            if len(row_values) < len(headers):
                row_values.extend([None] * (len(headers) - len(row_values)))
            for column_index in clear_indices:
                row_values[column_index] = None
            output_sheet.append(row_values)
        cleared.save(destination_path)
    finally:
        workbook.close()


def project_normalized_test_workbook_across_all_markers(
    normalized_workbook: Path,
    destination_path: Path,
    selected_germplasm_names: list[str],
    template_workbook: Path | None = None,
    base_override_columns: list[str] | None = None,
) -> None:
    ordered_names: list[str] = []
    seen_names: set[str] = set()
    for raw_name in selected_germplasm_names:
        name = normalize_selected_germplasm_name(raw_name)
        if name and name not in seen_names:
            ordered_names.append(name)
            seen_names.add(name)
    if not ordered_names:
        shutil.copy2(normalized_workbook, destination_path)
        return

    workbook = load_workbook(normalized_workbook, read_only=True, data_only=True)
    template_source_path = template_workbook or normalized_workbook
    template_source_workbook = load_workbook(template_source_path, read_only=True, data_only=True)
    projected = Workbook(write_only=True)
    try:
        source_sheet = workbook.worksheets[0]
        output_sheet = projected.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        if "Name" not in headers:
            raise ValueError("The normalized workbook does not contain the required Name column.")
        header_to_index = {header: idx for idx, header in enumerate(headers)}
        name_index = header_to_index["Name"]
        base_rows = [tuple(row) for row in rows]

        template_rows: dict[str, tuple[object, ...]] = {}
        template_sheet = template_source_workbook.worksheets[0]
        template_iter = template_sheet.iter_rows(values_only=True)
        template_headers = [
            str(value).strip() if value is not None else "" for value in next(template_iter)
        ]
        if "Name" not in template_headers:
            raise ValueError(
                "The template workbook used for the all-markers projection does not contain the required Name column."
            )
        template_header_to_index = {header: idx for idx, header in enumerate(template_headers)}
        shared_headers = [header for header in headers if header in template_header_to_index]
        template_name_index = template_header_to_index["Name"]
        for row in template_iter:
            if template_name_index >= len(row):
                continue
            raw_name = row[template_name_index]
            name = normalize_selected_germplasm_name(raw_name)
            if name in seen_names and name not in template_rows:
                template_rows[name] = row

        missing_names = [name for name in ordered_names if name not in template_rows]
        if missing_names:
            raise ValueError(
                "The uploaded workbook does not contain the selected germplasm Name values required for the all-markers projection:\n- "
                + "\n- ".join(missing_names)
            )

        override_columns = base_override_columns or ALL_MARKERS_BASE_LOCATION_COLUMNS
        base_location_indices = [
            header_to_index[column]
            for column in override_columns
            if column in header_to_index
        ]
        output_sheet.append(headers)
        for base_row in base_rows:
            for selected_name in ordered_names:
                template_row = template_rows[selected_name]
                projected_row = list(base_row)
                if len(projected_row) < len(headers):
                    projected_row.extend([None] * (len(headers) - len(projected_row)))
                for header in shared_headers:
                    base_index = header_to_index[header]
                    template_index = template_header_to_index[header]
                    projected_row[base_index] = (
                        template_row[template_index]
                        if template_index < len(template_row)
                        else None
                    )
                for column_index in base_location_indices:
                    if len(projected_row) <= column_index:
                        projected_row.extend([None] * (column_index + 1 - len(projected_row)))
                    projected_row[column_index] = (
                        base_row[column_index] if column_index < len(base_row) else None
                    )
                output_sheet.append(projected_row)

        projected.save(destination_path)
    finally:
        workbook.close()
        template_source_workbook.close()


def override_normalized_workbook_forecast_dates(
    normalized_workbook: Path,
    *,
    forecast_planting_date: str | None = None,
    forecast_harvesting_date: str | None = None,
) -> None:
    planting_value = str(forecast_planting_date or "").strip()
    harvesting_value = str(forecast_harvesting_date or "").strip()
    if not planting_value and not harvesting_value:
        return

    workbook = load_workbook(normalized_workbook)
    try:
        worksheet = workbook.worksheets[0]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        header_to_index = {header: idx + 1 for idx, header in enumerate(headers)}
        planting_column = header_to_index.get("Date of planting")
        harvesting_column = header_to_index.get("Date_of_harvesting")

        if planting_value and planting_column:
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=planting_column, value=planting_value)
        if harvesting_value and harvesting_column:
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=harvesting_column, value=harvesting_value)

        workbook.save(normalized_workbook)
    finally:
        workbook.close()


@contextmanager
def override_phase_dirs(output_root: Path):
    original_phase_dirs = ea_pipeline.PHASE_DIRS
    try:
        ea_pipeline.PHASE_DIRS = {
            phase_name: output_root / folder_name
            for phase_name, folder_name in PHASE_DIR_NAMES.items()
        }
        yield ea_pipeline.PHASE_DIRS
    finally:
        ea_pipeline.PHASE_DIRS = original_phase_dirs


def build_preprocess_summary(
    phase_dirs: dict[str, Path],
    input_workbook: Path,
    phase01_metadata: dict[str, object],
    phase02_metadata: dict[str, object],
) -> dict[str, object]:
    phase06_dir = phase_dirs["phase06"]
    preprocess_log = {
        "overview": {
            "input_rows": phase01_metadata.get("input_row_count", 0),
            "quality_output_rows": phase01_metadata.get("output_row_count", 0),
            "climate_exportable_rows": phase02_metadata.get("climate_input_row_count", 0),
            "phase02_source_rows": phase02_metadata.get("phase02_source_row_count", 0),
            "nasa_unique_queries": phase02_metadata.get("nasa_unique_queries", 0),
            "nasa_cache_fresh_queries_this_run": phase02_metadata.get(
                "nasa_cache_fresh_queries_this_run", 0
            ),
            "nasa_cache_hits": phase02_metadata.get("nasa_cache_hits", 0),
            "soil_enriched_cells": phase02_metadata.get("soil_enriched_cells", 0),
            "soil_fallback_cells": phase02_metadata.get("soil_fallback_cells", 0),
            "soil_cache_hits": phase02_metadata.get("soil_cache_hits", 0),
            "soil_service_queries": phase02_metadata.get("soil_service_queries", 0),
            "soil_service_success_cells": phase02_metadata.get("soil_service_success_cells", 0),
        },
        "quality_filters": {
            "name_updates": phase01_metadata.get("name_updates", 0),
            "deleted_by_name_rule": phase01_metadata.get("deleted_by_name_rule", 0),
            "excluded_missing_yield": phase01_metadata.get("excluded_missing_yield", 0),
            "excluded_rank_yield_zero": phase01_metadata.get("excluded_rank_yield_zero", 0),
            "wrong_coordinates": phase01_metadata.get("wrong_coordinates", 0),
        },
        "thinning": {
            "imputed": phase01_metadata.get("thinning_imputed", 0),
            "pending": phase01_metadata.get("thinning_pending", 0),
            "group_trial_country": phase01_metadata.get("thinning_group_trial_country", 0),
            "group_trial": phase01_metadata.get("thinning_group_trial", 0),
            "group_country": phase01_metadata.get("thinning_group_country", 0),
            "group_global": phase01_metadata.get("thinning_group_global", 0),
        },
        "climate": {
            "source_rows": phase02_metadata.get("phase02_source_row_count", 0),
            "climate_input_rows": phase02_metadata.get("climate_input_row_count", 0),
            "skipped_rows_missing_climate_fields": phase02_metadata.get(
                "skipped_rows_missing_climate_fields", 0
            ),
            "matched_rows": phase02_metadata.get("matched_rows", 0),
            "unmatched_rows": phase02_metadata.get("unmatched_rows", 0),
            "nasa_unique_queries": phase02_metadata.get("nasa_unique_queries", 0),
            "nasa_cache_initial_entries": phase02_metadata.get("nasa_cache_initial_entries", 0),
            "nasa_cache_final_entries": phase02_metadata.get("nasa_cache_final_entries", 0),
            "nasa_cache_fresh_queries_this_run": phase02_metadata.get(
                "nasa_cache_fresh_queries_this_run", 0
            ),
            "nasa_cache_hits": phase02_metadata.get("nasa_cache_hits", 0),
            "nasa_cache_misses": phase02_metadata.get("nasa_cache_misses", 0),
            "locality_count": phase02_metadata.get(
                "locality_count",
                phase02_metadata.get("manual_bbox_locality_count", 0),
            ),
            "expanded_row_count": phase02_metadata.get("expanded_row_count", 0),
            "projected_germplasm_count": phase02_metadata.get("projected_germplasm_count", 0),
            "soil_enriched_rows": phase02_metadata.get("soil_enriched_rows", 0),
            "soil_enriched_cells": phase02_metadata.get("soil_enriched_cells", 0),
            "soil_fallback_cells": phase02_metadata.get("soil_fallback_cells", 0),
            "soil_cache_hits": phase02_metadata.get("soil_cache_hits", 0),
            "soil_service_queries": phase02_metadata.get("soil_service_queries", 0),
            "soil_service_success_cells": phase02_metadata.get("soil_service_success_cells", 0),
        },
        "operations": [
            {
                "phase": "phase00",
                "title": "Input snapshot",
                "summary": "Selected the workbook and created the normalized preprocess workspace.",
                "metrics": {
                    "input_rows": phase01_metadata.get("input_row_count", 0),
                },
            },
            {
                "phase": "phase01",
                "title": "Quality preparation",
                "summary": "Applied name rules, excluded invalid rows, checked coordinates, and imputed thinning dates when possible.",
                "metrics": {
                    "output_rows": phase01_metadata.get("output_row_count", 0),
                    "deleted_by_name_rule": phase01_metadata.get("deleted_by_name_rule", 0),
                    "excluded_missing_yield": phase01_metadata.get("excluded_missing_yield", 0),
                    "excluded_rank_yield_zero": phase01_metadata.get("excluded_rank_yield_zero", 0),
                    "wrong_coordinates": phase01_metadata.get("wrong_coordinates", 0),
                    "thinning_imputed": phase01_metadata.get("thinning_imputed", 0),
                    "thinning_pending": phase01_metadata.get("thinning_pending", 0),
                },
            },
            {
                "phase": "phase02",
                "title": "Climate enrichment",
                "summary": "Built the NASA POWER climate input rows and joined the resulting climate windows back into the preprocess workbook.",
                "metrics": {
                    "source_rows": phase02_metadata.get("phase02_source_row_count", 0),
                    "climate_input_rows": phase02_metadata.get("climate_input_row_count", 0),
                    "skipped_rows_missing_climate_fields": phase02_metadata.get(
                        "skipped_rows_missing_climate_fields", 0
                    ),
                    "nasa_unique_queries": phase02_metadata.get("nasa_unique_queries", 0),
                    "nasa_cache_fresh_queries_this_run": phase02_metadata.get(
                        "nasa_cache_fresh_queries_this_run", 0
                    ),
                    "nasa_cache_hits": phase02_metadata.get("nasa_cache_hits", 0),
                    "expanded_row_count": phase02_metadata.get("expanded_row_count", 0),
                },
            },
            {
                "phase": "phase03",
                "title": "Management derivation",
                "summary": "Derived management indicators from the climate-enriched workbook.",
                "metrics": {},
            },
            {
                "phase": "phase04",
                "title": "Agronomic consolidation",
                "summary": "Consolidated fertilization and agronomic metrics into the preprocess output.",
                "metrics": {},
            },
            {
                "phase": "phase05",
                "title": "Primary-key assignment",
                "summary": "Assigned record identifiers before the phase06 projection.",
                "metrics": {},
            },
            {
                "phase": "phase06",
                "title": "Phase06 projection",
                "summary": "Projected the final preprocess workbook used downstream for modeling and map generation.",
                "metrics": {},
            },
        ],
    }
    return {
        "input_workbook": str(input_workbook),
        "phase06_xlsx": str(phase06_dir / "phase06_phase1_compatible.xlsx"),
        "preprocess_log": preprocess_log,
        "manual_bbox_grid": phase02_metadata.get("manual_bbox_grid"),
        "phase_directories": {name: str(path) for name, path in phase_dirs.items()},
        "task_report": [
            {
                "phase": "phase00",
                "title": "Loading uploaded test workbook",
                "output_file": str(phase_dirs["phase00"] / "selected_input.xlsx"),
            },
            {
                "phase": "phase01",
                "title": "Preparing quality-controlled records",
                "row_count": phase01_metadata.get("output_row_count"),
                "output_file": str(phase_dirs["phase01"] / "phase01_quality_prepared.xlsx"),
            },
            {
                "phase": "phase02",
                "title": "Enriching climate windows",
                "exportable_rows": phase02_metadata.get("climate_input_row_count"),
                "nasa_unique_queries": phase02_metadata.get("nasa_unique_queries"),
                "output_file": str(phase_dirs["phase02"] / "phase02_climate_enriched.xlsx"),
            },
            {
                "phase": "phase03",
                "title": "Deriving management indicators",
                "output_file": str(phase_dirs["phase03"] / "phase03_management_derived.xlsx"),
            },
            {
                "phase": "phase04",
                "title": "Consolidating agronomic metrics",
                "output_file": str(
                    phase_dirs["phase04"] / "phase04_fertilization_consolidated.xlsx"
                ),
            },
            {
                "phase": "phase05",
                "title": "Assigning primary keys",
                "output_file": str(phase_dirs["phase05"] / "phase05_final_with_idpk.xlsx"),
            },
            {
                "phase": "phase06",
                "title": "Projecting the final preprocessing workbook",
                "output_file": str(phase06_dir / "phase06_phase1_compatible.xlsx"),
            },
        ],
    }


def run_preprocess_pipeline(
    uploaded_test_xlsx: Path,
    run_dir: Path,
    progress_callback: ProgressCallback | None = None,
    forecast_planting_date: str | None = None,
    forecast_harvesting_date: str | None = None,
    climate_scope: str | None = None,
    regional_country: str | None = None,
    selected_model_id: str | None = None,
    selected_germplasm_names: list[str] | None = None,
    selected_germplasm_projection_mode: str | None = None,
    projection_template_workbook: Path | None = None,
    regional_bounds_label: str | None = None,
    regional_bounds_latitude_min: float | None = None,
    regional_bounds_latitude_max: float | None = None,
    regional_bounds_longitude_min: float | None = None,
    regional_bounds_longitude_max: float | None = None,
    nasa_grid_resolution_km: int | None = None,
) -> PreprocessOutputs:
    run_dir.mkdir(parents=True, exist_ok=True)
    preprocess_validation_enabled = get_preprocess_validation_enabled()
    saved_model_manual_bbox_prediction = is_saved_model_manual_bbox_prediction(
        climate_scope=climate_scope,
        selected_model_id=selected_model_id,
    )
    saved_model_point_prediction = bool(selected_model_id) and (climate_scope or "point") == "point"
    is_bridge_projection = (
        selected_germplasm_projection_mode == "all_markers"
        and projection_template_workbook is not None
    )
    is_prediction_point_projection = is_bridge_projection and (climate_scope or "point") == "point"
    bridge_uploaded_workbook_mode: str | None = None
    if is_bridge_projection:
        bridge_uploaded_workbook_mode, validation_errors = resolve_point_uploaded_workbook_mode(
            uploaded_test_xlsx
        )
    else:
        validation_errors = validate_uploaded_test_workbook(uploaded_test_xlsx)
    if preprocess_validation_enabled and validation_errors:
        raise ValueError(
            (
                "The uploaded workbook is not a valid coordinate workbook for point-by-uploaded-coordinates:\n- "
                if is_prediction_point_projection
                else "The uploaded workbook is not a valid app test.xlsx contract:\n- "
            )
            + "\n- ".join(validation_errors)
        )

    input_dir = run_dir / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    input_workbook = input_dir / "test.xlsx"
    if is_bridge_projection and bridge_uploaded_workbook_mode == "coordinate":
        if is_prediction_point_projection:
            normalize_uploaded_coordinate_workbook(uploaded_test_xlsx, input_workbook)
        else:
            if projection_template_workbook is None or not projection_template_workbook.exists():
                raise FileNotFoundError(
                    f"The projection template workbook was not found: {projection_template_workbook}"
                )
            normalize_uploaded_test_workbook(projection_template_workbook, input_workbook)
    else:
        normalize_uploaded_test_workbook(uploaded_test_xlsx, input_workbook)
    if selected_germplasm_names:
        if selected_germplasm_projection_mode == "all_markers":
            if projection_template_workbook is not None and (climate_scope or "point") == "point":
                if not projection_template_workbook.exists():
                    raise FileNotFoundError(
                        f"The projection template workbook was not found: {projection_template_workbook}"
                    )
                bridge_row_limit = get_prediction_bridge_dev_row_limit()
                if bridge_row_limit:
                    limited_input_workbook = input_dir / "test_bridge_source.xlsx"
                    if bridge_uploaded_workbook_mode == "coordinate":
                        limit_normalized_workbook_rows(
                            input_workbook,
                            limited_input_workbook,
                            row_limit=bridge_row_limit,
                        )
                    else:
                        if saved_model_point_prediction:
                            limit_normalized_workbook_rows(
                                input_workbook,
                                limited_input_workbook,
                                row_limit=bridge_row_limit,
                            )
                        else:
                            filter_normalized_test_workbook_for_prediction_bridge(
                                input_workbook,
                                limited_input_workbook,
                                row_limit=bridge_row_limit,
                            )
                    input_workbook = limited_input_workbook
                template_source_workbook = projection_template_workbook
            elif projection_template_workbook is not None and bridge_uploaded_workbook_mode == "coordinate":
                filtered_input_workbook = input_dir / "test_bridge_manual_filtered.xlsx"
                filter_normalized_test_workbook_by_germplasm_names(
                    input_workbook,
                    filtered_input_workbook,
                    selected_germplasm_names,
                )
                input_workbook = filtered_input_workbook
                template_source_workbook = None
            else:
                full_bridge_source_workbook = input_dir / "test_bridge_source_full.xlsx"
                bridge_source_workbook = input_dir / "test_bridge_source.xlsx"
                bridge_row_limit = get_prediction_bridge_dev_row_limit()
                if saved_model_manual_bbox_prediction or saved_model_point_prediction:
                    shutil.copy2(input_workbook, full_bridge_source_workbook)
                    limit_normalized_workbook_rows(
                        full_bridge_source_workbook,
                        bridge_source_workbook,
                        row_limit=bridge_row_limit or None,
                    )
                else:
                    filter_normalized_test_workbook_for_prediction_bridge(
                        input_workbook,
                        full_bridge_source_workbook,
                        row_limit=None,
                    )
                    filter_normalized_test_workbook_for_prediction_bridge(
                        full_bridge_source_workbook,
                        bridge_source_workbook,
                        row_limit=bridge_row_limit or None,
                    )
                input_workbook = bridge_source_workbook
                template_source_workbook = full_bridge_source_workbook
            if template_source_workbook is not None:
                projected_input_workbook = input_dir / "test_projected_selected_germplasm.xlsx"
                base_override_columns = (
                    COORDINATE_UPLOAD_OVERRIDE_COLUMNS
                    if bridge_uploaded_workbook_mode == "coordinate"
                    else (
                        POINT_SAVED_MODEL_BASE_OVERRIDE_COLUMNS
                        if saved_model_point_prediction
                        else None
                    )
                )
                project_normalized_test_workbook_across_all_markers(
                    input_workbook,
                    projected_input_workbook,
                    selected_germplasm_names,
                    template_workbook=template_source_workbook,
                    base_override_columns=base_override_columns,
                )
                if (
                    projection_template_workbook is not None
                    and (climate_scope or "point") == "point"
                    and bridge_uploaded_workbook_mode == "coordinate"
                ):
                    override_normalized_workbook_forecast_dates(
                        projected_input_workbook,
                        forecast_planting_date=forecast_planting_date,
                        forecast_harvesting_date=forecast_harvesting_date,
                    )
                if saved_model_manual_bbox_prediction:
                    cleared_projected_input_workbook = input_dir / "test_projected_selected_germplasm_bbox.xlsx"
                    clear_normalized_workbook_columns(
                        projected_input_workbook,
                        cleared_projected_input_workbook,
                        columns_to_clear=MANUAL_BBOX_CLEAR_COLUMNS,
                    )
                    projected_input_workbook = cleared_projected_input_workbook
                input_workbook = projected_input_workbook
        else:
            filtered_input_workbook = input_dir / "test_filtered.xlsx"
            filter_normalized_test_workbook_by_germplasm_names(
                input_workbook,
                filtered_input_workbook,
                selected_germplasm_names,
            )
            input_workbook = filtered_input_workbook

    with override_phase_dirs(run_dir) as phase_dirs:
        for phase_name in REQUIRED_PHASES:
            phase_dirs[phase_name].mkdir(parents=True, exist_ok=True)

        if progress_callback:
            progress_callback(
                12,
                "Loading uploaded workbook",
                "Validating the uploaded test.xlsx contract and preparing the preprocessing input.",
                None,
            )
        headers, records = ea_pipeline.build_phase00_snapshot(input_workbook, 0)
        headers, records, dg_headers, dg_by_row_id = split_dg_columns_from_records(headers, records)
        if progress_callback:
            progress_callback(
                22,
                "Preparing quality records",
                "Applying app preprocess quality rules, exclusions, coordinate checks, and thinning cleanup.",
                None,
            )
        phase01_headers, phase01_records, phase01_metadata = ea_pipeline.build_phase01_prepared(
            headers,
            records,
            allow_missing_yield=saved_model_manual_bbox_prediction,
            allow_missing_coordinates=saved_model_manual_bbox_prediction,
        )
        if progress_callback:
            progress_callback(
                34,
                "Enriching climate windows",
                "Running the preprocess climate stage and querying NASA POWER for valid phenological windows.",
                None,
            )
        phase02_cache = phase_dirs["phase02"] / "nasa_power_cache.json"
        if APP_SHARED_NASA_CACHE.exists() and not phase02_cache.exists():
            phase02_cache.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(APP_SHARED_NASA_CACHE, phase02_cache)
        phase02_headers, phase02_records, phase02_metadata = ea_pipeline.build_phase02_records(
            phase01_headers,
            phase01_records,
            progress_callback=progress_callback,
            climate_override=(
                {
                    "forecast_planting_date": (
                        forecast_planting_date if not saved_model_point_prediction else None
                    ),
                    "forecast_harvesting_date": (
                        forecast_harvesting_date if not saved_model_point_prediction else None
                    ),
                    "forecast_years_back": 5,
                    "climate_scope": climate_scope or "point",
                    "regional_country": regional_country or "",
                    "selected_model_id": selected_model_id or "",
                    "regional_bounds_label": regional_bounds_label or "",
                    "regional_bounds_latitude_min": regional_bounds_latitude_min,
                    "regional_bounds_latitude_max": regional_bounds_latitude_max,
                    "regional_bounds_longitude_min": regional_bounds_longitude_min,
                    "regional_bounds_longitude_max": regional_bounds_longitude_max,
                    "nasa_grid_resolution_km": nasa_grid_resolution_km,
                }
                if (
                    forecast_planting_date
                    and forecast_harvesting_date
                    and not saved_model_point_prediction
                )
                else (
                    {
                        "climate_scope": climate_scope or "point",
                        "regional_country": regional_country or "",
                        "selected_model_id": selected_model_id or "",
                        "regional_bounds_label": regional_bounds_label or "",
                        "regional_bounds_latitude_min": regional_bounds_latitude_min,
                        "regional_bounds_latitude_max": regional_bounds_latitude_max,
                        "regional_bounds_longitude_min": regional_bounds_longitude_min,
                        "regional_bounds_longitude_max": regional_bounds_longitude_max,
                        "nasa_grid_resolution_km": nasa_grid_resolution_km,
                    }
                    if climate_scope
                    or regional_country
                    or regional_bounds_label
                    or regional_bounds_latitude_min is not None
                    or regional_bounds_latitude_max is not None
                    or regional_bounds_longitude_min is not None
                    or regional_bounds_longitude_max is not None
                    else None
                )
            ),
        )
        if phase02_cache.exists():
            merge_nasa_cache_files(phase02_cache, APP_SHARED_NASA_CACHE)
        if saved_model_manual_bbox_prediction:
            if progress_callback:
                progress_callback(
                    40,
                    "Extracting soil data",
                    "Fetching SoilGrids/ISRIC soil texture and depth values for each generated bounding-box grid cell.",
                    None,
                )
            phase02_soil_cache = phase_dirs["phase02"] / "soilgrids_cache.json"
            if APP_SHARED_SOIL_CACHE.exists() and not phase02_soil_cache.exists():
                phase02_soil_cache.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(APP_SHARED_SOIL_CACHE, phase02_soil_cache)
            phase02_records, soil_metadata = enrich_manual_bbox_prediction_soils(
                phase02_records,
                climate_scope=climate_scope or "point",
                selected_model_id=selected_model_id,
                cache_path=phase02_soil_cache,
            )
            if phase02_soil_cache.exists():
                merge_nasa_cache_files(phase02_soil_cache, APP_SHARED_SOIL_CACHE)
            phase02_metadata.update(soil_metadata)
            if progress_callback:
                progress_callback(
                    41,
                    "Soil data extracted",
                    (
                        f"Soil extraction finished for {soil_metadata.get('soil_enriched_cells', 0)} grid cells. "
                        f"ISRIC queries: {soil_metadata.get('soil_service_queries', 0)}. "
                        f"Cache hits: {soil_metadata.get('soil_cache_hits', 0)}. "
                        f"Fallback cells: {soil_metadata.get('soil_fallback_cells', 0)}."
                    ),
                    {
                        "soil_enriched_cells": soil_metadata.get("soil_enriched_cells", 0),
                        "soil_service_queries": soil_metadata.get("soil_service_queries", 0),
                        "soil_cache_hits": soil_metadata.get("soil_cache_hits", 0),
                        "soil_fallback_cells": soil_metadata.get("soil_fallback_cells", 0),
                    },
                )
        if progress_callback:
            progress_callback(
                42,
                "Deriving management indicators",
                "Creating preprocess phase03 management variables from the climate-enriched workbook.",
                None,
            )
        phase03_headers, phase03_records = ea_pipeline.build_phase03_records(
            phase02_headers, phase02_records
        )
        if progress_callback:
            progress_callback(
                46,
                "Consolidating agronomic preprocess output",
                "Building preprocess phases 04 to 06 and projecting the final phase06 workbook.",
                None,
            )
        phase04_headers, phase04_records = ea_pipeline.build_phase04_records(
            phase03_headers, phase03_records
        )
        phase05_headers, phase05_records = ea_pipeline.build_phase05_records(
            phase04_headers, phase04_records
        )
        phase05_headers, phase05_records = restore_dg_columns_into_phase05(
            phase05_headers,
            phase05_records,
            dg_headers,
            dg_by_row_id,
        )
        ea_pipeline.build_phase06_records(phase05_headers, phase05_records)

        summary = build_preprocess_summary(
            phase_dirs,
            input_workbook,
            phase01_metadata,
            phase02_metadata,
        )
        summary_file = run_dir / "summary.json"
        summary_file.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    phase06_xlsx = run_dir / PHASE_DIR_NAMES["phase06"] / "phase06_phase1_compatible.xlsx"
    return PreprocessOutputs(
        run_dir=run_dir,
        input_workbook=input_workbook,
        phase06_xlsx=phase06_xlsx,
        summary_file=summary_file,
        summary=summary,
    )
