#!/usr/bin/env python3
from __future__ import annotations

import argparse
import calendar
import csv
import json
import re
import shutil
import statistics
import time
from functools import lru_cache
from threading import RLock
from datetime import date, datetime, time as dt_time, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook

from preprocess.phase1_schema import OUTPUT_HEADERS as PHASE1_COMPAT_HEADERS
from pipeline.africa_country_localities import (
    load_country_localities,
    resolve_african_country,
)
from pipeline.nasa_country_region import (
    build_manual_grid_cells,
    get_country_bounds,
    iter_country_tile_bounds,
    iter_manual_tile_bounds,
    normalize_country_key,
)


ROOT_DIR = Path(__file__).resolve().parent
REPO_ROOT = ROOT_DIR.parent.parent
DEFAULT_SOURCE_WORKBOOK = REPO_ROOT / "template" / "test.xlsx"
DEFAULT_SOURCE_SHEET_INDEX = 0

NASA_POWER_URL = "https://power.larc.nasa.gov/api/temporal/daily/point"
NASA_POWER_REGIONAL_URL = "https://power.larc.nasa.gov/api/temporal/daily/regional"
NASA_PARAMETERS = ("T2M", "PRECTOTCORR")
NASA_FILL_VALUE = -999.0
HTTP_TIMEOUT_SECONDS = 60
MAX_RETRIES = 8
RETRY_SLEEP_SECONDS = 5
MAX_RETRY_SLEEP_SECONDS = 120
HARVEST_BACK_WEEKS = 8
SIMILARITY_THRESHOLD = 0.86

_NASA_SERIES_AGGREGATION_CACHE: dict[int, dict[str, object]] = {}
_NASA_WINDOW_METRICS_CACHE: dict[tuple[int, str, str], dict[str, float | None]] = {}
_NASA_AGGREGATION_LOCK = RLock()

PHASE_DIRS = {
    "phase00": ROOT_DIR / "phase00_selected_input",
    "phase01": ROOT_DIR / "phase01_quality_prepared",
    "phase02": ROOT_DIR / "phase02_climate_enriched",
    "phase03": ROOT_DIR / "phase03_management_derived",
    "phase04": ROOT_DIR / "phase04_fertilization_consolidated",
    "phase05": ROOT_DIR / "phase05_final_with_idpk",
    "phase06": ROOT_DIR / "phase06_phase1_compatible",
}

COUNTRY_BOUNDS_BY_SOURCE_SET = {
    "01": {
        "KENYA": (-5.5, 5.5, 33.0, 42.5),
        "TANZANIA": (-12.5, -0.5, 28.0, 41.5),
        "UGANDA": (-2.0, 5.0, 29.0, 35.5),
    },
    "02": {
        "KENYA": (-5.5, 5.5, 33.0, 42.5),
        "TANZANIA": (-12.5, -0.5, 29.0, 40.5),
        "UGANDA": (-1.6, 4.5, 29.0, 35.5),
    },
}

NAME_REPLACEMENTS = {
    "101 HARAKA": "HARAKA101",
    "10H": "LONGE10H",
    "4141": "WE4141",
    "508": "WH508",
    "511": "H511",
    "520": "H520",
    "6213": "H6213",
    "6232": "H6232",
    "8031": "DK8031",
    "8031DK": "DK8031",
    "BABY CON": "BABYCON",
    "BABYCON": "BABYCON",
    "BAZOOKA": "BAZOOKA",
    "CKDHH211040": "CKDHH211040",
    "CKDHH211196": "CKDHH211196",
    "CKDHH211274": "CKDHH211274",
    "CKDHH211544": "CKDHH211544",
    "CKH211940": "CKH211940",
    "CKH220194": "CKH220194",
    "CKH220291": "CKH220291",
    "CKH220969": "CKH220969",
    "CKHMLN210109": "CKHMLN210109",
    "CKHMLN210643": "CKHMLN210643",
    "CKHMLN221104": "CKHMLN221104",
    "CKHMLN221106": "CKHMLN221106",
    "CKHMLN221109": "CKHMLN221109",
    "CKHMLN221152": "CKHMLN221152",
    "CKHMLN221181": "CKHMLN221181",
    "DH 02": "DH02",
    "DH 04": "DH04",
    "DH02": "DH02",
    "DH04": "DH04",
    "DK 777": "DK77",
    "DK 80-31": "DK8031",
    "DK 80-33": "DK8033",
    "DK 8031": "DK8031",
    "DK 9089": "DK9089",
    "DK777": "DK777",
    "DK8031": "DK8031",
    "DK9089": "DK9089",
    "DUMA": "DUMA43",
    "DUMA 43": "DUMA43",
    "FARMER'S VARIETY": "DELETE",
    "FARMER’S VARIETY": "DELETE",
    "FORMER TRIALS": "DELETE",
    "GRANARY": "GRANARY",
    "H 511": "H511",
    "H513": "H513",
    "H520": "H520",
    "H6213": "H6213",
    "HARAKA": "HARAKA101",
    "HARAKA 101": "HARAKA101",
    "HODARI": "HODARI",
    "HYBRID 511": "H511",
    "KH 500-43": "KH500-43",
    "KH 500-43A": "KH500-43",
    "KIMERU": "KIMERU",
    "LINE MUNAANA(KAKONGOLIRO)": "MUNAANA",
    "LOCAL": "DELETE",
    "LOCAL MAIZE": "DELETE",
    "LOCAL VARIETY": "DELETE",
    "LOCK MAIZE": "DELETE",
    "LONGE 10H": "LONGE10H",
    "LONGE 5": "LONGE5",
    "LONGE10H": "LONGE10H",
    "LONGER 10H": "LONGE10H",
    "MAKUENI": "MAKUENI",
    "MIX": "DELETE",
    "MIX 513& KIMERU": "DELETE",
    "MIXED CHECK": "DELETE",
    "MIXED VARIETY": "DELETE",
    "MIXED,513&DH04": "DELETE",
    "MUNANDI": "MUNANDI",
    "NATA06": "NATA06",
    "OVER RECYCLED BAZOOKA": "OVERRECYCLEDBAZOOKA",
    "PAN 53": "PAN53",
    "PAN-53": "PAN53",
    "PIONEER": "PIONEER",
    "RECYCLED LONGE 10H": "RECYCLEDLONGE10H",
    "RECYCLED LONGE 5": "RECYCLEDLONGE5",
    "SAWA": "SAWA",
    "SC 419": "SC419",
    "SC 513": "SC513",
    "SC 55": "SC555",
    "SC 555": "SC555",
    "SC419": "SC419",
    "SC555": "SC555",
    "TOSHEKA": "TOSHEKA",
    "TSAVO 4141": "TSAVO 4141",
    "TSAVO4141": "TSAVO4141",
    "UH5354": "UH5354",
    "UH5355": "UH5355",
    "WE 4141": "WE 4141",
    "WE2115": "WE2115",
    "WE3106": "WE3106",
    "WH508": "WH508",
    "WH509": "WH509",
    "WS508": "WS508",
    "YELLOW": "DELETE",
    "YELLOW MAIZE": "DELETE",
}

CLIMATE_INPUT_COLUMNS = [
    "trial_series_name",
    "farm",
    "plot",
    "entry_code",
    "name",
    "country",
    "latitude",
    "longitude",
    "date_of_planting",
    "date_of_thinning",
    "date_of_harvesting",
    "forecast_locality",
    "forecast_locality_geonameid",
    "forecast_locality_population",
    "forecast_country",
    "forecast_country_code",
    "forecast_grid_cell_id",
    "forecast_grid_cell_index",
    "selected_model_id",
]

CLIMATE_ID_COLUMNS = [
    "trial_series_name",
    "farm",
    "plot",
    "entry_code",
]

CLIMATE_FIXED_OUTPUT_COLUMNS = [
    "planting_week_1_total_prectotcorr",
    "planting_week_2_total_prectotcorr",
    "planting_week_2_avg_t2m",
    "intermediate_period_total_prectotcorr",
    "intermediate_period_avg_t2m",
]

CLIMATE_HEADERS_WITH_UNITS = [
    "planting_week_1_total_prectotcorr (mm)",
    "planting_week_2_total_prectotcorr (mm)",
    "planting_week_2_avg_t2m (C)",
    "intermediate_period_total_prectotcorr (mm)",
    "intermediate_period_avg_t2m (C)",
    "harvest_back_week_1_total_prectotcorr (mm)",
    "harvest_back_week_1_avg_t2m (C)",
    "harvest_back_week_2_total_prectotcorr (mm)",
    "harvest_back_week_2_avg_t2m (C)",
    "harvest_back_week_3_total_prectotcorr (mm)",
    "harvest_back_week_3_avg_t2m (C)",
    "harvest_back_week_4_total_prectotcorr (mm)",
    "harvest_back_week_4_avg_t2m (C)",
    "harvest_back_week_5_total_prectotcorr (mm)",
    "harvest_back_week_5_avg_t2m (C)",
    "harvest_back_week_6_total_prectotcorr (mm)",
    "harvest_back_week_6_avg_t2m (C)",
    "harvest_back_week_7_total_prectotcorr (mm)",
    "harvest_back_week_7_avg_t2m (C)",
    "harvest_back_week_8_total_prectotcorr (mm)",
    "harvest_back_week_8_avg_t2m (C)",
]

CLIMATE_AUDIT_COLUMNS = CLIMATE_ID_COLUMNS + [
    "latitude",
    "longitude",
    "date_of_planting",
    "date_of_harvesting",
    "forecast_mode",
    "forecast_years_back",
    "forecast_reference_year",
    "forecast_reference_years",
    "climate_scope",
    "regional_country",
    "regional_tile_count",
    "regional_bounds_label",
    "regional_bounds_latitude_min",
    "regional_bounds_latitude_max",
    "regional_bounds_longitude_min",
    "regional_bounds_longitude_max",
    "forecast_locality",
    "forecast_locality_geonameid",
    "forecast_locality_population",
    "forecast_country",
    "forecast_country_code",
    "forecast_grid_cell_id",
    "forecast_grid_cell_index",
    "selected_model_id",
    "planting_week_1_start",
    "planting_week_1_end",
    "planting_week_2_start",
    "planting_week_2_end",
    "intermediate_period_start",
    "intermediate_period_end",
    "harvest_back_week_8_start",
    "harvest_back_week_1_end",
]

NAME_HEADER = "Name"
RANK_HEADER = "Rank"
FRANK_HEADER = "Frank"
YIELD_HEADER = "Grain Yield (T/Ha)"
COUNTRY_HEADER = "Country"
GPS_HEADER = "GPS coordinates"
LAT_HEADER = "_GPS coordinates_latitude"
LON_HEADER = "_GPS coordinates_longitude"
DATE_PLANTING_HEADER = "Date of planting"
DATE_THINNING_HEADER = "Date of thinning"
DATE_HARVESTING_HEADER = "Date_of_harvesting"
TRIAL_SERIES_HEADER = "Trial series name"
FORECAST_LOCALITY_HEADER = "Forecast locality"
FORECAST_LOCALITY_GEONAMEID_HEADER = "Forecast locality geonameid"
FORECAST_LOCALITY_POPULATION_HEADER = "Forecast locality population"
FORECAST_COUNTRY_HEADER = "Forecast country"
FORECAST_COUNTRY_CODE_HEADER = "Forecast country code"
FORECAST_GRID_CELL_ID_HEADER = "Forecast grid cell id"
FORECAST_GRID_CELL_INDEX_HEADER = "Forecast grid cell index"
SELECTED_MODEL_ID_HEADER = "Selected model id"
ORIGINAL_GPS_HEADER = "Original GPS coordinates"
ORIGINAL_LAT_HEADER = "Original _GPS coordinates_latitude"
ORIGINAL_LON_HEADER = "Original _GPS coordinates_longitude"

FERTILIZER_TYPE_HEADER = "Planting fertiliser/s (type) used"
FERTILIZER_RATE_HEADER = "Rate fertilizer applied (kg per total trial unit area)"
IS_FETILIZE_USED_HEADER = "Is_fetilize_Used"
TOP_DRESSING_TYPE_HEADER = "Type/name of top dressing fertilizer/s used at 4 weeks"
TOP_DRESSING_RATE_HEADER = "First topdressing -rate of fertilizer applied (kg per unit area)"
IS_FERTILIZER_APPLIED_HEADER = "is_fertilizer_applied"
TOP_DRESSING_TIMES_HEADER = "Number_of_times_top_rtilizer_was_applied"
SECOND_TOP_DRESSING_APPLICATION_HEADER = "Type_of_top_dressing_ion_2nd_application"
SECOND_TOP_DRESSING_RATE_HEADER = "Second_topdressing_ed_kg_per_unit_area"
FIRST_HERBICIDE_HEADER = "Trial_management_by_ation_1st_herbicide"
FIRST_HERBICIDE_ALT_HEADER = "Trial_management_by_ation_1st_herbicide_001"
SECOND_HERBICIDE_HEADER = "Trial_management_by_ation_2nd_herbicide"
HERBICIDE_OUTPUT_HEADER = "is_Trial_management_by_action_herbicide"
INSECTICIDE_HEADER = "Trial_management_by_ation_Insecticide_1"
INSECTICIDE_ALT_HEADER = "Trial_management_by_ation_Insecticide_1_001"
INSECTICIDE_PESTS_HEADER = "Pests_sprayed_against_1st_insecticide"
INSECTICIDE_OUTPUT_HEADER = "is_Trial_management_by_action_Insecticide"

IS_FERTILIZER_HEADER = "is_fertilizer"
FERTILIZER_TIMES_HEADER = "Number_of_times_fertilizer_was_applied"
PERCENTAGE_PLANT_HARVEST_STAND_HEADER = "Percentage_Plant_harvest_stand_t_harvesting_Plot_1"
PERCENTAGE_ROOT_LODGING_HEADER = "Percentage_Number_of_plants_wit_root_lodging_Plot_1"
PERCENTAGE_STEM_LODGING_HEADER = "Percentage_Number_of_plants_wit_stem_lodging_Plot_1"
PERCENTAGE_POOR_HUSK_COVER_HEADER = "Percentage_Poor_husk_cover_Plot_1"
PERCENTAGE_EARS_HARVESTED_HEADER = "Percentage_Number_of_ears_harvested_Plot_1"
PERCENTAGE_ROTTEN_EARS_HEADER = "Percentage_Count_number_of_rotten_ears_Plot_1"
PLOT_AREA_HEADER = "Plot Area"
BAD_HUSK_COVER_HEADER = "Bad Husk Cover (%)"
EAR_POSITION_HEADER = "Ear Position"
EARS_PER_PLANT_HEADER = "Ears/Plant (#)"
ROOT_LODGING_HEADER = "Root Lodging (%)"
STEM_LODGING_HEADER = "Stem Logding (%)"
ROTTEN_EARS_HEADER = "Rotten Ears (%)"
PLANT_HARVEST_STAND_HEADER = "Plant_harvest_stand_t_harvesting_Plot_1"
ROOT_LODGING_COUNT_HEADER = "Number_of_plants_wit_root_lodging_Plot_1"
STEM_LODGING_COUNT_HEADER = "Number_of_plants_wit_stem_lodging_Plot_1"
POOR_HUSK_COVER_COUNT_HEADER = "Poor_husk_cover_Plot_1"
EARS_HARVESTED_COUNT_HEADER = "Number_of_ears_harvested_Plot_1"
ROTTEN_EARS_COUNT_HEADER = "Count_number_of_rotten_ears_Plot_1"
PLANT_STAND_HEADER = "Plant stand ; plot 1"
EAR_HEIGHT_HEADER = "Ear_height_Plot_1"
PLANT_HEIGHT_HEADER = "Plant_height_Plot_1"
ROW_LENGTH_HEADER = "Row length (meters)"
SPACING_HEADER = "Spacing (distance between rows) (meters)"
ROWS_PER_ENTRY_HEADER = "Number of rows planted per entry"
IDPK_HEADER = "idPK"
REP_HEADER = "Rep"

PHASE06_TRACE_HEADERS = [
    "Trial series name",
    REP_HEADER,
    "Farm",
    "Site Number",
    "Plot",
    "Entry",
    "EntryCode",
    "Name",
    FORECAST_LOCALITY_HEADER,
    FORECAST_LOCALITY_GEONAMEID_HEADER,
    FORECAST_LOCALITY_POPULATION_HEADER,
    FORECAST_COUNTRY_HEADER,
    FORECAST_COUNTRY_CODE_HEADER,
    SELECTED_MODEL_ID_HEADER,
    ORIGINAL_GPS_HEADER,
    ORIGINAL_LAT_HEADER,
    ORIGINAL_LON_HEADER,
]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_upper(value: object) -> str:
    return normalize_text(value).upper()


def parse_number(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value)
    if not text or text == ".":
        return None
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, dt_time):
        return None
    if isinstance(value, (int, float)):
        base = datetime(1899, 12, 30)
        return (base + timedelta(days=float(value))).date()
    text = normalize_text(value)
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def excel_value_to_iso(value: object) -> str:
    if isinstance(value, dt_time):
        return ""
    parsed = parse_date(value)
    return parsed.isoformat() if parsed is not None else normalize_text(value)


def normalize_key_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


def write_workbook(path: Path, sheet_name: str, headers: list[str], records: list[dict[str, object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    for record in records:
        worksheet.append([record.get(header) for header in headers])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_metadata(path: Path, payload: dict[str, object]) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def load_json_file(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def save_json_file(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def path_for_metadata(path: Path) -> str:
    try:
        return str(path.relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def load_first_sheet(path: Path, sheet_index: int = 0) -> tuple[str, list[str], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[sheet_index]
        rows = worksheet.iter_rows(values_only=True)
        headers = [value for value in next(rows) if value is not None]
        records = [dict(zip(headers, row[: len(headers)])) for row in rows]
        return worksheet.title, headers, records
    finally:
        workbook.close()


def clean_phase_outputs() -> None:
    for path in PHASE_DIRS.values():
        if path.exists():
            shutil.rmtree(path)


def format_number(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.4f}".rstrip("0").rstrip(".")


def format_nasa_date(value: date) -> str:
    return value.strftime("%Y%m%d")


def parse_iso_date(value: str) -> date:
    return datetime.strptime(value.strip(), "%Y-%m-%d").date()


def normalize_for_key(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().upper())


def normalize_for_similarity(value: str) -> str:
    return re.sub(r"[^A-Z0-9 ]+", "", normalize_for_key(value))


def infer_source_set_suffix(record: dict[str, object]) -> str:
    return "01" if normalize_text(record.get(REP_HEADER)) != "" else "02"


def has_real_value(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        text = value.strip()
        if text == "" or text.lower() == "null":
            return False
    return True


def is_blank(value: object) -> bool:
    return normalize_text(value) == ""


def has_text_value(value: object) -> bool:
    text = normalize_text(value)
    return text != "" and text.lower() != "null"


def safe_median(values: list[int]) -> int | None:
    if not values:
        return None
    return int(round(statistics.median(values)))


def build_phase00_snapshot(source_workbook: Path, source_sheet_index: int) -> tuple[list[str], list[dict[str, object]]]:
    sheet_name, headers, records = load_first_sheet(source_workbook, source_sheet_index)
    output_dir = PHASE_DIRS["phase00"]
    write_workbook(output_dir / "selected_input.xlsx", "selected_input", headers, records)
    write_metadata(
        output_dir / "metadata.json",
        {
            "phase": "phase00_selected_input",
            "source_workbook": path_for_metadata(source_workbook),
            "source_sheet_index": source_sheet_index,
            "source_sheet_name": sheet_name,
            "header_count": len(headers),
            "row_count": len(records),
            "headers": headers,
        },
    )
    return headers, records


def is_geographically_plausible(
    record: dict[str, object],
    latitude: float,
    longitude: float,
) -> bool:
    bounds = COUNTRY_BOUNDS_BY_SOURCE_SET[infer_source_set_suffix(record)].get(
        normalize_upper(record.get(COUNTRY_HEADER))
    )
    if bounds is None:
        return False
    min_lat, max_lat, min_lon, max_lon = bounds
    return min_lat <= latitude <= max_lat and min_lon <= longitude <= max_lon


def is_valid_coordinate_pair(record: dict[str, object]) -> bool:
    if not has_text_value(record.get(GPS_HEADER)):
        return False
    latitude = parse_number(record.get(LAT_HEADER))
    longitude = parse_number(record.get(LON_HEADER))
    if latitude is None or longitude is None:
        return False
    if not (-90 <= latitude <= 90):
        return False
    if not (-180 <= longitude <= 180):
        return False
    if latitude == 0 and longitude == 0:
        return False
    if not is_geographically_plausible(record, latitude, longitude):
        return False
    return True


def is_placeholder_thinning(value: object) -> bool:
    if isinstance(value, dt_time):
        return value == dt_time(0, 0)
    parsed = parse_date(value)
    if parsed is not None:
        return parsed <= date(1900, 1, 1)
    return normalize_text(value) in {"0/1/1900", "01/01/1900", "1/1/1900", "1900-01-01"}


def apply_name_rules(record: dict[str, object]) -> tuple[dict[str, object], str | None]:
    updated = dict(record)
    raw_name = updated.get(NAME_HEADER)
    if raw_name is None:
        return updated, None
    normalized = normalize_upper(raw_name)
    replacement = NAME_REPLACEMENTS.get(normalized, normalized)
    if replacement == "DELETE":
        return updated, "deleted_name_rule"
    updated[NAME_HEADER] = replacement
    return updated, replacement if replacement != normalize_text(raw_name) else None


def build_name_similarity_rows(records: list[dict[str, object]]) -> list[dict[str, object]]:
    unique_values: list[str] = []
    seen = set()
    for record in records:
        value = normalize_text(record.get(NAME_HEADER))
        if not value:
            continue
        key = normalize_for_key(value)
        if key in seen:
            continue
        seen.add(key)
        unique_values.append(value)

    normalized = [normalize_for_similarity(value) for value in unique_values]
    rows: list[dict[str, object]] = []
    for left in range(len(unique_values)):
        for right in range(left + 1, len(unique_values)):
            left_norm = normalized[left]
            right_norm = normalized[right]
            if not left_norm or not right_norm or left_norm == right_norm:
                continue
            ratio = SequenceMatcher(None, left_norm, right_norm).ratio()
            min_len = min(len(left_norm), len(right_norm))
            contained = left_norm in right_norm or right_norm in left_norm
            if ratio >= SIMILARITY_THRESHOLD and min_len >= 4:
                rows.append(
                    {
                        "Value 1": unique_values[left],
                        "Value 2": unique_values[right],
                        "Score": ratio,
                        "Rule": "ratio",
                    }
                )
            elif contained and min_len >= 5:
                rows.append(
                    {
                        "Value 1": unique_values[left],
                        "Value 2": unique_values[right],
                        "Score": ratio,
                        "Rule": "contains",
                    }
                )
    rows.sort(key=lambda item: float(item["Score"]), reverse=True)
    return rows


def build_plot_sort_key(plotid: object) -> tuple[int, object]:
    text = normalize_text(plotid)
    if text.lower().startswith("plot "):
        numeric_plotid = parse_number(text[5:])
        if numeric_plotid is not None:
            return (0, numeric_plotid)
    numeric_plotid = parse_number(plotid)
    if numeric_plotid is not None:
        return (0, numeric_plotid)
    return (1, text.lower())


def build_rank_list(row_values: tuple[object, object, object, object]) -> list[tuple[int, object]]:
    if not all(normalize_text(value) != "" for value in row_values):
        return []
    ranked_plots = [
        (1, row_values[0]),
        (2, row_values[1]),
        (3, row_values[2]),
        (4, row_values[3]),
    ]
    ranked_plots.sort(key=lambda item: build_plot_sort_key(item[1]))
    return ranked_plots


def build_joint_rank_values(
    list_rank: list[tuple[int, object]],
    list_frank: list[tuple[int, object]],
) -> list[int]:
    plot_scores: dict[object, list[int]] = {}
    plot_order: dict[object, tuple[int, object]] = {}
    for grade, plotid in list_rank:
        plot_scores.setdefault(plotid, []).append(grade)
        plot_order[plotid] = build_plot_sort_key(plotid)
    for grade, plotid in list_frank:
        plot_scores.setdefault(plotid, []).append(grade)
        plot_order[plotid] = build_plot_sort_key(plotid)
    ranked_by_average = sorted(
        plot_scores.items(),
        key=lambda item: (sum(item[1]) / len(item[1]), plot_order[item[0]]),
    )
    combined_grade_by_plotid = {
        plotid: new_grade for new_grade, (plotid, _) in enumerate(ranked_by_average, start=1)
    }
    plotids_sorted_for_rows = sorted(plot_scores, key=lambda plotid: plot_order[plotid])
    return [combined_grade_by_plotid[plotid] for plotid in plotids_sorted_for_rows]


def select_rank_values(
    gender_value: str,
    list_rank: list[tuple[int, object]],
    list_frank: list[tuple[int, object]],
) -> list[int]:
    if gender_value == "male":
        if list_rank:
            return [item[0] for item in list_rank]
        if list_frank:
            return [item[0] for item in list_frank]
    if gender_value in {"female", "famale"}:
        if list_frank:
            return [item[0] for item in list_frank]
        if list_rank:
            return [item[0] for item in list_rank]
    if gender_value == "joint":
        if list_rank and list_frank:
            return build_joint_rank_values(list_rank, list_frank)
        if list_rank:
            return [item[0] for item in list_rank]
        if list_frank:
            return [item[0] for item in list_frank]
    return []


def build_direct_rank_values(
    row_values: tuple[object, object, object, object],
) -> list[int]:
    ranked_plots = build_rank_list(row_values)
    return [item[0] for item in ranked_plots]


def derive_rank_and_frank_block(records: list[dict[str, object]]) -> list[dict[str, object]]:
    if not records:
        return []

    first_record = records[0]
    primary_values = tuple(first_record.get(header) for header in ("rank1", "rank2", "rank3", "rank4"))
    secondary_values = tuple(first_record.get(header) for header in ("Frank1", "Frank2", "Frank3", "Frank4"))
    list_rank = build_rank_list(primary_values)
    list_frank = build_rank_list(secondary_values)
    gender_value = normalize_text(first_record.get("Gender of the farmer (plot manager)")).lower()
    rank_values = select_rank_values(gender_value, list_rank, list_frank)
    frank_values = build_direct_rank_values(secondary_values)

    ranked_records: list[dict[str, object]] = []
    for offset, record in enumerate(records):
        updated = dict(record)
        updated[RANK_HEADER] = rank_values[offset] if offset < len(rank_values) else None
        updated[FRANK_HEADER] = frank_values[offset] if offset < len(frank_values) else None
        ranked_records.append(updated)
    return ranked_records


def derive_ranked_records(records: list[dict[str, object]]) -> list[dict[str, object]]:
    ranked_records: list[dict[str, object]] = []
    buffered_records: list[dict[str, object]] = []
    valid_genders = {"male", "female", "famale", "joint"}

    for record in records:
        buffered_records.append(record)

        while buffered_records:
            gender_value = normalize_text(
                buffered_records[0].get("Gender of the farmer (plot manager)")
            ).lower()
            if gender_value not in valid_genders:
                updated = dict(buffered_records.pop(0))
                updated[RANK_HEADER] = None
                updated[FRANK_HEADER] = None
                ranked_records.append(updated)
                continue

            if len(buffered_records) < 4:
                break

            ranked_records.extend(derive_rank_and_frank_block(buffered_records[:4]))
            del buffered_records[:4]

    while buffered_records:
        gender_value = normalize_text(
            buffered_records[0].get("Gender of the farmer (plot manager)")
        ).lower()
        if gender_value in valid_genders:
            block_size = min(4, len(buffered_records))
            ranked_records.extend(derive_rank_and_frank_block(buffered_records[:block_size]))
            del buffered_records[:block_size]
        else:
            updated = dict(buffered_records.pop(0))
            updated[RANK_HEADER] = None
            updated[FRANK_HEADER] = None
            ranked_records.append(updated)

    return ranked_records


def build_phase01_headers(headers: list[str]) -> list[str]:
    output_headers = list(headers)
    insert_after = "Gender of the farmer (plot manager)"
    if RANK_HEADER not in output_headers:
        output_headers.insert(output_headers.index(insert_after) + 1, RANK_HEADER)
    if FRANK_HEADER not in output_headers:
        output_headers.insert(output_headers.index(RANK_HEADER) + 1, FRANK_HEADER)
    return output_headers


ARCHIVE_FORWARD_FILL_HEADERS = [
    "Type/name of top dressing fertilizer/s used at 4 weeks",
    "First topdressing -rate of fertilizer applied (kg per unit area)",
    "Weed control practices",
    "Number_of_times_top_rtilizer_was_applied",
    "Type_of_top_dressing_ion_2nd_application",
    "Second_topdressing_ed_kg_per_unit_area",
    "Trial_management_by_ation_1st_herbicide",
    "Trial_management_by_ation_1st_herbicide_001",
    "Trial_management_by_ation_2nd_herbicide",
    "Trial_management_by_ation_Insecticide_1",
    "Trial_management_by_ation_Insecticide_1_001",
    "Pests_sprayed_against_1st_insecticide",
    "_8_1_BM_Who_in_your_h_e_on_farm_experiment",
    "_8_2_BM_Who_in_your_h_e_on_farm_experiment",
]


def normalize_archive_forward_fill_value(value: object) -> object:
    if isinstance(value, str) and value.strip().upper() == "NONE":
        return None
    return value


def has_archive_forward_fill_value(value: object) -> bool:
    return normalize_text(value) != ""


def apply_archive_forward_fill_to_block(records: list[dict[str, object]]) -> list[dict[str, object]]:
    updated_records = [dict(record) for record in records]
    for header in ARCHIVE_FORWARD_FILL_HEADERS:
        last_value: object = None
        remaining_rows = 0
        for record in updated_records:
            value = normalize_archive_forward_fill_value(record.get(header))
            record[header] = value
            if has_archive_forward_fill_value(value):
                last_value = value
                remaining_rows = 3
                continue
            if has_archive_forward_fill_value(last_value) and remaining_rows > 0:
                record[header] = last_value
                remaining_rows -= 1
    return updated_records


def apply_archive_management_forward_fill(records: list[dict[str, object]]) -> list[dict[str, object]]:
    updated_records: list[dict[str, object]] = []
    current_2025_block: list[dict[str, object]] = []
    current_block_key: tuple[str, str, str] | None = None

    def flush_current_block() -> None:
        nonlocal current_2025_block, current_block_key
        if not current_2025_block:
            return
        updated_records.extend(apply_archive_forward_fill_to_block(current_2025_block))
        current_2025_block = []
        current_block_key = None

    def build_2025_block_key(record: dict[str, object]) -> tuple[str, str, str]:
        return (
            normalize_key_value(record.get(TRIAL_SERIES_HEADER)),
            normalize_key_value(record.get("Rep")),
            normalize_key_value(record.get("Site Number")),
        )

    for record in records:
        # `Rep` is present in 2025 merged rows and absent in the 2024 workbook.
        if has_archive_forward_fill_value(record.get("Rep")):
            block_key = build_2025_block_key(record)
            if current_block_key is not None and block_key != current_block_key:
                flush_current_block()
            current_block_key = block_key
            current_2025_block.append(record)
            if len(current_2025_block) == 4:
                flush_current_block()
            continue
        flush_current_block()
        updated_records.append(dict(record))

    flush_current_block()
    return updated_records


def write_names_audit(output_dir: Path, records: list[dict[str, object]]) -> None:
    unique_rows: list[dict[str, object]] = []
    seen_sets: dict[str, set[str]] = {}
    unique_map: dict[str, str] = {}
    for record in records:
        value = normalize_text(record.get(NAME_HEADER))
        if not value:
            continue
        key = normalize_for_key(value)
        unique_map.setdefault(key, value)
        seen_sets.setdefault(key, set()).add(infer_source_set_suffix(record))

    for idx, (key, value) in enumerate(unique_map.items(), start=1):
        unique_rows.append(
            {
                "ID": idx,
                "Value": value,
                "Normalized key": key,
                "Seen in source sets": ", ".join(sorted(seen_sets[key])),
            }
        )

    workbook = Workbook()
    unique_ws = workbook.active
    unique_ws.title = "unique_names"
    unique_headers = ["ID", "Value", "Normalized key", "Seen in source sets"]
    unique_ws.append(unique_headers)
    for row in unique_rows:
        unique_ws.append([row[header] for header in unique_headers])

    similar_ws = workbook.create_sheet("similar_candidates")
    similar_headers = ["Value 1", "Value 2", "Score", "Rule"]
    similar_ws.append(similar_headers)
    for row in build_name_similarity_rows(records):
        similar_ws.append([row[header] for header in similar_headers])

    workbook.save(output_dir / "names_audit.xlsx")


def process_pre_climate_records(
    records: list[dict[str, object]],
    *,
    allow_missing_yield: bool = False,
    allow_missing_coordinates: bool = False,
) -> tuple[list[dict[str, object]], dict[str, list[dict[str, object]]], dict[str, object]]:
    ranked_records = derive_ranked_records(records)
    prepared_records: list[dict[str, object]] = []
    excluded_missing_yield: list[dict[str, object]] = []
    excluded_rank_yield_zero: list[dict[str, object]] = []
    deleted_by_name_rule: list[dict[str, object]] = []
    wrong_coordinates: list[dict[str, object]] = []

    name_updates = 0
    for ranked_record in ranked_records:
        named_record, name_result = apply_name_rules(ranked_record)
        if name_result == "deleted_name_rule":
            deleted_by_name_rule.append(named_record)
            continue
        if name_result is not None:
            name_updates += 1

        yield_value = named_record.get(YIELD_HEADER)
        numeric_yield = parse_number(yield_value)
        if not allow_missing_yield and not has_real_value(yield_value):
            excluded_missing_yield.append(named_record)
            continue
        if (
            not allow_missing_yield
            and is_blank(named_record.get(RANK_HEADER))
            and numeric_yield == 0
        ):
            excluded_rank_yield_zero.append(named_record)
            continue
        if not allow_missing_coordinates and not is_valid_coordinate_pair(named_record):
            wrong_coordinates.append(named_record)
            continue
        prepared_records.append(named_record)

    metadata = {
        "input_row_count": len(records),
        "output_row_count": len(prepared_records),
        "name_updates": name_updates,
        "deleted_by_name_rule": len(deleted_by_name_rule),
        "excluded_missing_yield": len(excluded_missing_yield),
        "excluded_rank_yield_zero": len(excluded_rank_yield_zero),
        "wrong_coordinates": len(wrong_coordinates),
    }
    audits = {
        "deleted_name_rule": deleted_by_name_rule,
        "excluded_missing_yield": excluded_missing_yield,
        "excluded_rank_yield_zero": excluded_rank_yield_zero,
        "wrong_coordinates": wrong_coordinates,
    }
    return prepared_records, audits, metadata


def impute_date_of_thinning(records: list[dict[str, object]]) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, int]]:
    deltas_by_trial_country: dict[tuple[str, str], list[int]] = {}
    deltas_by_trial: dict[str, list[int]] = {}
    deltas_by_country: dict[str, list[int]] = {}
    global_deltas: list[int] = []

    for record in records:
        planting_date = parse_date(record.get(DATE_PLANTING_HEADER))
        thinning_date = parse_date(record.get(DATE_THINNING_HEADER))
        if planting_date is None or thinning_date is None:
            continue
        if is_placeholder_thinning(record.get(DATE_THINNING_HEADER)):
            continue
        delta_days = (thinning_date - planting_date).days
        if delta_days < 0:
            continue
        trial = normalize_upper(record.get(TRIAL_SERIES_HEADER))
        country = normalize_upper(record.get(COUNTRY_HEADER))
        deltas_by_trial_country.setdefault((trial, country), []).append(delta_days)
        deltas_by_trial.setdefault(trial, []).append(delta_days)
        deltas_by_country.setdefault(country, []).append(delta_days)
        global_deltas.append(delta_days)

    stats = {
        "imputed": 0,
        "pending": 0,
        "group_trial_country": 0,
        "group_trial": 0,
        "group_country": 0,
        "group_global": 0,
    }
    audit_rows: list[dict[str, object]] = []
    updated_records: list[dict[str, object]] = []

    for row_number, record in enumerate(records, start=2):
        updated = dict(record)
        if not is_placeholder_thinning(updated.get(DATE_THINNING_HEADER)):
            updated_records.append(updated)
            continue

        planting_date = parse_date(updated.get(DATE_PLANTING_HEADER))
        if planting_date is None:
            stats["pending"] += 1
            audit_rows.append(
                {
                    "Row": row_number,
                    "Trial series name": updated.get(TRIAL_SERIES_HEADER),
                    "Country": updated.get(COUNTRY_HEADER),
                    "Date of planting": updated.get(DATE_PLANTING_HEADER),
                    "Date of thinning original": updated.get(DATE_THINNING_HEADER),
                    "Delta days used": None,
                    "Method used": "pending_no_planting_date",
                    "Date of thinning imputed": None,
                }
            )
            updated_records.append(updated)
            continue

        trial = normalize_upper(updated.get(TRIAL_SERIES_HEADER))
        country = normalize_upper(updated.get(COUNTRY_HEADER))
        delta = (
            safe_median(deltas_by_trial_country.get((trial, country), []))
            or safe_median(deltas_by_trial.get(trial, []))
            or safe_median(deltas_by_country.get(country, []))
            or safe_median(global_deltas)
        )
        if delta is None:
            stats["pending"] += 1
            audit_rows.append(
                {
                    "Row": row_number,
                    "Trial series name": updated.get(TRIAL_SERIES_HEADER),
                    "Country": updated.get(COUNTRY_HEADER),
                    "Date of planting": updated.get(DATE_PLANTING_HEADER),
                    "Date of thinning original": updated.get(DATE_THINNING_HEADER),
                    "Delta days used": None,
                    "Method used": "pending_no_reference_group",
                    "Date of thinning imputed": None,
                }
            )
            updated_records.append(updated)
            continue

        method_used = "global"
        if deltas_by_trial_country.get((trial, country)):
            stats["group_trial_country"] += 1
            method_used = "trial+country"
        elif deltas_by_trial.get(trial):
            stats["group_trial"] += 1
            method_used = "trial"
        elif deltas_by_country.get(country):
            stats["group_country"] += 1
            method_used = "country"
        else:
            stats["group_global"] += 1

        updated[DATE_THINNING_HEADER] = planting_date + timedelta(days=delta)
        stats["imputed"] += 1
        audit_rows.append(
            {
                "Row": row_number,
                "Trial series name": updated.get(TRIAL_SERIES_HEADER),
                "Country": updated.get(COUNTRY_HEADER),
                "Date of planting": updated.get(DATE_PLANTING_HEADER),
                "Date of thinning original": record.get(DATE_THINNING_HEADER),
                "Delta days used": delta,
                "Method used": method_used,
                "Date of thinning imputed": updated.get(DATE_THINNING_HEADER),
            }
        )
        updated_records.append(updated)

    return updated_records, audit_rows, stats


def build_phase01_prepared(
    headers: list[str],
    records: list[dict[str, object]],
    *,
    allow_missing_yield: bool = False,
    allow_missing_coordinates: bool = False,
) -> tuple[list[str], list[dict[str, object]], dict[str, object]]:
    output_dir = PHASE_DIRS["phase01"]
    phase01_headers = build_phase01_headers(headers)
    records = apply_archive_management_forward_fill(records)
    prepared_records, audits, quality_stats = process_pre_climate_records(
        records,
        allow_missing_yield=allow_missing_yield,
        allow_missing_coordinates=allow_missing_coordinates,
    )
    prepared_records, thinning_audit_rows, thinning_stats = impute_date_of_thinning(prepared_records)

    write_workbook(
        output_dir / "phase01_quality_prepared.xlsx",
        "phase01_prepared",
        phase01_headers,
        prepared_records,
    )

    for slug, audit_records in audits.items():
        if audit_records:
            write_workbook(output_dir / f"{slug}.xlsx", slug[:31], phase01_headers, audit_records)

    write_names_audit(output_dir, prepared_records)

    audit_headers = [
        "Row",
        "Trial series name",
        "Country",
        "Date of planting",
        "Date of thinning original",
        "Delta days used",
        "Method used",
        "Date of thinning imputed",
    ]
    write_workbook(
        output_dir / "date_thinning_audit.xlsx",
        "date_thinning_audit",
        audit_headers,
        thinning_audit_rows,
    )

    metadata = {
        "phase": "phase01_quality_prepared",
        **quality_stats,
        **{f"thinning_{key}": value for key, value in thinning_stats.items()},
    }
    write_metadata(output_dir / "metadata.json", metadata)
    return phase01_headers, prepared_records, metadata


def build_climate_input_row(record: dict[str, object]) -> dict[str, str]:
    return {
        "trial_series_name": normalize_key_value(record.get(TRIAL_SERIES_HEADER)),
        "farm": normalize_key_value(record.get("Site Number")),
        "plot": normalize_key_value(record.get("Plot")),
        "entry_code": normalize_key_value(record.get("EntryCode")),
        "name": normalize_text(record.get(NAME_HEADER)),
        "country": normalize_text(record.get(COUNTRY_HEADER)),
        "latitude": normalize_key_value(record.get(LAT_HEADER)),
        "longitude": normalize_key_value(record.get(LON_HEADER)),
        "date_of_planting": excel_value_to_iso(record.get(DATE_PLANTING_HEADER)),
        "date_of_thinning": excel_value_to_iso(record.get(DATE_THINNING_HEADER)),
        "date_of_harvesting": excel_value_to_iso(record.get(DATE_HARVESTING_HEADER)),
        "forecast_locality": normalize_text(record.get(FORECAST_LOCALITY_HEADER)),
        "forecast_locality_geonameid": normalize_key_value(record.get(FORECAST_LOCALITY_GEONAMEID_HEADER)),
        "forecast_locality_population": normalize_key_value(record.get(FORECAST_LOCALITY_POPULATION_HEADER)),
        "forecast_country": normalize_text(record.get(FORECAST_COUNTRY_HEADER)),
        "forecast_country_code": normalize_key_value(record.get(FORECAST_COUNTRY_CODE_HEADER)),
        "forecast_grid_cell_id": normalize_key_value(record.get(FORECAST_GRID_CELL_ID_HEADER)),
        "forecast_grid_cell_index": normalize_key_value(record.get(FORECAST_GRID_CELL_INDEX_HEADER)),
        "selected_model_id": normalize_key_value(record.get(SELECTED_MODEL_ID_HEADER)),
    }


def build_data_input_rows(records: list[dict[str, object]]) -> tuple[list[dict[str, str]], int]:
    output_rows: list[dict[str, str]] = []
    skipped_rows = 0
    for record in records:
        output_row = build_climate_input_row(record)
        required = (
            output_row["latitude"],
            output_row["longitude"],
            output_row["date_of_planting"],
            output_row["date_of_harvesting"],
        )
        if any(not value for value in required):
            skipped_rows += 1
            continue
        output_rows.append(output_row)
    return output_rows, skipped_rows


def build_country_locality_phase02_headers(headers: list[str], target_header: str = YIELD_HEADER) -> list[str]:
    output_headers = build_phase02_headers(headers, target_header=target_header)
    extra_headers = [
        FORECAST_LOCALITY_HEADER,
        FORECAST_LOCALITY_GEONAMEID_HEADER,
        FORECAST_LOCALITY_POPULATION_HEADER,
        FORECAST_COUNTRY_HEADER,
        FORECAST_COUNTRY_CODE_HEADER,
        SELECTED_MODEL_ID_HEADER,
        ORIGINAL_GPS_HEADER,
        ORIGINAL_LAT_HEADER,
        ORIGINAL_LON_HEADER,
    ]
    for header in extra_headers:
        if header not in output_headers:
            output_headers.append(header)
    return output_headers


def collapse_records_for_locality_projection(records: list[dict[str, object]]) -> list[dict[str, object]]:
    collapsed_records: list[dict[str, object]] = []
    seen_projection_keys: set[str] = set()
    for record in records:
        explicit_profile_key = normalize_text(record.get("Prediction Profile Key"))
        if explicit_profile_key:
            projection_key = explicit_profile_key
        else:
            normalized_name = normalize_upper(record.get(NAME_HEADER))
            projection_key = NAME_REPLACEMENTS.get(normalized_name, normalized_name)
        if projection_key:
            if projection_key in seen_projection_keys:
                continue
            seen_projection_keys.add(projection_key)
        collapsed_records.append(record)
    return collapsed_records


def expand_records_for_country_localities(
    records: list[dict[str, object]],
    climate_override: dict[str, object],
) -> tuple[list[dict[str, object]], list[dict[str, str]], dict[str, object]]:
    selected_country = resolve_african_country(str(climate_override.get("regional_country") or ""))
    localities = load_country_localities(selected_country.iso2)
    if not localities:
        raise ValueError(f"No localities were found for {selected_country.country}.")

    forecast_planting_date = str(climate_override.get("forecast_planting_date") or "").strip()
    forecast_harvesting_date = str(climate_override.get("forecast_harvesting_date") or "").strip()
    selected_model_id = str(climate_override.get("selected_model_id") or "").strip()
    use_source_row_dates = bool(climate_override.get("use_source_row_dates"))

    expanded_records: list[dict[str, object]] = []
    data_input_rows: list[dict[str, str]] = []
    projected_records = records if use_source_row_dates else collapse_records_for_locality_projection(records)
    for record in projected_records:
        for locality in localities:
            expanded = dict(record)
            expanded[ORIGINAL_GPS_HEADER] = record.get(GPS_HEADER, "")
            expanded[ORIGINAL_LAT_HEADER] = record.get(LAT_HEADER, "")
            expanded[ORIGINAL_LON_HEADER] = record.get(LON_HEADER, "")
            expanded[COUNTRY_HEADER] = selected_country.country
            expanded[GPS_HEADER] = f"{locality['latitude']} {locality['longitude']}"
            expanded[LAT_HEADER] = locality["latitude"]
            expanded[LON_HEADER] = locality["longitude"]
            expanded["_GPS coordinates_altitude"] = ""
            expanded["_GPS coordinates_precision"] = ""
            if forecast_planting_date:
                expanded[DATE_PLANTING_HEADER] = forecast_planting_date
            if forecast_harvesting_date:
                expanded[DATE_HARVESTING_HEADER] = forecast_harvesting_date
            expanded[FORECAST_LOCALITY_HEADER] = locality["name"]
            expanded[FORECAST_LOCALITY_GEONAMEID_HEADER] = locality["geoname_id"]
            expanded[FORECAST_LOCALITY_POPULATION_HEADER] = locality["population"]
            expanded[FORECAST_COUNTRY_HEADER] = selected_country.country
            expanded[FORECAST_COUNTRY_CODE_HEADER] = selected_country.iso2
            expanded[SELECTED_MODEL_ID_HEADER] = selected_model_id
            expanded_records.append(expanded)
            data_input_rows.append(build_climate_input_row(expanded))

    return expanded_records, data_input_rows, {
        "forecast_country": selected_country.country,
        "forecast_country_code": selected_country.iso2,
        "locality_count": len(localities),
        "projected_germplasm_count": len(projected_records),
        "source_row_dates_preserved": use_source_row_dates,
        "expanded_row_count": len(expanded_records),
        "selected_model_id": selected_model_id,
    }


def expand_records_for_manual_bbox_localities(
    records: list[dict[str, object]],
    climate_override: dict[str, object],
) -> tuple[list[dict[str, object]], list[dict[str, str]], dict[str, object]]:
    try:
        latitude_min = float(climate_override.get("regional_bounds_latitude_min"))
        latitude_max = float(climate_override.get("regional_bounds_latitude_max"))
        longitude_min = float(climate_override.get("regional_bounds_longitude_min"))
        longitude_max = float(climate_override.get("regional_bounds_longitude_max"))
    except (TypeError, ValueError) as error:
        raise ValueError("Manual bounding-box localities mode requires valid latitude and longitude bounds.") from error

    latitude_min, latitude_max = min(latitude_min, latitude_max), max(latitude_min, latitude_max)
    longitude_min, longitude_max = min(longitude_min, longitude_max), max(longitude_min, longitude_max)

    grid_resolution_km = float(climate_override.get("nasa_grid_resolution_km") or 10)
    grid_cells = build_manual_grid_cells(
        latitude_min,
        latitude_max,
        longitude_min,
        longitude_max,
        label=str(climate_override.get("regional_bounds_label") or "Manual bounds").strip() or "Manual bounds",
        cell_size_km=grid_resolution_km,
    )
    if not grid_cells:
        raise ValueError(
            "No NASA grid cells could be generated inside the selected manual bounding box "
            f"(lat {latitude_min} to {latitude_max}, lon {longitude_min} to {longitude_max})."
        )

    forecast_planting_date = str(climate_override.get("forecast_planting_date") or "").strip()
    forecast_harvesting_date = str(climate_override.get("forecast_harvesting_date") or "").strip()
    selected_model_id = str(climate_override.get("selected_model_id") or "").strip()
    use_source_row_dates = bool(climate_override.get("use_source_row_dates"))

    expanded_records: list[dict[str, object]] = []
    data_input_rows: list[dict[str, str]] = []
    projected_records = records if use_source_row_dates else collapse_records_for_locality_projection(records)
    for record in projected_records:
        for grid_cell in grid_cells:
            expanded = dict(record)
            expanded[ORIGINAL_GPS_HEADER] = record.get(GPS_HEADER, "")
            expanded[ORIGINAL_LAT_HEADER] = record.get(LAT_HEADER, "")
            expanded[ORIGINAL_LON_HEADER] = record.get(LON_HEADER, "")
            expanded[COUNTRY_HEADER] = str(grid_cell["label"])
            expanded[GPS_HEADER] = (
                f"{grid_cell['center_latitude']} {grid_cell['center_longitude']}"
            )
            expanded[LAT_HEADER] = grid_cell["center_latitude"]
            expanded[LON_HEADER] = grid_cell["center_longitude"]
            expanded["_GPS coordinates_altitude"] = ""
            expanded["_GPS coordinates_precision"] = ""
            if forecast_planting_date:
                expanded[DATE_PLANTING_HEADER] = forecast_planting_date
            if forecast_harvesting_date:
                expanded[DATE_HARVESTING_HEADER] = forecast_harvesting_date
            expanded[FORECAST_LOCALITY_HEADER] = str(grid_cell["grid_cell_id"])
            expanded[FORECAST_LOCALITY_GEONAMEID_HEADER] = ""
            expanded[FORECAST_LOCALITY_POPULATION_HEADER] = ""
            expanded[FORECAST_COUNTRY_HEADER] = str(grid_cell["label"])
            expanded[FORECAST_COUNTRY_CODE_HEADER] = ""
            expanded[FORECAST_GRID_CELL_ID_HEADER] = str(grid_cell["grid_cell_id"])
            expanded[FORECAST_GRID_CELL_INDEX_HEADER] = grid_cell["grid_cell_index"]
            expanded[SELECTED_MODEL_ID_HEADER] = selected_model_id
            expanded_records.append(expanded)
            data_input_rows.append(build_climate_input_row(expanded))

    return expanded_records, data_input_rows, {
        "manual_bbox_locality_count": len(grid_cells),
        "manual_bbox_grid_cell_count": len(grid_cells),
        "manual_bbox_grid": {
            "label": str(climate_override.get("regional_bounds_label") or "Manual bounds").strip()
            or "Manual bounds",
            "grid_resolution_km": grid_resolution_km,
            "grid_cell_count": len(grid_cells),
            "cells": grid_cells,
        },
        "projected_germplasm_count": len(projected_records),
        "source_row_dates_preserved": use_source_row_dates,
        "expanded_row_count": len(expanded_records),
        "selected_model_id": selected_model_id,
        "regional_bounds_latitude_min": latitude_min,
        "regional_bounds_latitude_max": latitude_max,
        "regional_bounds_longitude_min": longitude_min,
        "regional_bounds_longitude_max": longitude_max,
    }


def build_api_url(latitude: str, longitude: str, start_date: date, end_date: date) -> str:
    query = urlencode(
        {
            "parameters": ",".join(NASA_PARAMETERS),
            "community": "AG",
            "longitude": longitude,
            "latitude": latitude,
            "start": format_nasa_date(start_date),
            "end": format_nasa_date(end_date),
            "format": "JSON",
        }
    )
    return f"{NASA_POWER_URL}?{query}"


def build_regional_api_url(
    latitude_min: float,
    latitude_max: float,
    longitude_min: float,
    longitude_max: float,
    parameter: str,
    start_date: date,
    end_date: date,
) -> str:
    query = urlencode(
        {
            "parameters": parameter,
            "community": "AG",
            "latitude-min": latitude_min,
            "latitude-max": latitude_max,
            "longitude-min": longitude_min,
            "longitude-max": longitude_max,
            "start": format_nasa_date(start_date),
            "end": format_nasa_date(end_date),
            "format": "JSON",
        }
    )
    return f"{NASA_POWER_REGIONAL_URL}?{query}"


def normalize_nasa_value(value: Any) -> float | None:
    if value is None:
        return None
    numeric = float(value)
    if numeric == NASA_FILL_VALUE:
        return None
    return numeric


def load_nasa_cache(
    path: Path,
) -> dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]]:
    raw_payload = load_json_file(path, default={})
    cache: dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]] = {}
    if not isinstance(raw_payload, dict):
        return cache
    for raw_key, value in raw_payload.items():
        if not isinstance(raw_key, str):
            continue
        parts = raw_key.split("|")
        if len(parts) != 4:
            continue
        cache[(parts[0], parts[1], parts[2], parts[3])] = value
    return cache


def save_nasa_cache(
    path: Path,
    cache: dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]],
) -> None:
    serialized = {"|".join(key): value for key, value in sorted(cache.items())}
    save_json_file(path, serialized)


def parse_retry_after_seconds(error: HTTPError) -> float | None:
    retry_after = error.headers.get("Retry-After")
    if retry_after is None:
        return None
    retry_after = retry_after.strip()
    if not retry_after:
        return None
    try:
        return max(float(retry_after), 0.0)
    except ValueError:
        return None


def fetch_nasa_series(
    latitude: str,
    longitude: str,
    start_date: date,
    end_date: date,
    cache: dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]],
    cache_path: Path | None = None,
    stats: dict[str, int] | None = None,
) -> dict[str, dict[str, float | None]]:
    cache_key = (latitude, longitude, start_date.isoformat(), end_date.isoformat())
    if cache_key in cache:
        if stats is not None:
            stats["cache_hits"] = stats.get("cache_hits", 0) + 1
        return cache[cache_key]

    url = build_api_url(latitude, longitude, start_date, end_date)
    request = Request(url, headers={"User-Agent": "app preprocess pipeline"})
    last_error: Exception | None = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            with urlopen(request, timeout=HTTP_TIMEOUT_SECONDS) as response:
                payload = json.load(response)
            parameter_block = payload["properties"]["parameter"]
            result = {
                parameter: {
                    key: normalize_nasa_value(value)
                    for key, value in parameter_block.get(parameter, {}).items()
                }
                for parameter in NASA_PARAMETERS
            }
            cache[cache_key] = result
            if stats is not None:
                stats["cache_misses"] = stats.get("cache_misses", 0) + 1
            if cache_path is not None:
                save_nasa_cache(cache_path, cache)
            return result
        except (HTTPError, URLError, TimeoutError, KeyError, json.JSONDecodeError) as exc:
            last_error = exc
            if attempt == MAX_RETRIES:
                break
            sleep_seconds = min(RETRY_SLEEP_SECONDS * (2 ** (attempt - 1)), MAX_RETRY_SLEEP_SECONDS)
            if isinstance(exc, HTTPError) and exc.code == 429:
                retry_after = parse_retry_after_seconds(exc)
                if retry_after is not None:
                    sleep_seconds = max(sleep_seconds, retry_after)
            time.sleep(sleep_seconds)
    raise RuntimeError(
        "No se pudo consultar NASA POWER para "
        f"lat={latitude}, lon={longitude}, start={start_date}, end={end_date}"
    ) from last_error


def aggregate_regional_payload(
    payload: dict[str, object],
    parameter: str,
) -> dict[str, float | None]:
    features = payload.get("features", [])
    if not isinstance(features, list):
        return {}
    by_date: dict[str, list[float]] = {}
    for feature in features:
        if not isinstance(feature, dict):
            continue
        properties = feature.get("properties", {})
        if not isinstance(properties, dict):
            continue
        parameter_block = properties.get("parameter", {})
        if not isinstance(parameter_block, dict):
            continue
        series = parameter_block.get(parameter, {})
        if not isinstance(series, dict):
            continue
        for day_key, raw_value in series.items():
            value = normalize_nasa_value(raw_value)
            if value is None:
                continue
            by_date.setdefault(day_key, []).append(value)
    return {
        day_key: (sum(values) / len(values)) if values else None
        for day_key, values in by_date.items()
    }


def fetch_nasa_regional_series(
    region_label: str,
    tile_bounds: list[tuple[float, float, float, float]],
    start_date: date,
    end_date: date,
    cache: dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]],
    cache_path: Path | None = None,
    stats: dict[str, int] | None = None,
) -> tuple[dict[str, dict[str, float | None]], int]:
    combined = {parameter: {} for parameter in NASA_PARAMETERS}
    tile_count = 0
    normalized_region = normalize_country_key(region_label)
    for latitude_min, latitude_max, longitude_min, longitude_max in tile_bounds:
        tile_count += 1
        for parameter in NASA_PARAMETERS:
            cache_key = (
                f"regional:{normalized_region}:{parameter}",
                f"{latitude_min}:{latitude_max}",
                f"{longitude_min}:{longitude_max}",
                f"{start_date.isoformat()}:{end_date.isoformat()}",
            )
            if cache_key in cache:
                if stats is not None:
                    stats["cache_hits"] = stats.get("cache_hits", 0) + 1
                regional_series = cache[cache_key]
            else:
                url = build_regional_api_url(
                    latitude_min,
                    latitude_max,
                    longitude_min,
                    longitude_max,
                    parameter,
                    start_date,
                    end_date,
                )
                request = Request(url, headers={"User-Agent": "app preprocess pipeline"})
                last_error: Exception | None = None
                for attempt in range(1, MAX_RETRIES + 1):
                    try:
                        with urlopen(request, timeout=HTTP_TIMEOUT_SECONDS) as response:
                            payload = json.load(response)
                        spatial_average = aggregate_regional_payload(payload, parameter)
                        regional_series = {parameter: spatial_average}
                        cache[cache_key] = regional_series
                        if stats is not None:
                            stats["cache_misses"] = stats.get("cache_misses", 0) + 1
                        if cache_path is not None:
                            save_nasa_cache(cache_path, cache)
                        break
                    except (HTTPError, URLError, TimeoutError, KeyError, json.JSONDecodeError) as exc:
                        last_error = exc
                        if attempt == MAX_RETRIES:
                            raise RuntimeError(
                                "No se pudo consultar NASA POWER regional para "
                                f"region={region_label}, parameter={parameter}, "
                                f"lat=({latitude_min},{latitude_max}), lon=({longitude_min},{longitude_max}), "
                                f"start={start_date}, end={end_date}"
                            ) from last_error
                        sleep_seconds = min(RETRY_SLEEP_SECONDS * (2 ** (attempt - 1)), MAX_RETRY_SLEEP_SECONDS)
                        if isinstance(exc, HTTPError) and exc.code == 429:
                            retry_after = parse_retry_after_seconds(exc)
                            if retry_after is not None:
                                sleep_seconds = max(sleep_seconds, retry_after)
                        time.sleep(sleep_seconds)
            for day_key, value in regional_series.get(parameter, {}).items():
                combined[parameter].setdefault(day_key, []).append(value)
    aggregated = {
        parameter: {
            day_key: average_metric_values([item for item in values if item is not None])
            for day_key, values in day_values.items()
        }
        for parameter, day_values in combined.items()
    }
    return aggregated, tile_count


def fetch_nasa_regional_country_series(
    country: str,
    start_date: date,
    end_date: date,
    cache: dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]],
    cache_path: Path | None = None,
    stats: dict[str, int] | None = None,
) -> tuple[dict[str, dict[str, float | None]], int]:
    bounds = get_country_bounds(country)
    return fetch_nasa_regional_series(
        bounds.country,
        list(iter_country_tile_bounds(bounds.country)),
        start_date,
        end_date,
        cache,
        cache_path=cache_path,
        stats=stats,
    )


def daterange(start_date: date, end_date: date) -> list[date]:
    return [start_date + timedelta(days=offset) for offset in range((end_date - start_date).days + 1)]


def _build_nasa_series_aggregation_cache(
    nasa_series: dict[str, dict[str, float | None]],
) -> dict[str, object]:
    cache_key = id(nasa_series)
    with _NASA_AGGREGATION_LOCK:
        cached = _NASA_SERIES_AGGREGATION_CACHE.get(cache_key)
        if cached is not None:
            return cached

    keys = sorted(set(nasa_series.get("T2M", {}).keys()) | set(nasa_series.get("PRECTOTCORR", {}).keys()))
    key_to_index = {key: index for index, key in enumerate(keys)}
    temperature_prefix_sum = [0.0]
    temperature_prefix_count = [0]
    precipitation_prefix_sum = [0.0]

    running_temperature_sum = 0.0
    running_temperature_count = 0
    running_precipitation_sum = 0.0
    for key in keys:
        temperature = nasa_series.get("T2M", {}).get(key)
        precipitation = nasa_series.get("PRECTOTCORR", {}).get(key)
        if temperature is not None:
            running_temperature_sum += temperature
            running_temperature_count += 1
        if precipitation is not None:
            running_precipitation_sum += precipitation
        temperature_prefix_sum.append(running_temperature_sum)
        temperature_prefix_count.append(running_temperature_count)
        precipitation_prefix_sum.append(running_precipitation_sum)

    prepared = {
        "key_to_index": key_to_index,
        "temperature_prefix_sum": temperature_prefix_sum,
        "temperature_prefix_count": temperature_prefix_count,
        "precipitation_prefix_sum": precipitation_prefix_sum,
    }
    with _NASA_AGGREGATION_LOCK:
        _NASA_SERIES_AGGREGATION_CACHE[cache_key] = prepared
    return prepared


def aggregate_metrics(
    nasa_series: dict[str, dict[str, float | None]],
    start_date: date,
    end_date: date,
) -> dict[str, float | None]:
    if end_date < start_date:
        return {"avg_t2m": None, "total_prectotcorr": None}

    window_cache_key = (id(nasa_series), start_date.isoformat(), end_date.isoformat())
    with _NASA_AGGREGATION_LOCK:
        cached_window_metrics = _NASA_WINDOW_METRICS_CACHE.get(window_cache_key)
    if cached_window_metrics is not None:
        return cached_window_metrics

    prepared = _build_nasa_series_aggregation_cache(nasa_series)
    key_to_index = prepared["key_to_index"]
    start_key = format_nasa_date(start_date)
    end_key = format_nasa_date(end_date)
    start_index = key_to_index.get(start_key)
    end_index = key_to_index.get(end_key)

    if start_index is None or end_index is None or end_index < start_index:
        temperatures: list[float] = []
        precipitations: list[float] = []
        for current_date in daterange(start_date, end_date):
            key = format_nasa_date(current_date)
            temperature = nasa_series["T2M"].get(key)
            precipitation = nasa_series["PRECTOTCORR"].get(key)
            if temperature is not None:
                temperatures.append(temperature)
            if precipitation is not None:
                precipitations.append(precipitation)
        metrics = {
            "avg_t2m": (sum(temperatures) / len(temperatures)) if temperatures else None,
            "total_prectotcorr": sum(precipitations) if precipitations else None,
        }
    else:
        temperature_prefix_sum = prepared["temperature_prefix_sum"]
        temperature_prefix_count = prepared["temperature_prefix_count"]
        precipitation_prefix_sum = prepared["precipitation_prefix_sum"]
        temperature_sum = temperature_prefix_sum[end_index + 1] - temperature_prefix_sum[start_index]
        temperature_count = temperature_prefix_count[end_index + 1] - temperature_prefix_count[start_index]
        precipitation_sum = precipitation_prefix_sum[end_index + 1] - precipitation_prefix_sum[start_index]
        metrics = {
            "avg_t2m": (temperature_sum / temperature_count) if temperature_count else None,
            "total_prectotcorr": precipitation_sum if precipitation_sum != 0.0 else 0.0,
        }

    with _NASA_AGGREGATION_LOCK:
        _NASA_WINDOW_METRICS_CACHE[window_cache_key] = metrics
    return metrics


def build_harvest_back_windows(harvesting_date: date) -> list[tuple[str, date, date]]:
    windows: list[tuple[str, date, date]] = []
    for week_idx in range(1, HARVEST_BACK_WEEKS + 1):
        end_date = harvesting_date - timedelta(days=(week_idx - 1) * 7)
        start_date = end_date - timedelta(days=6)
        windows.append((f"harvest_back_week_{week_idx}", start_date, end_date))
    return windows


def build_intermediate_window(planting_date: date, harvesting_date: date) -> tuple[date, date] | None:
    harvest_back_windows = build_harvest_back_windows(harvesting_date)
    start_date = planting_date + timedelta(days=14)
    end_date = harvest_back_windows[-1][1] - timedelta(days=1)
    if end_date < start_date:
        return None
    return start_date, end_date


def clamp_date_to_year(year: int, month: int, day: int) -> date:
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, min(day, last_day))


def build_climate_windows(
    planting_date: date,
    harvesting_date: date,
) -> tuple[tuple[date, date], tuple[date, date], list[tuple[str, date, date]], tuple[date, date] | None]:
    planting_week_1 = (planting_date, planting_date + timedelta(days=6))
    planting_week_2 = (planting_date + timedelta(days=7), planting_date + timedelta(days=13))
    harvest_back_windows = build_harvest_back_windows(harvesting_date)
    intermediate_window = build_intermediate_window(planting_date, harvesting_date)
    return planting_week_1, planting_week_2, harvest_back_windows, intermediate_window


@lru_cache(maxsize=4096)
def resolve_cached_climate_window_payload(
    planting_date_iso: str,
    harvesting_date_iso: str,
) -> tuple[
    tuple[date, date],
    tuple[date, date],
    tuple[tuple[str, date, date], ...],
    tuple[date, date] | None,
    date,
    date,
]:
    planting_date = parse_iso_date(planting_date_iso)
    harvesting_date = parse_iso_date(harvesting_date_iso)
    planting_week_1, planting_week_2, harvest_back_windows, intermediate_window = build_climate_windows(
        planting_date,
        harvesting_date,
    )
    all_windows = _build_all_climate_windows(
        planting_week_1,
        planting_week_2,
        harvest_back_windows,
        intermediate_window,
    )
    start_date = min(window_start for window_start, _ in all_windows)
    end_date = max(window_end for _, window_end in all_windows)
    return (
        planting_week_1,
        planting_week_2,
        tuple(harvest_back_windows),
        intermediate_window,
        start_date,
        end_date,
    )


def average_metric_values(values: list[float | None]) -> float | None:
    numeric_values = [value for value in values if value is not None]
    if not numeric_values:
        return None
    return sum(numeric_values) / len(numeric_values)


def resolve_forecast_reference_dates(
    forecast_planting_date: str,
    forecast_harvesting_date: str,
    years_back: int,
) -> list[tuple[date, date]]:
    planting_template = parse_iso_date(forecast_planting_date)
    harvesting_template = parse_iso_date(forecast_harvesting_date)
    current_year = date.today().year
    reference_years = [current_year - offset for offset in range(1, years_back + 1)]
    return [
        (
            clamp_date_to_year(year, planting_template.month, planting_template.day),
            clamp_date_to_year(year, harvesting_template.month, harvesting_template.day),
        )
        for year in reference_years
    ]


def _build_all_climate_windows(
    planting_week_1: tuple[date, date],
    planting_week_2: tuple[date, date],
    harvest_back_windows: list[tuple[str, date, date]],
    intermediate_window: tuple[date, date] | None,
) -> list[tuple[date, date]]:
    windows = [planting_week_1, planting_week_2]
    windows.extend((start_date, end_date) for _, start_date, end_date in harvest_back_windows)
    if intermediate_window is not None:
        windows.append(intermediate_window)
    return windows


def _compute_climate_window_metrics(
    nasa_series: dict[str, dict[str, float | None]],
    planting_week_1: tuple[date, date],
    planting_week_2: tuple[date, date],
    harvest_back_windows: list[tuple[str, date, date]],
    intermediate_window: tuple[date, date] | None,
) -> dict[str, float | None]:
    year_metrics: dict[str, float | None] = {}

    planting_week_1_metrics = aggregate_metrics(nasa_series, *planting_week_1)
    year_metrics["planting_week_1_total_prectotcorr"] = planting_week_1_metrics["total_prectotcorr"]

    planting_week_2_metrics = aggregate_metrics(nasa_series, *planting_week_2)
    year_metrics["planting_week_2_total_prectotcorr"] = planting_week_2_metrics["total_prectotcorr"]
    year_metrics["planting_week_2_avg_t2m"] = planting_week_2_metrics["avg_t2m"]

    year_metrics["intermediate_period_total_prectotcorr"] = None
    year_metrics["intermediate_period_avg_t2m"] = None
    if intermediate_window is not None:
        intermediate_metrics = aggregate_metrics(nasa_series, *intermediate_window)
        intermediate_days = (intermediate_window[1] - intermediate_window[0]).days + 1
        weekly_precipitation: float | None = None
        if intermediate_metrics["total_prectotcorr"] is not None and intermediate_days > 0:
            weekly_precipitation = (intermediate_metrics["total_prectotcorr"] / intermediate_days) * 7
        year_metrics["intermediate_period_total_prectotcorr"] = weekly_precipitation
        year_metrics["intermediate_period_avg_t2m"] = intermediate_metrics["avg_t2m"]

    for week_name, week_start, week_end in harvest_back_windows:
        metrics = aggregate_metrics(nasa_series, week_start, week_end)
        year_metrics[f"{week_name}_total_prectotcorr"] = metrics["total_prectotcorr"]
        year_metrics[f"{week_name}_avg_t2m"] = metrics["avg_t2m"]

    return year_metrics


def build_output_columns() -> list[str]:
    headers = list(CLIMATE_ID_COLUMNS)
    headers.extend(CLIMATE_FIXED_OUTPUT_COLUMNS)
    for week_idx in range(1, HARVEST_BACK_WEEKS + 1):
        headers.append(f"harvest_back_week_{week_idx}_total_prectotcorr")
        headers.append(f"harvest_back_week_{week_idx}_avg_t2m")
    return headers


def build_output_row(
    row: dict[str, str],
    cache: dict[tuple[str, str, str, str], dict[str, dict[str, float | None]]],
    cache_path: Path | None = None,
    climate_override: dict[str, object] | None = None,
    cache_stats: dict[str, int] | None = None,
) -> tuple[dict[str, str], dict[str, str]]:
    output_row = {column: row[column] for column in CLIMATE_ID_COLUMNS}
    audit_row = {column: row[column] for column in CLIMATE_ID_COLUMNS}
    climate_scope = str((climate_override or {}).get("climate_scope", "point") or "point").strip() or "point"
    regional_country = ""
    regional_tile_count = ""
    regional_bounds_label = ""
    regional_bounds_latitude_min = ""
    regional_bounds_latitude_max = ""
    regional_bounds_longitude_min = ""
    regional_bounds_longitude_max = ""
    forecast_locality = str(row.get("forecast_locality") or "").strip()
    forecast_locality_geonameid = str(row.get("forecast_locality_geonameid") or "").strip()
    forecast_locality_population = str(row.get("forecast_locality_population") or "").strip()
    forecast_country = str(row.get("forecast_country") or "").strip()
    forecast_country_code = str(row.get("forecast_country_code") or "").strip()
    forecast_grid_cell_id = str(row.get("forecast_grid_cell_id") or "").strip()
    forecast_grid_cell_index = str(row.get("forecast_grid_cell_index") or "").strip()
    selected_model_id = str(row.get("selected_model_id") or (climate_override or {}).get("selected_model_id") or "").strip()
    manual_grid_point_mode = climate_scope == "regional_manual" and bool(forecast_grid_cell_id)

    def resolve_regional_source() -> tuple[str, list[tuple[float, float, float, float]]]:
        nonlocal regional_country
        nonlocal regional_bounds_label
        nonlocal regional_bounds_latitude_min
        nonlocal regional_bounds_latitude_max
        nonlocal regional_bounds_longitude_min
        nonlocal regional_bounds_longitude_max
        if climate_scope == "regional_country":
            regional_country = str((climate_override or {}).get("regional_country") or row.get("country") or "").strip()
            if not regional_country:
                raise ValueError("Regional country climate mode requires a supported country.")
            return regional_country, list(iter_country_tile_bounds(regional_country))
        if climate_scope == "regional_manual":
            regional_bounds_label = str((climate_override or {}).get("regional_bounds_label") or "Manual bounds").strip() or "Manual bounds"
            try:
                latitude_min = float((climate_override or {}).get("regional_bounds_latitude_min"))
                latitude_max = float((climate_override or {}).get("regional_bounds_latitude_max"))
                longitude_min = float((climate_override or {}).get("regional_bounds_longitude_min"))
                longitude_max = float((climate_override or {}).get("regional_bounds_longitude_max"))
            except (TypeError, ValueError) as error:
                raise ValueError("Regional manual climate mode requires valid latitude and longitude bounds.") from error
            regional_bounds_latitude_min = format_number(latitude_min)
            regional_bounds_latitude_max = format_number(latitude_max)
            regional_bounds_longitude_min = format_number(longitude_min)
            regional_bounds_longitude_max = format_number(longitude_max)
            return regional_bounds_label, list(
                iter_manual_tile_bounds(
                    latitude_min,
                    latitude_max,
                    longitude_min,
                    longitude_max,
                    label=regional_bounds_label,
                )
            )
        return "", []

    if climate_override and climate_override.get("forecast_planting_date") and climate_override.get("forecast_harvesting_date"):
        regional_label = ""
        regional_tiles: list[tuple[float, float, float, float]] = []
        if climate_scope == "regional_country":
            regional_label, regional_tiles = resolve_regional_source()
        reference_pairs = resolve_forecast_reference_dates(
            str(climate_override["forecast_planting_date"]),
            str(climate_override["forecast_harvesting_date"]),
            int(climate_override.get("forecast_years_back", 5)),
        )
        per_year_metrics: list[dict[str, float | None]] = []
        audit_windows: list[dict[str, str]] = []
        for planting_date, harvesting_date in reference_pairs:
            planting_week_1, planting_week_2, harvest_back_windows, intermediate_window = build_climate_windows(
                planting_date,
                harvesting_date,
            )
            all_windows = _build_all_climate_windows(
                planting_week_1,
                planting_week_2,
                harvest_back_windows,
                intermediate_window,
            )

            start_date = min(window_start for window_start, _ in all_windows)
            end_date = max(window_end for _, window_end in all_windows)
            if climate_scope == "regional_country":
                nasa_series, tile_count = fetch_nasa_regional_series(
                    regional_label,
                    regional_tiles,
                    start_date,
                    end_date,
                    cache,
                    cache_path=cache_path,
                    stats=cache_stats,
                )
                regional_tile_count = str(tile_count)
            else:
                nasa_series = fetch_nasa_series(
                    latitude=row["latitude"],
                    longitude=row["longitude"],
                    start_date=start_date,
                    end_date=end_date,
                    cache=cache,
                    cache_path=cache_path,
                    stats=cache_stats,
                )
            year_metrics = _compute_climate_window_metrics(
                nasa_series,
                planting_week_1,
                planting_week_2,
                harvest_back_windows,
                intermediate_window,
            )

            per_year_metrics.append(year_metrics)
            audit_windows.append(
                {
                    "forecast_reference_year": str(planting_date.year),
                    "planting_week_1_start": planting_week_1[0].isoformat(),
                    "planting_week_1_end": planting_week_1[1].isoformat(),
                    "planting_week_2_start": planting_week_2[0].isoformat(),
                    "planting_week_2_end": planting_week_2[1].isoformat(),
                    "intermediate_period_start": intermediate_window[0].isoformat() if intermediate_window else "",
                    "intermediate_period_end": intermediate_window[1].isoformat() if intermediate_window else "",
                    "harvest_back_week_8_start": harvest_back_windows[-1][1].isoformat(),
                    "harvest_back_week_1_end": harvest_back_windows[0][2].isoformat(),
                }
            )

        output_columns = build_output_columns()
        for header in output_columns:
            if header in CLIMATE_ID_COLUMNS:
                continue
            output_row[header] = format_number(average_metric_values([metrics.get(header) for metrics in per_year_metrics]))

        latest_audit = audit_windows[0]
        audit_row.update(
            {
                "latitude": row["latitude"],
                "longitude": row["longitude"],
                "date_of_planting": str(climate_override["forecast_planting_date"]),
                "date_of_harvesting": str(climate_override["forecast_harvesting_date"]),
                **latest_audit,
            }
        )
        audit_row["forecast_mode"] = "five_year_average"
        audit_row["forecast_years_back"] = str(climate_override.get("forecast_years_back", 5))
        audit_row["forecast_reference_years"] = ",".join(window["forecast_reference_year"] for window in audit_windows)
        audit_row["climate_scope"] = climate_scope
        audit_row["regional_country"] = regional_country
        audit_row["regional_tile_count"] = regional_tile_count
        audit_row["regional_bounds_label"] = regional_bounds_label
        audit_row["regional_bounds_latitude_min"] = regional_bounds_latitude_min
        audit_row["regional_bounds_latitude_max"] = regional_bounds_latitude_max
        audit_row["regional_bounds_longitude_min"] = regional_bounds_longitude_min
        audit_row["regional_bounds_longitude_max"] = regional_bounds_longitude_max
        audit_row["forecast_locality"] = forecast_locality
        audit_row["forecast_locality_geonameid"] = forecast_locality_geonameid
        audit_row["forecast_locality_population"] = forecast_locality_population
        audit_row["forecast_country"] = forecast_country
        audit_row["forecast_country_code"] = forecast_country_code
        audit_row["forecast_grid_cell_id"] = forecast_grid_cell_id
        audit_row["forecast_grid_cell_index"] = forecast_grid_cell_index
        audit_row["selected_model_id"] = selected_model_id
        return output_row, audit_row

    (
        planting_week_1,
        planting_week_2,
        harvest_back_windows,
        intermediate_window,
        start_date,
        end_date,
    ) = resolve_cached_climate_window_payload(
        str(row["date_of_planting"]),
        str(row["date_of_harvesting"]),
    )
    if climate_scope == "regional_country" and not manual_grid_point_mode:
        regional_label, regional_tiles = resolve_regional_source()
        nasa_series, tile_count = fetch_nasa_regional_series(
            regional_label,
            regional_tiles,
            start_date,
            end_date,
            cache,
            cache_path=cache_path,
            stats=cache_stats,
        )
        regional_tile_count = str(tile_count)
    else:
        nasa_series = fetch_nasa_series(
            latitude=row["latitude"],
            longitude=row["longitude"],
            start_date=start_date,
            end_date=end_date,
            cache=cache,
            cache_path=cache_path,
            stats=cache_stats,
        )

    year_metrics = _compute_climate_window_metrics(
        nasa_series,
        planting_week_1,
        planting_week_2,
        harvest_back_windows,
        intermediate_window,
    )
    output_row["planting_week_1_total_prectotcorr"] = format_number(year_metrics["planting_week_1_total_prectotcorr"])
    output_row["planting_week_2_total_prectotcorr"] = format_number(year_metrics["planting_week_2_total_prectotcorr"])
    output_row["planting_week_2_avg_t2m"] = format_number(year_metrics["planting_week_2_avg_t2m"])
    output_row["intermediate_period_total_prectotcorr"] = format_number(year_metrics["intermediate_period_total_prectotcorr"])
    output_row["intermediate_period_avg_t2m"] = format_number(year_metrics["intermediate_period_avg_t2m"])

    for week_idx in range(1, HARVEST_BACK_WEEKS + 1):
        week_name = f"harvest_back_week_{week_idx}"
        output_row[f"{week_name}_total_prectotcorr"] = format_number(year_metrics[f"{week_name}_total_prectotcorr"])
        output_row[f"{week_name}_avg_t2m"] = format_number(year_metrics[f"{week_name}_avg_t2m"])

    audit_row.update(
        {
            "latitude": row["latitude"],
            "longitude": row["longitude"],
            "date_of_planting": row["date_of_planting"],
            "date_of_harvesting": row["date_of_harvesting"],
            "planting_week_1_start": planting_week_1[0].isoformat(),
            "planting_week_1_end": planting_week_1[1].isoformat(),
            "planting_week_2_start": planting_week_2[0].isoformat(),
            "planting_week_2_end": planting_week_2[1].isoformat(),
            "intermediate_period_start": intermediate_window[0].isoformat() if intermediate_window else "",
            "intermediate_period_end": intermediate_window[1].isoformat() if intermediate_window else "",
            "harvest_back_week_8_start": harvest_back_windows[-1][1].isoformat(),
            "harvest_back_week_1_end": harvest_back_windows[0][2].isoformat(),
            "climate_scope": climate_scope,
            "regional_country": regional_country,
            "regional_tile_count": regional_tile_count,
            "regional_bounds_label": regional_bounds_label,
            "regional_bounds_latitude_min": regional_bounds_latitude_min,
            "regional_bounds_latitude_max": regional_bounds_latitude_max,
            "regional_bounds_longitude_min": regional_bounds_longitude_min,
            "regional_bounds_longitude_max": regional_bounds_longitude_max,
            "forecast_locality": forecast_locality,
            "forecast_locality_geonameid": forecast_locality_geonameid,
            "forecast_locality_population": forecast_locality_population,
            "forecast_country": forecast_country,
            "forecast_country_code": forecast_country_code,
            "forecast_grid_cell_id": forecast_grid_cell_id,
            "forecast_grid_cell_index": forecast_grid_cell_index,
            "selected_model_id": selected_model_id,
        }
    )
    return output_row, audit_row


def climate_header_to_label(header: str) -> str:
    if header.endswith("_avg_t2m"):
        return f"{header} (C)"
    if header.endswith("_total_prectotcorr"):
        return f"{header} (mm)"
    return header


def build_phase02_headers(headers: list[str], target_header: str = YIELD_HEADER) -> list[str]:
    output_headers = list(headers)
    insert_after = str(target_header or "").strip() or YIELD_HEADER
    if insert_after not in output_headers:
        raise ValueError(f"No se encontro la columna objetivo '{insert_after}'.")
    insert_at = output_headers.index(insert_after) + 1
    return output_headers[:insert_at] + CLIMATE_HEADERS_WITH_UNITS + output_headers[insert_at:]


def build_phase02_records(
    headers: list[str],
    records: list[dict[str, object]],
    progress_callback: Callable[[int, str, str, dict[str, object] | None], None] | None = None,
    climate_override: dict[str, object] | None = None,
    target_header: str = YIELD_HEADER,
) -> tuple[list[str], list[dict[str, object]], dict[str, object]]:
    output_dir = PHASE_DIRS["phase02"]
    climate_scope_override = str((climate_override or {}).get("climate_scope") or "").strip()
    use_source_row_dates = bool((climate_override or {}).get("use_source_row_dates"))
    locality_mode = (
        climate_scope_override in {"country_localities", "regional_manual"}
        and (
            use_source_row_dates
            or (
                bool((climate_override or {}).get("forecast_planting_date"))
                and bool((climate_override or {}).get("forecast_harvesting_date"))
            )
        )
    )
    locality_metadata: dict[str, object] = {}
    source_records = records
    if locality_mode:
        climate_scope = str((climate_override or {}).get("climate_scope") or "").strip()
        if climate_scope == "country_localities":
            source_records, data_input_rows, locality_metadata = expand_records_for_country_localities(
                records,
                climate_override or {},
            )
        else:
            source_records, data_input_rows, locality_metadata = expand_records_for_manual_bbox_localities(
                records,
                climate_override or {},
            )
        skipped_rows = 0
    else:
        data_input_rows, skipped_rows = build_data_input_rows(records)
    if not data_input_rows:
        raise ValueError("No hay filas exportables para la fase climatica.")

    write_csv(output_dir / "dataInput.csv", CLIMATE_INPUT_COLUMNS, data_input_rows)

    cache_path = output_dir / "nasa_power_cache.json"
    cache = load_nasa_cache(cache_path)
    output_rows: list[dict[str, str]] = []
    audit_rows: list[dict[str, str]] = []
    initial_cache_size = len(cache)
    cache_stats = {"cache_hits": 0, "cache_misses": 0}
    total_rows = len(data_input_rows)
    for index, row in enumerate(data_input_rows, start=1):
        effective_climate_override = dict(climate_override or {})
        if locality_mode:
            effective_climate_override["climate_scope"] = "point"
        output_row, audit_row = build_output_row(
            row,
            cache,
            cache_path=cache_path,
            climate_override=effective_climate_override,
            cache_stats=cache_stats,
        )
        output_rows.append(output_row)
        audit_rows.append(audit_row)
        if progress_callback:
            fresh_queries = max(len(cache) - initial_cache_size, 0)
            cached_rows = max(0, index - fresh_queries)
            phase_percent = 34 + min(7, round((index / total_rows) * 7))
            progress_callback(
                phase_percent,
                "Enriching climate windows",
                (
                    f"Processing NASA POWER climate rows {index}/{total_rows}. "
                    f"New queries this run: {fresh_queries}. Cached rows served: {cached_rows}."
                ),
                {
                    "phase": "phase02",
                    "processed_rows": index,
                    "total_rows": total_rows,
                    "fresh_queries_this_run": fresh_queries,
                    "cached_rows_served": cached_rows,
                    "nasa_unique_queries_available": len(cache),
                },
            )

    output_columns = build_output_columns()
    write_csv(output_dir / "output.csv", output_columns, output_rows)
    write_csv(output_dir / "weekly_windows_audit.csv", CLIMATE_AUDIT_COLUMNS, audit_rows)

    climate_by_key = {
        tuple(normalize_key_value(row[column]) for column in CLIMATE_ID_COLUMNS): row
        for row in output_rows
    }

    phase02_headers = (
        build_country_locality_phase02_headers(headers, target_header=target_header)
        if locality_mode
        else build_phase02_headers(headers, target_header=target_header)
    )
    phase02_records: list[dict[str, object]] = []
    matched_rows = 0
    unmatched_rows = 0
    if locality_mode:
        matched_rows = len(source_records)
        for record, climate_row in zip(source_records, output_rows):
            enriched = dict(record)
            for header in output_columns:
                if header in CLIMATE_ID_COLUMNS:
                    continue
                enriched[climate_header_to_label(header)] = climate_row.get(header, "")
            phase02_records.append(enriched)
    else:
        for record in source_records:
            key = (
                normalize_key_value(record.get(TRIAL_SERIES_HEADER)),
                normalize_key_value(record.get("Site Number")),
                normalize_key_value(record.get("Plot")),
                normalize_key_value(record.get("EntryCode")),
            )
            climate_row = climate_by_key.get(key, {})
            if climate_row:
                matched_rows += 1
            else:
                unmatched_rows += 1

            enriched = dict(record)
            for header in output_columns:
                if header in CLIMATE_ID_COLUMNS:
                    continue
                enriched[climate_header_to_label(header)] = climate_row.get(header, "")
            phase02_records.append(enriched)

    write_workbook(
        output_dir / "phase02_climate_enriched.xlsx",
        "phase02_climate",
        phase02_headers,
        phase02_records,
    )

    metadata = {
        "phase": "phase02_climate_enriched",
        "input_row_count": len(records),
        "phase02_source_row_count": len(source_records),
        "climate_input_row_count": len(data_input_rows),
        "skipped_rows_missing_climate_fields": skipped_rows,
        "matched_rows": matched_rows,
        "unmatched_rows": unmatched_rows,
        "nasa_unique_queries": len(cache),
        "nasa_cache_initial_entries": initial_cache_size,
        "nasa_cache_final_entries": len(cache),
        "nasa_cache_fresh_queries_this_run": cache_stats["cache_misses"],
        "nasa_cache_hits": cache_stats["cache_hits"],
        "nasa_cache_misses": cache_stats["cache_misses"],
        "nasa_cache_file": str(cache_path),
        "header_count": len(phase02_headers),
        "climate_headers_with_units": CLIMATE_HEADERS_WITH_UNITS,
        "climate_override": climate_override or {},
        **locality_metadata,
    }
    write_metadata(output_dir / "metadata.json", metadata)
    return phase02_headers, phase02_records, metadata


def is_empty_or_zero_or_none(value: object) -> bool:
    normalized = normalize_upper(value)
    if normalized in {"", "0", "NONE"}:
        return True
    numeric = parse_number(value)
    return numeric == 0 if numeric is not None else False


def build_planting_fertilizer_value(fertilizer_type: object, fertilizer_rate: object) -> int:
    rate_value = parse_number(fertilizer_rate)
    return 0 if is_empty_or_zero_or_none(fertilizer_type) and rate_value == 0 else 1


def build_top_dressing_output_value(fertilizer_type: object, fertilizer_rate: object) -> int:
    if normalize_text(fertilizer_type) == "":
        return 0
    rate_value = parse_number(fertilizer_rate)
    if rate_value is None or rate_value == 0:
        return 0
    return 1


def build_fertilizer_times_output_value(
    top_dressing_times: object,
    second_application_type: object,
    second_application_rate: object,
) -> object:
    for candidate in (
        top_dressing_times,
        second_application_type,
        second_application_rate,
    ):
        numeric_value = parse_number(candidate)
        if numeric_value is None or numeric_value <= 0:
            continue
        return int(numeric_value) if numeric_value.is_integer() else numeric_value
    return None


def build_herbicide_output_value(
    first_herbicide: object,
    first_herbicide_alt: object,
    second_herbicide: object,
) -> int:
    first = normalize_upper(first_herbicide)
    first_alt = normalize_upper(first_herbicide_alt)
    second = normalize_upper(second_herbicide)
    if first not in {"", "0", "NA", "N/A", "NOT USED BY THE FARMER", "NOT USED"}:
        return 1
    if first_alt not in {"", "0"}:
        return 1
    if second not in {"", "0", "NOT APPLICABLE", "NO"}:
        return 1
    return 0


def build_insecticide_output_value(
    insecticide: object,
    insecticide_alt: object,
    insecticide_pests: object,
) -> int:
    left = normalize_upper(insecticide)
    right = normalize_upper(insecticide_alt)
    pests = normalize_upper(insecticide_pests)
    if left not in {"", "0"}:
        return 1
    if right not in {"", "0"}:
        return 1
    if pests not in {"", "0", "NOT APPLICABLE", "NO"}:
        return 1
    return 0


def build_phase03_headers(headers: list[str]) -> list[str]:
    output_headers: list[str] = []
    for header in headers:
        if header == FERTILIZER_TYPE_HEADER:
            output_headers.append(IS_FETILIZE_USED_HEADER)
        if header == TOP_DRESSING_TYPE_HEADER:
            output_headers.append(IS_FERTILIZER_APPLIED_HEADER)
        if header == TOP_DRESSING_TIMES_HEADER:
            output_headers.append(FERTILIZER_TIMES_HEADER)
        if header == FIRST_HERBICIDE_HEADER:
            output_headers.append(HERBICIDE_OUTPUT_HEADER)
        if header == INSECTICIDE_HEADER:
            output_headers.append(INSECTICIDE_OUTPUT_HEADER)
        if header in {
            FERTILIZER_TYPE_HEADER,
            FERTILIZER_RATE_HEADER,
            TOP_DRESSING_TYPE_HEADER,
            TOP_DRESSING_RATE_HEADER,
            TOP_DRESSING_TIMES_HEADER,
            SECOND_TOP_DRESSING_APPLICATION_HEADER,
            SECOND_TOP_DRESSING_RATE_HEADER,
            FIRST_HERBICIDE_HEADER,
            FIRST_HERBICIDE_ALT_HEADER,
            SECOND_HERBICIDE_HEADER,
            INSECTICIDE_HEADER,
            INSECTICIDE_ALT_HEADER,
            INSECTICIDE_PESTS_HEADER,
        }:
            continue
        output_headers.append(header)
    return output_headers


def build_phase03_records(headers: list[str], records: list[dict[str, object]]) -> tuple[list[str], list[dict[str, object]]]:
    phase03_headers = build_phase03_headers(headers)
    phase03_records: list[dict[str, object]] = []
    for record in records:
        derived = dict(record)
        derived[IS_FETILIZE_USED_HEADER] = build_planting_fertilizer_value(
            record.get(FERTILIZER_TYPE_HEADER),
            record.get(FERTILIZER_RATE_HEADER),
        )
        derived[IS_FERTILIZER_APPLIED_HEADER] = build_top_dressing_output_value(
            record.get(TOP_DRESSING_TYPE_HEADER),
            record.get(TOP_DRESSING_RATE_HEADER),
        )
        derived[FERTILIZER_TIMES_HEADER] = build_fertilizer_times_output_value(
            record.get(TOP_DRESSING_TIMES_HEADER),
            record.get(SECOND_TOP_DRESSING_APPLICATION_HEADER),
            record.get(SECOND_TOP_DRESSING_RATE_HEADER),
        )
        derived[HERBICIDE_OUTPUT_HEADER] = build_herbicide_output_value(
            record.get(FIRST_HERBICIDE_HEADER),
            record.get(FIRST_HERBICIDE_ALT_HEADER),
            record.get(SECOND_HERBICIDE_HEADER),
        )
        derived[INSECTICIDE_OUTPUT_HEADER] = build_insecticide_output_value(
            record.get(INSECTICIDE_HEADER),
            record.get(INSECTICIDE_ALT_HEADER),
            record.get(INSECTICIDE_PESTS_HEADER),
        )
        phase03_records.append(derived)

    output_dir = PHASE_DIRS["phase03"]
    write_workbook(
        output_dir / "phase03_management_derived.xlsx",
        "phase03_manage",
        phase03_headers,
        phase03_records,
    )
    write_metadata(
        output_dir / "metadata.json",
        {
            "phase": "phase03_management_derived",
            "header_count": len(phase03_headers),
            "row_count": len(phase03_records),
            "added_headers": [
                IS_FETILIZE_USED_HEADER,
                IS_FERTILIZER_APPLIED_HEADER,
                FERTILIZER_TIMES_HEADER,
                HERBICIDE_OUTPUT_HEADER,
                INSECTICIDE_OUTPUT_HEADER,
            ],
        },
    )
    return phase03_headers, phase03_records


def normalize_binary_value(value: object) -> int | None:
    if value is None:
        return None
    text = normalize_text(value)
    if text == "" or text.lower() in {"none", "null", "nan", "missing"}:
        return None
    try:
        numeric = float(text)
    except ValueError:
        return None
    if numeric == 0:
        return 0
    if numeric == 1:
        return 1
    return None


def build_is_fertilizer_value(planting_fertilizer_value: object, top_dressing_fertilizer_value: object) -> int | None:
    left = normalize_binary_value(planting_fertilizer_value)
    right = normalize_binary_value(top_dressing_fertilizer_value)
    if left is None and right is None:
        return None
    if left == 1 or right == 1:
        return 1
    if left == 0 or right == 0:
        return 0
    return None


def build_percentage_value(numerator: object, denominator: object) -> float | None:
    left = parse_number(numerator)
    right = parse_number(denominator)
    if left is None or right is None or right == 0:
        return 0
    value = (left / right) * 100
    if value > 100:
        return None
    return value


def build_bounded_ratio(numerator: object, denominator: object) -> float | None:
    left = parse_number(numerator)
    right = parse_number(denominator)
    if left is None or right is None or right == 0:
        return None
    value = left / right
    if value == 0 or value > 1:
        return None
    return value


def build_ears_per_plant_value(
    plant_harvest_stand_value: object,
    plant_stand_value: object,
) -> float | None:
    numerator = parse_number(plant_harvest_stand_value)
    denominator = parse_number(plant_stand_value)
    if numerator is None or denominator is None or denominator == 0:
        return None
    value = numerator / denominator
    if value == 0 or value > 1:
        return None
    return value


def build_plot_area(
    row_length: object,
    spacing: object,
    rows_per_entry: object,
) -> float | None:
    row_length_value = parse_number(row_length)
    spacing_value = parse_number(spacing)
    rows_value = parse_number(rows_per_entry)
    if row_length_value is None or spacing_value is None or rows_value is None:
        return None
    return row_length_value * spacing_value * rows_value


def build_phase04_headers(headers: list[str]) -> list[str]:
    output_headers: list[str] = []
    for header in headers:
        if header == IS_FETILIZE_USED_HEADER:
            continue
        if header == IS_FERTILIZER_APPLIED_HEADER:
            output_headers.append(IS_FERTILIZER_HEADER)
            continue
        if header in {ROW_LENGTH_HEADER, SPACING_HEADER, ROWS_PER_ENTRY_HEADER}:
            continue
        output_headers.append(header)
    for header in [
        PERCENTAGE_PLANT_HARVEST_STAND_HEADER,
        PERCENTAGE_ROOT_LODGING_HEADER,
        PERCENTAGE_STEM_LODGING_HEADER,
        PERCENTAGE_POOR_HUSK_COVER_HEADER,
        PERCENTAGE_EARS_HARVESTED_HEADER,
        PERCENTAGE_ROTTEN_EARS_HEADER,
        PLOT_AREA_HEADER,
        ROOT_LODGING_HEADER,
        STEM_LODGING_HEADER,
        BAD_HUSK_COVER_HEADER,
        EARS_PER_PLANT_HEADER,
        ROTTEN_EARS_HEADER,
        EAR_POSITION_HEADER,
    ]:
        if header not in output_headers:
            output_headers.append(header)
    return output_headers


def build_phase04_records(headers: list[str], records: list[dict[str, object]]) -> tuple[list[str], list[dict[str, object]]]:
    phase04_headers = build_phase04_headers(headers)
    phase04_records: list[dict[str, object]] = []
    for record in records:
        consolidated = dict(record)
        source_set_suffix = infer_source_set_suffix(record)
        consolidated[IS_FERTILIZER_HEADER] = build_is_fertilizer_value(
            record.get(IS_FETILIZE_USED_HEADER),
            record.get(IS_FERTILIZER_APPLIED_HEADER),
        )
        consolidated[PERCENTAGE_PLANT_HARVEST_STAND_HEADER] = build_percentage_value(
            record.get(PLANT_HARVEST_STAND_HEADER),
            record.get(PLANT_STAND_HEADER),
        )
        consolidated[PERCENTAGE_ROOT_LODGING_HEADER] = build_percentage_value(
            record.get(ROOT_LODGING_COUNT_HEADER),
            record.get(PLANT_HARVEST_STAND_HEADER),
        )
        consolidated[PERCENTAGE_STEM_LODGING_HEADER] = build_percentage_value(
            record.get(STEM_LODGING_COUNT_HEADER),
            record.get(PLANT_HARVEST_STAND_HEADER),
        )
        consolidated[PERCENTAGE_POOR_HUSK_COVER_HEADER] = build_percentage_value(
            record.get(POOR_HUSK_COVER_COUNT_HEADER),
            record.get(EARS_HARVESTED_COUNT_HEADER),
        )
        consolidated[PERCENTAGE_EARS_HARVESTED_HEADER] = build_percentage_value(
            record.get(EARS_HARVESTED_COUNT_HEADER),
            record.get(PLANT_HARVEST_STAND_HEADER),
        )
        consolidated[PERCENTAGE_ROTTEN_EARS_HEADER] = build_percentage_value(
            record.get(ROTTEN_EARS_COUNT_HEADER),
            record.get(EARS_HARVESTED_COUNT_HEADER),
        )
        computed_plot_area = build_plot_area(
            record.get(ROW_LENGTH_HEADER),
            record.get(SPACING_HEADER),
            record.get(ROWS_PER_ENTRY_HEADER),
        )
        consolidated[PLOT_AREA_HEADER] = (
            record.get(PLOT_AREA_HEADER)
            if source_set_suffix == "01" and not is_blank(record.get(PLOT_AREA_HEADER))
            else computed_plot_area
        )
        consolidated[ROOT_LODGING_HEADER] = consolidated[PERCENTAGE_ROOT_LODGING_HEADER]
        consolidated[STEM_LODGING_HEADER] = consolidated[PERCENTAGE_STEM_LODGING_HEADER]
        consolidated[BAD_HUSK_COVER_HEADER] = consolidated[PERCENTAGE_POOR_HUSK_COVER_HEADER]
        consolidated[EARS_PER_PLANT_HEADER] = record.get(EARS_PER_PLANT_HEADER)
        consolidated[ROTTEN_EARS_HEADER] = consolidated[PERCENTAGE_ROTTEN_EARS_HEADER]
        consolidated[EAR_POSITION_HEADER] = build_bounded_ratio(
            record.get(EAR_HEIGHT_HEADER),
            record.get(PLANT_HEIGHT_HEADER),
        )
        phase04_records.append(consolidated)

    output_dir = PHASE_DIRS["phase04"]
    write_workbook(
        output_dir / "phase04_fertilization_consolidated.xlsx",
        "phase04_fertilize",
        phase04_headers,
        phase04_records,
    )
    write_metadata(
        output_dir / "metadata.json",
        {
            "phase": "phase04_fertilization_consolidated",
            "header_count": len(phase04_headers),
            "row_count": len(phase04_records),
            "added_headers": [
                IS_FERTILIZER_HEADER,
                PERCENTAGE_PLANT_HARVEST_STAND_HEADER,
                PERCENTAGE_ROOT_LODGING_HEADER,
                PERCENTAGE_STEM_LODGING_HEADER,
                PERCENTAGE_POOR_HUSK_COVER_HEADER,
                PERCENTAGE_EARS_HARVESTED_HEADER,
                PERCENTAGE_ROTTEN_EARS_HEADER,
                PLOT_AREA_HEADER,
                ROOT_LODGING_HEADER,
                STEM_LODGING_HEADER,
                BAD_HUSK_COVER_HEADER,
                EARS_PER_PLANT_HEADER,
                ROTTEN_EARS_HEADER,
                EAR_POSITION_HEADER,
            ],
        },
    )
    return phase04_headers, phase04_records


def build_phase05_records(headers: list[str], records: list[dict[str, object]]) -> tuple[list[str], list[dict[str, object]]]:
    phase05_headers = [IDPK_HEADER] + list(headers)
    counters = {"01": 0, "02": 0}
    phase05_records: list[dict[str, object]] = []
    for record in records:
        suffix = infer_source_set_suffix(record)
        counters[suffix] += 1
        enriched = dict(record)
        enriched[IDPK_HEADER] = f"{counters[suffix]:04d}-Set-{suffix}"
        phase05_records.append(enriched)

    output_dir = PHASE_DIRS["phase05"]
    write_workbook(
        output_dir / "phase05_final_with_idpk.xlsx",
        "phase05_final",
        phase05_headers,
        phase05_records,
    )
    write_metadata(
        output_dir / "metadata.json",
        {
            "phase": "phase05_final_with_idpk",
            "header_count": len(phase05_headers),
            "row_count": len(phase05_records),
            "idpk_counters": counters,
        },
    )
    return phase05_headers, phase05_records


def build_phase06_headers(headers: list[str]) -> list[str]:
    missing_headers = [header for header in PHASE1_COMPAT_HEADERS if header not in headers]
    if missing_headers:
        raise ValueError(
            "Faltan columnas para la salida compatible con phase1: "
            + ", ".join(missing_headers)
        )
    trace_headers = [
        header for header in PHASE06_TRACE_HEADERS if header in headers and header not in PHASE1_COMPAT_HEADERS
    ]
    dg_headers = [
        header
        for header in headers
        if header.startswith("DG") and header[2:].isdigit() and header not in PHASE1_COMPAT_HEADERS
    ]
    phase1_headers = list(PHASE1_COMPAT_HEADERS)
    idpk_position = phase1_headers.index(IDPK_HEADER) + 1
    return phase1_headers[:idpk_position] + trace_headers + dg_headers + phase1_headers[idpk_position:]


def build_phase06_records(headers: list[str], records: list[dict[str, object]]) -> tuple[list[str], list[dict[str, object]]]:
    phase06_headers = build_phase06_headers(headers)
    date_headers = {DATE_PLANTING_HEADER, DATE_HARVESTING_HEADER}
    phase06_records: list[dict[str, object]] = []
    for record in records:
        projected = {header: record.get(header) for header in phase06_headers}
        for header in date_headers:
            if header not in projected:
                continue
            parsed = parse_date(projected.get(header))
            if parsed is not None:
                projected[header] = parsed
        phase06_records.append(projected)

    output_dir = PHASE_DIRS["phase06"]
    write_workbook(
        output_dir / "phase06_phase1_compatible.xlsx",
        "phase06_phase1",
        phase06_headers,
        phase06_records,
    )
    write_metadata(
        output_dir / "metadata.json",
        {
            "phase": "phase06_phase1_compatible",
            "header_count": len(phase06_headers),
            "row_count": len(phase06_records),
            "source_phase": "phase05_final_with_idpk",
            "phase1_compatible_headers": list(PHASE1_COMPAT_HEADERS),
            "trace_headers": [
                header for header in phase06_headers if header not in PHASE1_COMPAT_HEADERS
            ],
        },
    )
    return phase06_headers, phase06_records


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Construye el pipeline interno de preprocess de app desde test.xlsx hoja 0."
    )
    parser.add_argument(
        "--source-workbook",
        default=str(DEFAULT_SOURCE_WORKBOOK),
        help="Ruta al workbook fuente consolidado.",
    )
    parser.add_argument(
        "--source-sheet-index",
        type=int,
        default=DEFAULT_SOURCE_SHEET_INDEX,
        help="Indice de la hoja fuente consolidada.",
    )
    parser.add_argument(
        "--keep-existing",
        action="store_true",
        help="No elimina salidas previas antes de reconstruir.",
    )
    args = parser.parse_args()

    source_workbook = Path(args.source_workbook)
    if not source_workbook.exists():
        raise FileNotFoundError(f"No existe el workbook fuente: {source_workbook}")

    if not args.keep_existing:
        clean_phase_outputs()

    headers, records = build_phase00_snapshot(source_workbook, args.source_sheet_index)
    phase01_headers, phase01_records, phase01_metadata = build_phase01_prepared(headers, records)
    phase02_headers, phase02_records, phase02_metadata = build_phase02_records(phase01_headers, phase01_records)
    phase03_headers, phase03_records = build_phase03_records(phase02_headers, phase02_records)
    phase04_headers, phase04_records = build_phase04_records(phase03_headers, phase03_records)
    phase05_headers, phase05_records = build_phase05_records(phase04_headers, phase04_records)
    build_phase06_records(phase05_headers, phase05_records)

    print("Pipeline interno construido en cimmyt_app/preprocess")
    print(f"Entrada usada: {source_workbook}")
    print(f"Filas phase00: {len(records)}")
    print(f"Filas phase01 preparadas: {phase01_metadata['output_row_count']}")
    print(f"Filas con clima exportable: {phase02_metadata['climate_input_row_count']}")
    print(f"Consultas unicas a NASA POWER: {phase02_metadata['nasa_unique_queries']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
