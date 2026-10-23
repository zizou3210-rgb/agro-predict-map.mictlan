import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from ce_pipeline.phase01 import phase1
from pipeline.progress import work_progress
from preprocess.ea_pipeline import write_workbook


class WorkbookProgressTests(unittest.TestCase):
    def test_streamed_workbook_supports_zip64_and_completes_after_close(self):
        from zipfile import ZipFile
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / 'large-sheet.xlsx'
            records = [{'value': 'soil & climate ' * 100} for _ in range(32)]
            events = []
            with patch('zipfile.ZIP64_LIMIT', 2048):
                phase1.write_xlsx_rows(path, ['value'], ((row['value'],) for row in records), len(records), lambda *args: events.append(args))
            with ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
                self.assertGreater(archive.getinfo('xl/worksheets/sheet1.xml').file_size, 2048)
                self.assertGreaterEqual(archive.getinfo('xl/worksheets/sheet1.xml').extract_version, 45)
            self.assertEqual(phase1.read_sheet_headers_and_rows(path)[1], [{'value': row['value'].strip()} for row in records])
            self.assertEqual(events[-1], (32, 32, True))

    def test_default_writer_without_callback_roundtrips(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / 'default.xlsx'
            records = [{'name': 'maíz & suelo', 'value': '12'}]
            phase1.write_xlsx(path, ['name', 'value'], records)
            self.assertEqual(phase1.read_sheet_headers_and_rows(path), (['name', 'value'], records))

    def test_streamed_workbook_reports_completion_after_close_and_roundtrips(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / 'rows.xlsx'
            events = []
            def report(done, total, finished=False):
                events.append((done, total, finished))
                if finished:
                    self.assertEqual(phase1.read_sheet_headers_and_rows(path)[1][-1]['value'], '1999')
            phase1.write_xlsx_rows(path, ['value'], ((i,) for i in range(2000)), 2000, report)
            self.assertEqual(events[0], (0, 2000, False))
            self.assertEqual(events[-1], (2000, 2000, True))
            read_events = []
            headers, rows = phase1.read_sheet_headers_and_rows(path, lambda *args: read_events.append(args))
            self.assertEqual(len(rows), 2000)
            self.assertEqual(headers, ['value'])
            self.assertTrue(read_events[-1][2])
            self.assertEqual(read_events[-1][0], read_events[-1][1])

    def test_climate_workbook_preserves_types_sheet_and_rows(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            records = [{'name': 'á & <x>', 'number': 1.5, 'date': datetime(2026, 4, 7)}, {'name': None}]
            events = []
            for name, callback in [('original', None), ('progress', lambda *args: events.append(args))]:
                write_workbook(root / (name + '.xlsx'), 'climate', ['name', 'number', 'date'], records, callback)
            books = [load_workbook(root / (name + '.xlsx')) for name in ['original', 'progress']]
            try:
                self.assertEqual(books[0].sheetnames, books[1].sheetnames)
                self.assertEqual(list(books[0].active.values), list(books[1].active.values))
            finally:
                for book in books:
                    book.close()
            self.assertEqual(events[-1], (2, 2, True))

    def test_progress_throttles_updates_and_does_not_finish_early(self):
        events = []
        with patch('pipeline.progress.time.monotonic', side_effect=[0, 0, .1, 1.2, 1.3, 1.4]):
            report = work_progress(lambda *args: events.append(args), 90, 'Saving', phase='phase03')
            report(0, 10)
            report(1, 10)
            report(5, 10)
            report(10, 10)
            report(10, 10, True)
        self.assertEqual([event[3]['stage_percent'] for event in events], [0, 50, 100])
        self.assertEqual(events[1][3]['stage_completed'], 5)
        self.assertGreater(events[1][3]['stage_eta_seconds'], 0)

    def test_failed_write_never_reports_completion(self):
        def broken_rows():
            yield ('first',)
            raise ValueError('source failed')
        events = []
        with tempfile.TemporaryDirectory() as root:
            with self.assertRaisesRegex(ValueError, 'source failed'):
                phase1.write_xlsx_rows(Path(root) / 'failed.xlsx', ['name'], broken_rows(), 2, lambda *args: events.append(args))
        self.assertFalse(any(len(event) > 2 and event[2] for event in events))
