from __future__ import annotations

import unittest
from pathlib import Path

from cimmyt_app.workbook_preview import (
    build_original_feature_collection_from_xlsx,
    describe_distinct_germplasm_names_from_xlsx,
    list_distinct_germplasm_names_from_xlsx,
    validate_required_workbook_headers,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEST_WORKBOOK = PROJECT_ROOT / "cimmyt_app" / "template" / "test.xlsx"
ORIGINAL_CATEGORY_WORKBOOK = PROJECT_ROOT / "cimmyt_app" / "template" / "original_category.xlsx"


class TestWorkbookPreview(unittest.TestCase):
    def test_list_distinct_germplasm_names_from_xlsx_reads_name_column(self) -> None:
        names = list_distinct_germplasm_names_from_xlsx(TEST_WORKBOOK)

        self.assertIn("CIM-RT-EAPP1-E -001", names)
        self.assertIn("CIM-RT-EAPP1-E -007", names)
        self.assertEqual(names, sorted(names))


    def test_describe_distinct_germplasm_names_from_xlsx_returns_profile_metadata(self) -> None:
        description = describe_distinct_germplasm_names_from_xlsx(TEST_WORKBOOK)

        self.assertIn("germplasm_profiles_by_name", description)
        self.assertIsInstance(description["germplasm_profiles_by_name"], dict)
        self.assertEqual(description["count"], len(description["germplasm_names"]))
        self.assertIn("headers", description)
        self.assertIn("id_field_candidates", description)
        self.assertNotIn("Name", description["id_field_candidates"])

    def test_build_original_feature_collection_from_xlsx_returns_mappable_features(self) -> None:
        geojson = build_original_feature_collection_from_xlsx(TEST_WORKBOOK)

        self.assertEqual(geojson["type"], "FeatureCollection")
        self.assertTrue(geojson["features"])
        self.assertGreater(geojson["metadata"]["attribute_count"], 0)
        first_feature = geojson["features"][0]
        self.assertEqual(first_feature["geometry"]["type"], "Point")
        self.assertIn("Name", first_feature["properties"])
        self.assertIn("Grain Yield (T/Ha)", first_feature["properties"])
        self.assertIn("Rank", first_feature["properties"])
        self.assertIn("_GPS coordinates_latitude", first_feature["properties"])
        self.assertIn("_GPS coordinates_longitude", first_feature["properties"])

    def test_original_category_template_exposes_pioneer_multi_profile_demo(self) -> None:
        description = describe_distinct_germplasm_names_from_xlsx(ORIGINAL_CATEGORY_WORKBOOK)

        self.assertIn("6232", description["germplasm_profiles_by_name"])
        self.assertIn("Duma 43", description["germplasm_profiles_by_name"])
        self.assertIn("Makueni", description["germplasm_profiles_by_name"])

        pioneer_demo = description["germplasm_profiles_by_name"].get("Pioneer")
        self.assertIsNotNone(pioneer_demo)
        self.assertTrue(pioneer_demo["multi_record"])
        self.assertTrue(pioneer_demo["multi_profile_mode"])
        self.assertEqual(pioneer_demo["profile_count"], 2)
        self.assertEqual(
            pioneer_demo["varying_columns"],
            [
                {
                    "header": "Education/Training of the farmer",
                    "values": ["Post-secondary", "Primary"],
                    "unique_count": 2,
                }
            ],
        )

    def test_describe_distinct_germplasm_names_from_xlsx_can_return_selected_id_values(self) -> None:
        description = describe_distinct_germplasm_names_from_xlsx(
            ORIGINAL_CATEGORY_WORKBOOK,
            selected_id_header="Name",
        )

        self.assertEqual(description["selected_id_header"], "Name")
        self.assertIn("germplasm_id_values_by_name", description)
        self.assertEqual(description["germplasm_id_values_by_name"].get("Pioneer"), ["Pioneer"])

    def test_describe_distinct_germplasm_names_from_xlsx_fills_missing_selected_id_for_repeated_rows(self) -> None:
        from openpyxl import Workbook
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir_name:
            workbook_path = Path(tmp_dir_name) / "missing_ids.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Nursery Id", "Education"])
            worksheet.append(["Pioneer", "", "Primary"])
            worksheet.append(["Pioneer", "null", "Post-secondary"])
            worksheet.append(["Single", "", "Primary"])
            workbook.save(workbook_path)
            workbook.close()

            description = describe_distinct_germplasm_names_from_xlsx(
                workbook_path,
                selected_id_header="Nursery Id",
            )

            self.assertEqual(description["germplasm_id_values_by_name"].get("Pioneer"), ["1", "2"])
            self.assertEqual(description["germplasm_id_values_by_name"].get("Single"), [])

    def test_describe_distinct_germplasm_names_from_xlsx_requires_farm_for_uploaded_mode(self) -> None:
        from openpyxl import Workbook
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir_name:
            workbook_path = Path(tmp_dir_name) / "missing_farm.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Education"])
            worksheet.append(["Pioneer", "Primary"])
            workbook.save(workbook_path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "Farm column"):
                describe_distinct_germplasm_names_from_xlsx(
                    workbook_path,
                    require_farm_column=True,
                )

    def test_describe_distinct_germplasm_names_from_xlsx_groups_uploaded_rows_by_name_and_farm(self) -> None:
        from openpyxl import Workbook
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir_name:
            workbook_path = Path(tmp_dir_name) / "with_farm.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Farm", "Education"])
            worksheet.append(["CKDHH211274", "Farm A", "Primary"])
            worksheet.append(["CKDHH211274", "Farm A", "Post-secondary"])
            worksheet.append(["CKDHH211274", "Farm B", "Primary"])
            workbook.save(workbook_path)
            workbook.close()

            description = describe_distinct_germplasm_names_from_xlsx(
                workbook_path,
                require_farm_column=True,
            )

            self.assertEqual(description["selected_id_header"], "Farm")
            self.assertEqual(
                description["germplasm_names"],
                ["CKDHH211274 - Farm A", "CKDHH211274 - Farm B"],
            )
            farm_a = description["germplasm_profiles_by_name"]["CKDHH211274 - Farm A"]
            self.assertEqual(farm_a["row_count"], 2)
            self.assertEqual(len(farm_a["internal_row_ids"]), 2)

    def test_validate_required_workbook_headers_requires_name_and_farm(self) -> None:
        from openpyxl import Workbook
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir_name:
            workbook_path = Path(tmp_dir_name) / "missing_name.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Farm", "Education"])
            worksheet.append(["Farm A", "Primary"])
            workbook.save(workbook_path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "required Name column"):
                validate_required_workbook_headers(workbook_path, ("Name", "Farm"))

    def test_validate_required_workbook_headers_returns_resolved_headers(self) -> None:
        from openpyxl import Workbook
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir_name:
            workbook_path = Path(tmp_dir_name) / "required_headers.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append([" name ", " farm ", "Education"])
            worksheet.append(["Pioneer", "Farm A", "Primary"])
            workbook.save(workbook_path)
            workbook.close()

            resolved = validate_required_workbook_headers(workbook_path, ("Name", "Farm"))

            self.assertEqual(resolved, {"Name": "name", "Farm": "farm"})


if __name__ == "__main__":
    unittest.main()
