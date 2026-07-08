from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook

from cimmyt_app.ce_pipeline.phase01.phase1 import read_sheet_headers_and_rows, write_xlsx_rows
from cimmyt_app.ce_pipeline.prediction.prediction import (
    PREDICTION_PROFILE_LABEL_HEADER,
    PREDICTION_PROFILE_ABOVE_MEAN_HEADER,
    PREDICTION_PROFILE_KEY_HEADER,
    PREDICTION_PREDICTED_COLUMN,
    _aggregate_manual_grid_prediction_dataframe_for_display,
    _build_lightweight_manual_grid_geojson_dataframe,
    _normalize_manual_grid_display_columns,
    _prepare_prediction_output_dataframe_for_display,
    _sort_prediction_output_files,
    _postprocess_multi_profile_prediction_results,
    _should_preserve_manual_grid_prediction_rows,
    build_prediction_phase02_workbook,
    build_required_prediction_headers,
    update_saved_model_prediction_metadata,
    validate_saved_model_prediction_workbook,
)


class CePredictionDefinitionTest(unittest.TestCase):
    def test_aggregate_manual_grid_prediction_dataframe_for_display_keeps_one_row_per_cell(self) -> None:
        prediction_df = pd.DataFrame(
            [
                {
                    "_GPS coordinates_latitude": -0.5,
                    "_GPS coordinates_longitude": 37.1,
                    "Forecast grid cell id": "Manual bounds:R1C1",
                    "Forecast grid cell index": 1,
                    "Farm": "Farm A",
                    PREDICTION_PREDICTED_COLUMN: 5.8,
                },
                {
                    "_GPS coordinates_latitude": -0.5,
                    "_GPS coordinates_longitude": 37.1,
                    "Forecast grid cell id": "Manual bounds:R1C1",
                    "Forecast grid cell index": 1,
                    "Farm": "Farm B",
                    PREDICTION_PREDICTED_COLUMN: 6.2,
                },
                {
                    "_GPS coordinates_latitude": -0.4,
                    "_GPS coordinates_longitude": 37.2,
                    "Forecast grid cell id": "Manual bounds:R1C2",
                    "Forecast grid cell index": 2,
                    "Farm": "Farm C",
                    PREDICTION_PREDICTED_COLUMN: 6.0,
                },
            ]
        )

        aggregated_df = _aggregate_manual_grid_prediction_dataframe_for_display(prediction_df)

        self.assertEqual(len(aggregated_df.index), 2)
        first_cell = aggregated_df[aggregated_df["Forecast grid cell id"] == "Manual bounds:R1C1"].iloc[0]
        self.assertEqual(first_cell["Farm"], "Farm B")
        self.assertEqual(float(first_cell[PREDICTION_PREDICTED_COLUMN]), 6.2)
        self.assertEqual(int(first_cell["Prediction cell row count"]), 2)

    def test_build_lightweight_manual_grid_geojson_dataframe_keeps_map_fields(self) -> None:
        prediction_df = pd.DataFrame(
            [
                {
                    "_GPS coordinates_latitude": -0.5,
                    "_GPS coordinates_longitude": 37.1,
                    "Forecast grid cell id": "Manual bounds:R1C1",
                    "Forecast grid cell index": 1,
                    "Name": "G1",
                    "Farm": "Farm A",
                    PREDICTION_PREDICTED_COLUMN: 5.8,
                    PREDICTION_PROFILE_KEY_HEADER: "profile-1",
                    PREDICTION_PROFILE_LABEL_HEADER: "Profile 1",
                    PREDICTION_PROFILE_ABOVE_MEAN_HEADER: "Yes",
                    "Very Heavy Column": "drop me",
                }
            ]
        )

        lightweight_df = _build_lightweight_manual_grid_geojson_dataframe(
            prediction_df,
            selected_id_header="Farm",
        )

        self.assertEqual(
            lightweight_df.columns.tolist(),
            [
                "_GPS coordinates_latitude",
                "_GPS coordinates_longitude",
                "Forecast grid cell id",
                "Forecast grid cell index",
                "Name",
                "Farm",
                PREDICTION_PREDICTED_COLUMN,
                PREDICTION_PROFILE_KEY_HEADER,
                PREDICTION_PROFILE_LABEL_HEADER,
                PREDICTION_PROFILE_ABOVE_MEAN_HEADER,
            ],
        )
        self.assertNotIn("Very Heavy Column", lightweight_df.columns)

    def test_normalize_manual_grid_display_columns_uses_forecast_locality_as_grid_cell_id(self) -> None:
        prediction_df = pd.DataFrame(
            [
                {
                    "_GPS coordinates_latitude": -0.5,
                    "_GPS coordinates_longitude": 37.1,
                    "Forecast locality": "Manual bounds:R1C1",
                    "Farm": "Farm A",
                    PREDICTION_PREDICTED_COLUMN: 5.8,
                }
            ]
        )

        normalized_df = _normalize_manual_grid_display_columns(
            prediction_df,
            manual_bbox_grid={
                "cells": [
                    {
                        "grid_cell_id": "Manual bounds:R1C1",
                        "grid_cell_index": 1,
                    }
                ]
            },
        )

        self.assertEqual(normalized_df.loc[0, "Forecast grid cell id"], "Manual bounds:R1C1")
        self.assertEqual(normalized_df.loc[0, "Forecast grid cell index"], 1)

    def test_prepare_prediction_output_dataframe_for_display_reorders_sorts_and_drops_report_only_fields(self) -> None:
        prediction_df = pd.DataFrame(
            [
                {
                    "Country": "Kenya",
                    "Farm": "Farm B",
                    "_GPS coordinates_latitude": -0.4,
                    "_GPS coordinates_longitude": 37.2,
                    PREDICTION_PREDICTED_COLUMN: 6.1,
                    PREDICTION_PROFILE_KEY_HEADER: "profile-2",
                    PREDICTION_PROFILE_LABEL_HEADER: "Profile 2",
                },
                {
                    "Country": "Kenya",
                    "Farm": "Farm A",
                    "_GPS coordinates_latitude": -0.5,
                    "_GPS coordinates_longitude": 37.1,
                    PREDICTION_PREDICTED_COLUMN: 5.8,
                    PREDICTION_PROFILE_KEY_HEADER: "profile-1",
                    PREDICTION_PROFILE_LABEL_HEADER: "Profile 1",
                },
            ]
        )

        prepared_df = _prepare_prediction_output_dataframe_for_display(
            prediction_df,
            latitude_column="_GPS coordinates_latitude",
            longitude_column="_GPS coordinates_longitude",
            selected_id_header="Farm",
        )

        self.assertEqual(prepared_df.columns.tolist()[0], "Farm")
        self.assertNotIn("Country", prepared_df.columns)
        self.assertNotIn(PREDICTION_PROFILE_KEY_HEADER, prepared_df.columns)
        self.assertNotIn(PREDICTION_PROFILE_LABEL_HEADER, prepared_df.columns)
        self.assertEqual(prepared_df.iloc[0]["Farm"], "Farm A")
        self.assertEqual(float(prepared_df.iloc[0]["_GPS coordinates_longitude"]), 37.1)

    def test_write_xlsx_rows_round_trips_manual_bbox_ready_values(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            output_file = Path(tmp_dir_name) / "streamed.xlsx"
            headers = ["Farm", "Forecast grid cell id", "_GPS coordinates_latitude", "_GPS coordinates_longitude"]
            rows = [
                ("Farm A", "Manual bounds:R1C1", -0.12345, 37.98765),
                ("Farm A", "Manual bounds:R1C2", "", ""),
            ]

            write_xlsx_rows(output_file, headers, rows, len(rows))

            loaded_headers, loaded_rows = read_sheet_headers_and_rows(output_file)
            self.assertEqual(loaded_headers, headers)
            self.assertEqual(len(loaded_rows), 2)
            self.assertEqual(loaded_rows[0]["Forecast grid cell id"], "Manual bounds:R1C1")
            self.assertEqual(loaded_rows[0]["_GPS coordinates_latitude"], "-0.12345")
            self.assertEqual(loaded_rows[1]["_GPS coordinates_longitude"], "")

    def test_build_required_prediction_headers_collects_stepper_columns(self) -> None:
        selection_summary = {
            "initial_settings": {
                "longitude_column": "Longitude",
                "latitude_column": "Latitude",
                "planting_date_column": "Planting",
                "harvesting_date_column": "Harvesting",
                "soil_texture_column": "SoilTexture",
                "soil_depth_column": "SoilDepth",
            },
            "divisions": {
                "germplams_identifiers": ["Name"],
                "DG": ["DG1", "DG2"],
                "categorical_data": ["Management"],
                "cuantitative_data": ["Nitrogen"],
            },
        }

        headers = build_required_prediction_headers(selection_summary)

        self.assertEqual(
            headers,
            [
                "Longitude",
                "Latitude",
                "Planting",
                "Harvesting",
                "SoilTexture",
                "SoilDepth",
                "Name",
                "DG1",
                "DG2",
                "Management",
                "Nitrogen",
            ],
        )

    def test_update_saved_model_prediction_metadata_copies_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            model_dir = tmp_dir / "model"
            model_dir.mkdir(parents=True, exist_ok=True)
            selection_summary_file = tmp_dir / "selection_summary.json"
            selection_summary = {
                "initial_settings": {"longitude_column": "Longitude", "latitude_column": "Latitude"},
                "divisions": {"germplams_identifiers": ["Name"]},
            }
            selection_summary_file.write_text(json.dumps(selection_summary), encoding="utf-8")
            selected_fields_csv = tmp_dir / "phase04_selected_fields.csv"
            selected_fields_csv.write_text("Name,Longitude\nA,1\n", encoding="utf-8")
            phase04_workbook = tmp_dir / "phase04.xlsx"
            phase04_workbook.write_bytes(b"xlsx")
            training_csv = tmp_dir / "phase04_training_input.csv"
            training_csv.write_text("feature_a\n1\n", encoding="utf-8")
            normalization_stats_file = tmp_dir / "phase04_normalization_stats.json"
            normalization_stats_file.write_text(
                json.dumps({"Nitrogen": {"apply_log": False, "mean": 10.0, "std": 2.0}}),
                encoding="utf-8",
            )

            metadata = update_saved_model_prediction_metadata(
                model_dir=model_dir,
                metadata={"model_id": "model-1"},
                selection_summary_file=selection_summary_file,
                phase04_selected_fields_csv=selected_fields_csv,
                phase04_workbook=phase04_workbook,
                phase04_training_input_csv=training_csv,
                phase04_normalization_stats_file=normalization_stats_file,
            )

            self.assertTrue((model_dir / "selection_summary.json").exists())
            self.assertTrue((model_dir / "phase04_selected_fields.csv").exists())
            self.assertTrue((model_dir / "phase04_normalization_stats.json").exists())
            self.assertEqual(metadata["required_prediction_headers"], ["Longitude", "Latitude", "Name"])
            self.assertEqual(metadata["phase04_normalization_stats"]["Nitrogen"]["std"], 2.0)

    def test_build_prediction_phase02_workbook_creates_initial_setting_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            workbook_path = tmp_dir / "uploaded.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Farm", "Management", "Nitrogen", "Planting", "Harvesting", "Yield"])
            worksheet.append(["PIONEER", "Farm A", "Hand", 45, "2024-05-01", "2024-09-15", 5.6])
            workbook.save(workbook_path)
            workbook.close()

            selection_summary = {
                "initial_settings": {
                    "target_column": "Yield",
                    "longitude_column": "Longitude",
                    "latitude_column": "Latitude",
                    "planting_date_column": "Planting",
                    "harvesting_date_column": "Harvesting",
                    "soil_texture_column": "SoilTexture",
                    "soil_depth_column": "SoilDepth",
                },
                "divisions": {
                    "germplams_identifiers": ["Name"],
                    "categorical_data": ["Management"],
                    "cuantitative_data": ["Nitrogen"],
                },
            }
            output_file = tmp_dir / "phase02prediction.xlsx"

            payload = build_prediction_phase02_workbook(
                workbook_path,
                selection_summary=selection_summary,
                output_file=output_file,
                forecast_planting_date="2026-05-01",
                forecast_harvesting_date="2026-09-15",
                regional_bounds_label="Manual bounds",
                regional_bounds_latitude_min=10.0,
                regional_bounds_latitude_max=14.0,
                regional_bounds_longitude_min=20.0,
                regional_bounds_longitude_max=28.0,
                selected_germplasm_names=["PIONEER"],
            )

            self.assertTrue(output_file.exists())
            headers, rows = read_sheet_headers_and_rows(output_file)
            self.assertIn("Longitude", headers)
            self.assertIn("Latitude", headers)
            self.assertIn("Planting", headers)
            self.assertIn("Harvesting", headers)
            self.assertEqual(rows[0]["Longitude"], "24.0")
            self.assertEqual(rows[0]["Latitude"], "12.0")
            self.assertEqual(rows[0]["Planting"], "2024-05-01")
            self.assertEqual(rows[0]["Harvesting"], "2024-09-15")
            self.assertEqual(payload["output_row_count"], 1)
            self.assertTrue(payload["uses_source_row_dates"])


    def test_build_prediction_phase02_workbook_detects_multi_profile_duplicates(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            workbook_path = tmp_dir / "uploaded_profiles.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Farm", "Education/Training of the farmer", "Planting", "Harvesting", "Yield"])
            worksheet.append(["CKDHH211040", "Farm A", "Post-secondary", "2024-05-01", "2024-09-01", 6.2])
            worksheet.append(["CKDHH211040", "Farm B", "Primary", "2024-05-03", "2024-09-04", 5.1])
            workbook.save(workbook_path)
            workbook.close()

            selection_summary = {
                "initial_settings": {
                    "target_column": "Yield",
                    "longitude_column": "Longitude",
                    "latitude_column": "Latitude",
                    "planting_date_column": "Planting",
                    "harvesting_date_column": "Harvesting",
                    "soil_texture_column": "SoilTexture",
                    "soil_depth_column": "SoilDepth",
                },
                "divisions": {
                    "germplams_identifiers": ["Name"],
                    "categorical_data": ["Education/Training of the farmer"],
                    "cuantitative_data": [],
                },
            }
            output_file = tmp_dir / "phase02prediction_profiles.xlsx"

            payload = build_prediction_phase02_workbook(
                workbook_path,
                selection_summary=selection_summary,
                output_file=output_file,
                forecast_planting_date="",
                forecast_harvesting_date="",
                regional_bounds_label="Manual bounds",
                regional_bounds_latitude_min=10.0,
                regional_bounds_latitude_max=14.0,
                regional_bounds_longitude_min=20.0,
                regional_bounds_longitude_max=28.0,
                selected_germplasm_names=["CKDHH211040"],
                selected_id_header="Education/Training of the farmer",
            )

            headers, rows = read_sheet_headers_and_rows(output_file)
            self.assertIn(PREDICTION_PROFILE_LABEL_HEADER, headers)
            self.assertEqual(len(rows), 2)
            self.assertTrue(payload["prediction_profile_analysis"]["multi_profile_mode"])
            self.assertEqual(payload["prediction_profile_analysis"]["profile_count"], 2)
            self.assertEqual(rows[0]["Planting"], "2024-05-01")
            self.assertEqual(rows[1]["Planting"], "2024-05-03")

    def test_build_prediction_phase02_workbook_keeps_single_profile_without_selected_id_header(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            workbook_path = tmp_dir / "uploaded_single_profile.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Farm", "Management", "Planting", "Harvesting", "Yield"])
            worksheet.append(["PIONEER", "Farm A", "Hand", "2024-05-01", "2024-09-15", 5.6])
            worksheet.append(["PIONEER", "Farm B", "Machine", "2024-05-03", "2024-09-17", 5.8])
            workbook.save(workbook_path)
            workbook.close()

            selection_summary = {
                "initial_settings": {
                    "target_column": "Yield",
                    "longitude_column": "Longitude",
                    "latitude_column": "Latitude",
                    "planting_date_column": "Planting",
                    "harvesting_date_column": "Harvesting",
                    "soil_texture_column": "SoilTexture",
                    "soil_depth_column": "SoilDepth",
                },
                "divisions": {
                    "germplams_identifiers": ["Name"],
                    "categorical_data": ["Management"],
                    "cuantitative_data": [],
                },
            }
            output_file = tmp_dir / "phase02prediction_single_profile.xlsx"

            payload = build_prediction_phase02_workbook(
                workbook_path,
                selection_summary=selection_summary,
                output_file=output_file,
                forecast_planting_date="",
                forecast_harvesting_date="",
                regional_bounds_label="Manual bounds",
                regional_bounds_latitude_min=10.0,
                regional_bounds_latitude_max=14.0,
                regional_bounds_longitude_min=20.0,
                regional_bounds_longitude_max=28.0,
                selected_germplasm_names=["PIONEER"],
                selected_id_header="",
            )

            self.assertEqual(payload["prediction_profile_analysis"]["selected_id_header"], "")
            self.assertFalse(payload["prediction_profile_analysis"]["multi_profile_mode"])
            self.assertEqual(payload["prediction_profile_analysis"]["profile_count"], 1)
            self.assertEqual(len(payload["prediction_profile_analysis"]["profile_summaries"]), 1)

    def test_build_prediction_phase02_workbook_fills_missing_selected_id_for_repeated_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            workbook_path = tmp_dir / "uploaded_missing_ids.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Name", "Farm", "Nursery Id", "Management", "Planting", "Harvesting", "Yield"])
            worksheet.append(["PIONEER", "Farm A", "", "Hand", "2024-05-01", "2024-09-15", 5.6])
            worksheet.append(["PIONEER", "Farm A", "null", "Machine", "2024-05-03", "2024-09-17", 5.8])
            worksheet.append(["SINGLE", "Farm B", "", "Hand", "2024-05-02", "2024-09-16", 4.9])
            workbook.save(workbook_path)
            workbook.close()

            selection_summary = {
                "initial_settings": {
                    "target_column": "Yield",
                    "longitude_column": "Longitude",
                    "latitude_column": "Latitude",
                    "planting_date_column": "Planting",
                    "harvesting_date_column": "Harvesting",
                    "soil_texture_column": "SoilTexture",
                    "soil_depth_column": "SoilDepth",
                },
                "divisions": {
                    "germplams_identifiers": ["Name"],
                    "categorical_data": ["Management"],
                    "cuantitative_data": [],
                },
            }
            output_file = tmp_dir / "phase02prediction_missing_ids.xlsx"

            build_prediction_phase02_workbook(
                workbook_path,
                selection_summary=selection_summary,
                output_file=output_file,
                forecast_planting_date="",
                forecast_harvesting_date="",
                regional_bounds_label="Manual bounds",
                regional_bounds_latitude_min=10.0,
                regional_bounds_latitude_max=14.0,
                regional_bounds_longitude_min=20.0,
                regional_bounds_longitude_max=28.0,
                selected_germplasm_names=["PIONEER", "SINGLE"],
                selected_id_header="Nursery Id",
            )

            from cimmyt_app.ce_pipeline.phase01.phase1 import read_sheet_headers_and_rows
            _, rows = read_sheet_headers_and_rows(output_file)
            pioneer_rows = [row for row in rows if row["Name"] == "PIONEER"]
            single_rows = [row for row in rows if row["Name"] == "SINGLE"]

            self.assertEqual([row["Nursery Id"] for row in pioneer_rows], ["1", "2"])
            self.assertEqual(single_rows[0]["Nursery Id"], "")

    def test_postprocess_multi_profile_prediction_results_keeps_all_profiles(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            geojson_file = tmp_dir / "prediction.geojson"
            prediction_csv = tmp_dir / "prediction.csv"
            prediction_xlsx = tmp_dir / "prediction.xlsx"

            geojson = {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "geometry": {"type": "Point", "coordinates": [20.0, 10.0]},
                        "properties": {
                            PREDICTION_PROFILE_KEY_HEADER: "Education=Post-secondary",
                            PREDICTION_PROFILE_LABEL_HEADER: "Education/Training of the farmer: Post-secondary",
                            PREDICTION_PREDICTED_COLUMN: 6.8,
                        },
                    },
                    {
                        "type": "Feature",
                        "geometry": {"type": "Point", "coordinates": [21.0, 11.0]},
                        "properties": {
                            PREDICTION_PROFILE_KEY_HEADER: "Education=Primary",
                            PREDICTION_PROFILE_LABEL_HEADER: "Education/Training of the farmer: Primary",
                            PREDICTION_PREDICTED_COLUMN: 5.2,
                        },
                    },
                ],
            }
            geojson_file.write_text(json.dumps(geojson), encoding="utf-8")
            prediction_csv.write_text(
                ",".join([PREDICTION_PROFILE_KEY_HEADER, PREDICTION_PROFILE_LABEL_HEADER, PREDICTION_PREDICTED_COLUMN]) + "\n"
                + "Education=Post-secondary,Education/Training of the farmer: Post-secondary,6.8\n"
                + "Education=Primary,Education/Training of the farmer: Primary,5.2\n",
                encoding="utf-8",
            )
            prediction_xlsx.write_bytes(b"xlsx")

            filtered_geojson, summary = _postprocess_multi_profile_prediction_results(
                geojson=geojson,
                geojson_file=geojson_file,
                prediction_xlsx=prediction_xlsx,
                prediction_csv=prediction_csv,
                profile_metadata={
                    "multi_profile_mode": True,
                    "target_mean_threshold": 6.0,
                    "profile_summaries": [
                        {"key": "Education=Post-secondary", "label": "Education/Training of the farmer: Post-secondary"},
                        {"key": "Education=Primary", "label": "Education/Training of the farmer: Primary"},
                    ],
                },
            )

            self.assertEqual(len(filtered_geojson["features"]), 2)
            self.assertEqual(summary["displayed_profile_count"], 1)
            above_mean_values = [
                feature["properties"][PREDICTION_PROFILE_ABOVE_MEAN_HEADER]
                for feature in filtered_geojson["features"]
            ]
            self.assertEqual(above_mean_values, ["Yes", "No"])

    def test_should_preserve_manual_grid_prediction_rows_for_single_selected_name(self) -> None:
        self.assertTrue(
            _should_preserve_manual_grid_prediction_rows(
                climate_scope="regional_manual",
                selected_id_header="",
                selected_germplasm_names=["PIONEER"],
            )
        )

    def test_should_not_preserve_manual_grid_prediction_rows_for_multi_name_or_identifier(self) -> None:
        self.assertFalse(
            _should_preserve_manual_grid_prediction_rows(
                climate_scope="regional_manual",
                selected_id_header="Farm",
                selected_germplasm_names=["PIONEER"],
            )
        )
        self.assertFalse(
            _should_preserve_manual_grid_prediction_rows(
                climate_scope="regional_manual",
                selected_id_header="",
                selected_germplasm_names=["PIONEER", "SINGLE"],
            )
        )
        self.assertFalse(
            _should_preserve_manual_grid_prediction_rows(
                climate_scope="country",
                selected_id_header="",
                selected_germplasm_names=["PIONEER"],
            )
        )

    def test_sort_prediction_output_files_places_selected_id_first(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            prediction_csv = tmp_dir / "prediction.csv"
            prediction_xlsx = tmp_dir / "prediction.xlsx"
            prediction_df = pd.DataFrame(
                [
                    {
                        "Country": "Kenya",
                        "Longitude": 35.2,
                        "Latitude": -1.4,
                        "Name": "A",
                        "Nursery Id": "ID-2",
                        PREDICTION_PROFILE_KEY_HEADER: "k2",
                        PREDICTION_PROFILE_LABEL_HEADER: "A - ID-2",
                        "Prediction Profile Columns": "Name|Nursery Id",
                        "Prediction Target Mean Threshold": 4.2,
                        PREDICTION_PREDICTED_COLUMN: 6.7,
                    },
                    {
                        "Country": "Kenya",
                        "Longitude": 35.2,
                        "Latitude": -1.4,
                        "Name": "A",
                        "Nursery Id": "ID-1",
                        PREDICTION_PROFILE_KEY_HEADER: "k1",
                        PREDICTION_PROFILE_LABEL_HEADER: "A - ID-1",
                        "Prediction Profile Columns": "Name|Nursery Id",
                        "Prediction Target Mean Threshold": 4.2,
                        PREDICTION_PREDICTED_COLUMN: 7.2,
                    },
                ]
            )
            prediction_df.to_excel(prediction_xlsx, index=False)
            prediction_df.to_csv(prediction_csv, index=False)

            _sort_prediction_output_files(
                prediction_xlsx=prediction_xlsx,
                prediction_csv=prediction_csv,
                latitude_column="Latitude",
                longitude_column="Longitude",
                selected_id_header="Nursery Id",
            )

            sorted_df = pd.read_excel(prediction_xlsx)
            self.assertEqual(sorted_df.columns[0], "Nursery Id")
            self.assertNotIn("Country", sorted_df.columns)
            self.assertNotIn(PREDICTION_PROFILE_KEY_HEADER, sorted_df.columns)
            self.assertEqual(sorted_df.iloc[0]["Nursery Id"], "ID-1")

    def test_validate_saved_model_prediction_workbook_requires_saved_definition_headers(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir_name:
            tmp_dir = Path(tmp_dir_name)
            workbook_path = tmp_dir / "prediction_input.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Latitude", "Longitude", "Name"])
            worksheet.append([1.0, 2.0, "PIONEER"])
            workbook.save(workbook_path)
            workbook.close()

            registered_model = SimpleNamespace(
                metadata={
                    "selection_summary": {
                        "initial_settings": {
                            "longitude_column": "Longitude",
                            "latitude_column": "Latitude",
                            "planting_date_column": "Planting",
                            "harvesting_date_column": "Harvesting",
                            "soil_texture_column": "SoilTexture",
                            "soil_depth_column": "SoilDepth",
                        },
                        "divisions": {
                            "germplams_identifiers": ["Name"],
                            "categorical_data": ["Management"],
                            "cuantitative_data": ["Nitrogen"],
                        },
                    }
                },
                model_dir=tmp_dir / "model",
            )

            with patch(
                "cimmyt_app.ce_pipeline.prediction.prediction.get_registered_model",
                return_value=registered_model,
            ):
                with self.assertRaisesRegex(ValueError, "Missing required columns"):
                    validate_saved_model_prediction_workbook(
                        workbook_path,
                        selected_model_id="model-1",
                        climate_scope="regional_manual",
                    )


if __name__ == "__main__":
    unittest.main()
