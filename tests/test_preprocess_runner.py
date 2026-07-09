from __future__ import annotations

import json
import os
import shutil
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from preprocess import ea_pipeline
from preprocess.runner import (
    MANUAL_BBOX_CLEAR_COLUMNS,
    POINT_SAVED_MODEL_BASE_OVERRIDE_COLUMNS,
    clear_normalized_workbook_columns,
    is_saved_model_manual_bbox_prediction,
    list_distinct_germplasm_names,
    merge_nasa_cache_files,
    project_normalized_test_workbook_across_all_markers,
    run_preprocess_pipeline,
)
from config_env import get_prediction_bridge_dev_row_limit


APP_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = APP_ROOT.parent
TEST_WORKBOOK = PROJECT_ROOT / "cimmyt_app" / "template" / "test.xlsx"
REFERENCE_PHASE06 = (
    PROJECT_ROOT
    / "cimmyt_app"
    / "template"
    / "test_phase06_output.xlsx"
)
REFERENCE_NASA_CACHE = (
    PROJECT_ROOT
    / "cimmyt_app"
    / ".cache"
    / "nasa_power_cache.json"
)


def normalize_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip()


def parse_numeric_text(value: str) -> float | None:
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def load_sheet_rows(workbook_path: Path) -> tuple[list[str], list[tuple[str, ...]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [normalize_cell_value(value) for value in next(iterator)]
        rows = [
            tuple(normalize_cell_value(value) for value in row)
            for row in iterator
        ]
        return headers, rows
    finally:
        workbook.close()


def build_legacy_plot_no_workbook(source_workbook: Path, destination_path: Path) -> None:
    source = load_workbook(source_workbook, read_only=True, data_only=True)
    legacy = Workbook(write_only=True)
    try:
        source_sheet = source.worksheets[0]
        output_sheet = legacy.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [normalize_cell_value(value) for value in next(rows)]
        plot_index = headers.index("Plot")
        legacy_headers = headers[: plot_index + 1] + ["Plot No."] + headers[plot_index + 1 :]
        output_sheet.append(legacy_headers)

        for row in rows:
            row_values = list(row)
            plot_value = row_values[plot_index]
            legacy_row = row_values[: plot_index + 1] + [plot_value] + row_values[plot_index + 1 :]
            output_sheet.append(legacy_row)

        legacy.save(destination_path)
    finally:
        source.close()


def build_workbook_with_extra_id_column(source_workbook: Path, destination_path: Path) -> None:
    source = load_workbook(source_workbook, read_only=True, data_only=True)
    modified = Workbook(write_only=True)
    try:
        source_sheet = source.worksheets[0]
        output_sheet = modified.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [normalize_cell_value(value) for value in next(rows)]
        output_sheet.append(["id", *headers])

        for index, row in enumerate(rows, start=1):
            output_sheet.append([index, *list(row)])

        modified.save(destination_path)
    finally:
        source.close()


def build_workbook_with_dg_columns(source_workbook: Path, destination_path: Path) -> None:
    source = load_workbook(source_workbook, read_only=True, data_only=True)
    modified = Workbook(write_only=True)
    try:
        source_sheet = source.worksheets[0]
        output_sheet = modified.create_sheet(source_sheet.title)
        rows = source_sheet.iter_rows(values_only=True)
        headers = [normalize_cell_value(value) for value in next(rows)]
        name_index = headers.index("Name")
        unique_names: list[str] = []
        seen: set[str] = set()
        cached_rows: list[list[object]] = []
        for row in rows:
            row_values = list(row)
            cached_rows.append(row_values)
            name = normalize_cell_value(row_values[name_index])
            if name and name not in seen:
                seen.add(name)
                unique_names.append(name)
        dg_headers = [f"DG{index}" for index in range(1, len(unique_names) + 1)]
        output_sheet.append([*headers, *dg_headers])
        for row_values in cached_rows:
            name = normalize_cell_value(row_values[name_index])
            encoded = [1 if name == unique_name else 0 for unique_name in unique_names]
            output_sheet.append([*row_values, *encoded])
        modified.save(destination_path)
    finally:
        source.close()


class TestPreprocessRunner(unittest.TestCase):
    maxDiff = None

    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_dir = tempfile.TemporaryDirectory(prefix="app_preprocess_test_")
        cls.run_dir = Path(cls.temp_dir.name) / "preprocess"
        phase02_dir = cls.run_dir / "phase02_climate_enriched"
        phase02_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")
        cls.outputs = run_preprocess_pipeline(TEST_WORKBOOK, cls.run_dir)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp_dir.cleanup()

    def test_run_preprocess_pipeline_matches_reference_phase06(self) -> None:
        self.assertTrue(
            self.outputs.phase06_xlsx.exists(),
            "The preprocess phase06 workbook was not created.",
        )
        self.assertTrue(
            self.outputs.summary_file.exists(),
            "The preprocess summary file was not created.",
        )

        expected_headers, expected_rows = load_sheet_rows(REFERENCE_PHASE06)
        actual_headers, actual_rows = load_sheet_rows(self.outputs.phase06_xlsx)

        self.assertEqual(actual_headers, expected_headers)
        self.assertEqual(actual_rows, expected_rows)

    def test_run_preprocess_pipeline_excludes_rows_whose_name_maps_to_delete(self) -> None:
        headers, rows = load_sheet_rows(self.outputs.phase06_xlsx)
        name_index = headers.index("Name")
        delete_rows = [row for row in rows if row[name_index] == "DELETE"]

        self.assertEqual(delete_rows, [])

    def test_run_preprocess_pipeline_writes_expected_summary(self) -> None:
        summary = json.loads(self.outputs.summary_file.read_text(encoding="utf-8"))

        self.assertEqual(Path(summary["input_workbook"]), self.outputs.input_workbook)
        self.assertEqual(Path(summary["phase06_xlsx"]), self.outputs.phase06_xlsx)
        self.assertEqual(
            [task["phase"] for task in summary["task_report"]],
            ["phase00", "phase01", "phase02", "phase03", "phase04", "phase05", "phase06"],
        )
        self.assertIn("preprocess_log", summary)
        self.assertIn("overview", summary["preprocess_log"])
        self.assertIn("operations", summary["preprocess_log"])
        self.assertIn("nasa_cache_fresh_queries_this_run", summary["preprocess_log"]["overview"])
        self.assertIn("nasa_cache_hits", summary["preprocess_log"]["overview"])
        self.assertIn("nasa_cache_hits", summary["preprocess_log"]["climate"])
        self.assertIn("nasa_cache_misses", summary["preprocess_log"]["climate"])

    def test_run_preprocess_pipeline_accepts_legacy_plot_no_upload(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_legacy_") as temp_dir:
            legacy_workbook = Path(temp_dir) / "legacy_test.xlsx"
            run_dir = Path(temp_dir) / "preprocess"
            phase02_dir = run_dir / "phase02_climate_enriched"
            build_legacy_plot_no_workbook(TEST_WORKBOOK, legacy_workbook)
            phase02_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")

            outputs = run_preprocess_pipeline(legacy_workbook, run_dir)

            normalized_headers, _ = load_sheet_rows(outputs.input_workbook)
            self.assertNotIn("Plot No.", normalized_headers)
            self.assertEqual(normalized_headers, load_sheet_rows(TEST_WORKBOOK)[0])

    def test_run_preprocess_pipeline_skips_strict_header_validation_when_env_disabled(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_validation_off_") as temp_dir:
            modified_workbook = Path(temp_dir) / "test_with_id.xlsx"
            run_dir = Path(temp_dir) / "preprocess"
            phase02_dir = run_dir / "phase02_climate_enriched"
            build_workbook_with_extra_id_column(TEST_WORKBOOK, modified_workbook)
            phase02_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")
            previous_value = os.environ.get("APP_ENABLE_PREPROCESS_VALIDATION")
            os.environ["APP_ENABLE_PREPROCESS_VALIDATION"] = "0"

            try:
                outputs = run_preprocess_pipeline(modified_workbook, run_dir)
            finally:
                if previous_value is None:
                    os.environ.pop("APP_ENABLE_PREPROCESS_VALIDATION", None)
                else:
                    os.environ["APP_ENABLE_PREPROCESS_VALIDATION"] = previous_value

            normalized_headers, _ = load_sheet_rows(outputs.input_workbook)
            self.assertEqual(normalized_headers, load_sheet_rows(TEST_WORKBOOK)[0])

    def test_run_preprocess_pipeline_accepts_and_preserves_dg_prefixed_columns(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_dg_") as temp_dir:
            dg_workbook = Path(temp_dir) / "test_with_dg.xlsx"
            run_dir = Path(temp_dir) / "preprocess"
            phase02_dir = run_dir / "phase02_climate_enriched"
            build_workbook_with_dg_columns(TEST_WORKBOOK, dg_workbook)
            phase02_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")

            outputs = run_preprocess_pipeline(dg_workbook, run_dir)

            normalized_headers, _ = load_sheet_rows(outputs.input_workbook)
            phase06_headers, _ = load_sheet_rows(outputs.phase06_xlsx)
            dg_headers = [header for header in normalized_headers if header.startswith("DG")]
            self.assertTrue(dg_headers)
            self.assertEqual(dg_headers, [header for header in phase06_headers if header.startswith("DG")])

    def test_merge_nasa_cache_files_preserves_existing_entries_and_adds_new_ones(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_cache_merge_") as temp_dir:
            temp_path = Path(temp_dir)
            source_cache = temp_path / "source_cache.json"
            destination_cache = temp_path / "destination_cache.json"
            source_cache.write_text(
                json.dumps(
                    {
                        "latA|lonA|2024-01-01|2024-02-01": {"T2M": {"20240101": 25.1}},
                        "latB|lonB|2024-03-01|2024-04-01": {"T2M": {"20240301": 26.2}},
                    }
                ),
                encoding="utf-8",
            )
            destination_cache.write_text(
                json.dumps(
                    {
                        "latA|lonA|2024-01-01|2024-02-01": {"T2M": {"20240101": 25.1}},
                        "latC|lonC|2024-05-01|2024-06-01": {"T2M": {"20240501": 27.3}},
                    }
                ),
                encoding="utf-8",
            )

            merge_nasa_cache_files(source_cache, destination_cache)

            merged_payload = json.loads(destination_cache.read_text(encoding="utf-8"))
            self.assertEqual(len(merged_payload), 3)
            self.assertIn("latA|lonA|2024-01-01|2024-02-01", merged_payload)
            self.assertIn("latB|lonB|2024-03-01|2024-04-01", merged_payload)
            self.assertIn("latC|lonC|2024-05-01|2024-06-01", merged_payload)

    def test_list_distinct_germplasm_names_returns_name_values_from_uploaded_workbook(self) -> None:
        names = list_distinct_germplasm_names(TEST_WORKBOOK)

        self.assertIn("CIM-RT-EAPP1-E -001", names)
        self.assertIn("CIM-RT-EAPP1-E -007", names)
        self.assertEqual(names, sorted(names))

    def test_run_preprocess_pipeline_filters_rows_by_selected_germplasm_names(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_filtered_") as temp_dir:
            run_dir = Path(temp_dir) / "preprocess"
            phase02_dir = run_dir / "phase02_climate_enriched"
            phase02_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")

            outputs = run_preprocess_pipeline(
                TEST_WORKBOOK,
                run_dir,
                selected_germplasm_names=["CIM-RT-EAPP1-E -001", "CIM-RT-EAPP1-E -007"],
            )

            headers, rows = load_sheet_rows(outputs.input_workbook)
            name_index = headers.index("Name")
            names = {row[name_index] for row in rows if row[name_index]}

            self.assertEqual(names, {"CIM-RT-EAPP1-E -001", "CIM-RT-EAPP1-E -007"})

    def test_run_preprocess_pipeline_filters_rows_by_normalized_selected_germplasm_names(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_filtered_normalized_") as temp_dir:
            run_dir = Path(temp_dir) / "preprocess"
            phase02_dir = run_dir / "phase02_climate_enriched"
            phase02_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")

            outputs = run_preprocess_pipeline(
                TEST_WORKBOOK,
                run_dir,
                selected_germplasm_names=["BABYCON", "OVERRECYCLEDBAZOOKA"],
            )

            headers, rows = load_sheet_rows(outputs.input_workbook)
            name_index = headers.index("Name")
            names = {row[name_index] for row in rows if row[name_index]}

            self.assertIn("Over recycled Bazooka", names)
            self.assertTrue({"Babycon", "Baby Con"} & names)
            self.assertLessEqual(names, {"Babycon", "Baby Con", "Over recycled Bazooka"})

    def test_run_preprocess_pipeline_projects_selected_germplasm_across_all_markers(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_projected_") as temp_dir:
            run_dir = Path(temp_dir) / "preprocess"
            phase02_dir = run_dir / "phase02_climate_enriched"
            phase02_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(REFERENCE_NASA_CACHE, phase02_dir / "nasa_power_cache.json")

            outputs = run_preprocess_pipeline(
                TEST_WORKBOOK,
                run_dir,
                selected_germplasm_names=["CIM-RT-EAPP1-E -001"],
                selected_germplasm_projection_mode="all_markers",
            )

            original_headers, original_rows = load_sheet_rows(TEST_WORKBOOK)
            projected_headers, projected_rows = load_sheet_rows(outputs.input_workbook)
            name_index = projected_headers.index("Name")
            planting_date_index = projected_headers.index("Date of planting")
            plot_index = projected_headers.index("Plot")
            latitude_index = projected_headers.index("_GPS coordinates_latitude")
            original_plot_index = original_headers.index("Plot")
            original_latitude_index = original_headers.index("_GPS coordinates_latitude")
            original_name_index = original_headers.index("Name")
            original_harvesting_index = original_headers.index("Date_of_harvesting")
            original_grain_yield_index = original_headers.index("Grain Yield (T/Ha)")
            template_row = next(
                row for row in original_rows if row[original_name_index] == "CIM-RT-EAPP1-E -001"
            )
            template_planting_date = template_row[planting_date_index]
            expected_base_rows = [
                row
                for row in original_rows
                if row[planting_date_index]
                and row[original_harvesting_index]
                and (
                    parse_numeric_text(row[original_grain_yield_index]) is None
                    or parse_numeric_text(row[original_grain_yield_index]) > 0
                )
            ]

            bridge_row_limit = get_prediction_bridge_dev_row_limit()
            expected_row_count = (
                min(len(expected_base_rows), bridge_row_limit)
                if bridge_row_limit
                else len(expected_base_rows)
            )

            self.assertEqual(len(projected_rows), expected_row_count)
            self.assertEqual({row[name_index] for row in projected_rows if row[name_index]}, {"CIM-RT-EAPP1-E -001"})
            self.assertEqual(
                {row[planting_date_index] for row in projected_rows if row[planting_date_index]},
                {template_planting_date},
            )
            self.assertEqual(
                [row[plot_index] for row in projected_rows[:25]],
                [row[original_plot_index] for row in expected_base_rows[:25]],
            )
            self.assertEqual(
                [row[latitude_index] for row in projected_rows[:25]],
                [row[original_latitude_index] for row in expected_base_rows[:25]],
            )

    def test_saved_model_manual_bbox_prediction_mode_is_detected(self) -> None:
        self.assertTrue(
            is_saved_model_manual_bbox_prediction(
                climate_scope="regional_manual",
                selected_model_id="model-123",
            )
        )
        self.assertFalse(
            is_saved_model_manual_bbox_prediction(
                climate_scope="point",
                selected_model_id="model-123",
            )
        )

    def test_phase01_can_preserve_blank_yield_and_coordinates_for_saved_model_manual_bbox(self) -> None:
        workbook = load_workbook(TEST_WORKBOOK, read_only=True, data_only=True)
        try:
            worksheet = workbook.worksheets[0]
            rows = worksheet.iter_rows(values_only=True)
            headers = [normalize_cell_value(value) for value in next(rows)]
            first_row = list(next(rows))
        finally:
            workbook.close()

        header_to_index = {header: index for index, header in enumerate(headers)}
        record = {
            header: first_row[index] if index < len(first_row) else None
            for index, header in enumerate(headers)
        }
        for header in (
            "Country",
            "GPS coordinates",
            "_GPS coordinates_latitude",
            "_GPS coordinates_longitude",
            "Date of planting",
            "Date of thinning",
            "Date_of_harvesting",
            "Grain Yield (T/Ha)",
        ):
            record[header] = ""
        for header in ("rank1", "rank2", "rank3", "rank4", "Frank1", "Frank2", "Frank3", "Frank4"):
            if header in header_to_index:
                record[header] = ""

        phase01_headers, phase01_records, metadata = ea_pipeline.build_phase01_prepared(
            headers,
            [record],
            allow_missing_yield=True,
            allow_missing_coordinates=True,
        )

        self.assertIn("Name", phase01_headers)
        self.assertEqual(len(phase01_records), 1)
        self.assertEqual(metadata["excluded_missing_yield"], 0)
        self.assertEqual(metadata["wrong_coordinates"], 0)

    def test_point_saved_model_projection_preserves_base_location_date_and_yield_columns(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_point_projection_") as temp_dir:
            projected_workbook = Path(temp_dir) / "projected.xlsx"
            project_normalized_test_workbook_across_all_markers(
                TEST_WORKBOOK,
                projected_workbook,
                ["CIM-RT-EAPP1-E -001"],
                template_workbook=TEST_WORKBOOK,
                base_override_columns=POINT_SAVED_MODEL_BASE_OVERRIDE_COLUMNS,
            )

            original_headers, original_rows = load_sheet_rows(TEST_WORKBOOK)
            projected_headers, projected_rows = load_sheet_rows(projected_workbook)
            original_first = original_rows[0]
            projected_first = projected_rows[0]
            for header in (
                "Country",
                "GPS coordinates",
                "_GPS coordinates_latitude",
                "_GPS coordinates_longitude",
                "_GPS coordinates_altitude",
                "_GPS coordinates_precision",
                "Date of planting",
                "Date of thinning",
                "Date_of_harvesting",
                "Grain Yield (T/Ha)",
            ):
                original_value = original_first[original_headers.index(header)]
                projected_value = projected_first[projected_headers.index(header)]
                self.assertEqual(projected_value, original_value, header)

    def test_manual_bbox_projection_clears_location_date_yield_and_soil_columns(self) -> None:
        with tempfile.TemporaryDirectory(prefix="app_preprocess_bbox_clear_") as temp_dir:
            cleared_workbook = Path(temp_dir) / "bbox_cleared.xlsx"
            clear_normalized_workbook_columns(
                TEST_WORKBOOK,
                cleared_workbook,
                columns_to_clear=MANUAL_BBOX_CLEAR_COLUMNS,
            )

            cleared_headers, cleared_rows = load_sheet_rows(cleared_workbook)
            cleared_first = cleared_rows[0]
            cleared_map = {
                header: (cleared_first[index] if index < len(cleared_first) else "")
                for index, header in enumerate(cleared_headers)
            }
            for header in MANUAL_BBOX_CLEAR_COLUMNS:
                value = cleared_map.get(header, "")
                self.assertEqual(value, "", header)


if __name__ == "__main__":
    unittest.main()
