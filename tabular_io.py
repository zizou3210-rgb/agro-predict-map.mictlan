"""CSV support for large pipeline tables, with explicit identifier preservation."""
from __future__ import annotations

import csv
import os
import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd


CSV_MIN_ROWS = 50_000
IDENTIFIER_COLUMN = re.compile(r'(^|[ _])(id|code|identifier)($|[ _])|^Farm$|^Name$|profile|grid.*cell.*id|entrycode|ce row', re.I)


def csv_threshold():
    value = int(os.environ.get('APP_LARGE_TABLE_MIN_ROWS', CSV_MIN_ROWS))
    if value < 1:
        raise ValueError('APP_LARGE_TABLE_MIN_ROWS must be at least 1.')
    return value


def large_table_path(path, row_count):
    path = Path(path)
    return path.with_suffix('.csv') if row_count >= csv_threshold() else path


def read_headers(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            return [value.strip() for value in next(csv.reader(handle), [])]
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [str(v).strip() for v in next(book.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True), ()) if v is not None]
    finally:
        book.close()


def read_records(path, progress_callback=None):
    path = Path(path)
    total = path.stat().st_size
    records = []
    if progress_callback:
        progress_callback(0, total)
    with path.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.reader(handle)
        headers = [value.strip() for value in next(reader, [])]
        for values in reader:
            if len(values) > len(headers):
                raise ValueError(f'CSV row has {len(values)} fields; expected {len(headers)}: {path.name}')
            records.append({header: (values[index] if index < len(values) else '').strip() for index, header in enumerate(headers)})
            if progress_callback:
                progress_callback(min(handle.buffer.tell(), total), total)
    if progress_callback:
        progress_callback(total, total, True)
    return headers, records


def _csv_value(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def write_rows(path, headers, rows, total, progress_callback=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if progress_callback:
        progress_callback(0, total)
    with path.open('w', encoding='utf-8', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for completed, row in enumerate(rows, 1):
            writer.writerow([_csv_value(value) for value in row])
            if progress_callback:
                progress_callback(completed, total)
    if progress_callback:
        progress_callback(total, total, True)


def read_dataframe(path, *, dtype=None, **kwargs):
    path = Path(path)
    if path.suffix.lower() != ".csv":
        frame = pd.read_excel(path, dtype=object, keep_default_na=False, na_filter=False, **kwargs)
    else:
        frame = pd.read_csv(path, dtype=object, keep_default_na=False, na_filter=False, **kwargs)
    # Literal NA/NULL and numeric-looking Farm/row identifiers are values, not
    # missing data. Only empty fields become missing values for pandas math.
    frame = frame.mask(frame.eq(''), float('nan'))
    if dtype is not None:
        return frame.astype(dtype)
    for column in frame.columns:
        if IDENTIFIER_COLUMN.search(str(column)):
            continue
        values = frame[column]
        present = values.notna()
        if present.any():
            numeric = pd.to_numeric(values, errors='coerce')
            if numeric[present].notna().all():
                frame[column] = numeric
    return frame


def estimated_manual_grid_rows(profile_count, bounds, cell_size_km):
    from pipeline.nasa_country_region import validate_regional_bounds, _build_manual_grid_steps
    region = validate_regional_bounds(*bounds, label='CSV planning')
    lat_step, lon_step = _build_manual_grid_steps(region, cell_size_km)
    counts = []
    for minimum, maximum, step in [(region.latitude_min, region.latitude_max, lat_step), (region.longitude_min, region.longitude_max, lon_step)]:
        count = 0
        while minimum < maximum:
            minimum = min(minimum + step, maximum)
            count += 1
        counts.append(count)
    return profile_count * counts[0] * counts[1]
